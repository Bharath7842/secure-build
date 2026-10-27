#!/usr/bin/env python3
"""Join the npg-pakistan review exports by finding identity and write the labels.

Three exports describe the same scan and are not interchangeable snapshots:

  R  the developer review workbook (xlsx) — original SB verdict and DEV label
  S  the scanner's own Excel export (xlsx) — the verdict it displayed
  P  the scanner's PDF (optional) — its numbered finding headings and pills

They are joined on what a finding *is* — (file path, CWE, title) — never on row
number: the PDF holds four findings the workbooks do not, so its numbering shifts
after #80 and a positional join would mislabel everything past it. The output is
`<out>.jsonl`, one record per finding *per affected file* (a grouped row expands to
one record per listed file, all sharing a `group_ref`), carrying `scanner_verdict_S`,
`scanner_label_P` and `dev_label_R` as separate fields. Every record is
`adjudication: dev_label_only`: developer review is evidence, not ground truth, and
nothing is adjudicated by this import. The file holds file paths, titles, CWEs and
labels only — no source code, and not the developers' free-text explanations.

The historical exports predate `occurrence_id`, so the join is heuristic. Each record
says how sure it is: `exact` (same CWE, same title, overlapping paths), `fuzzy` (a
truncated title or a near match — flagged `needs_manual_review`, never auto-resolved)
or `none` (the row exists in one source only).

    python scripts/accuracy/import_review_labels.py --review R.xlsx --scanner S.xlsx \\
        [--pdf P.pdf] --out backend/tests/accuracy_corpus/_labels/npg-pakistan.jsonl

Prints the join summary: rows that joined across every source, rows present in only
one, and the R/S verdict differences. Keep that table in the PR description when the
labels are refreshed.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

ADJUDICATION = "dev_label_only"

# Where the review workbook keeps what we need (columns A-J, rows 2-292 in the npg-pakistan file).
DEFAULT_REVIEW_COLUMNS = {
    "number": "A", "title": "B", "location": "C", "cwe": "D", "original": "F", "dev": "G",
}
DEFAULT_REVIEW_SHEET = "Secure Build Findings"
DEFAULT_SCANNER_SHEET = "Findings"
# Exports since the CWE roll-up list one row per CWE and verdict in "Findings";
# the per-issue rows this tool joins on are in "Issue Groups".
ISSUE_GROUPS_SHEET = "Issue Groups"

# The legacy 14-column Findings sheet of the scanner export, for a file with no header row.
_LEGACY_SCANNER_COLUMNS = (
    "number", "severity", "verdict", "owasp", "cwe", "cvss", "confidence", "files", "title",
    "location", "data_flow", "risk", "recommendation", "example_fix",
)
_SCANNER_HEADERS = {
    "#": "number", "severity": "severity", "scanner verdict": "verdict", "verdict": "verdict",
    "verdict (raw)": "verdict_raw", "cwe": "cwe", "confidence": "confidence", "files": "files",
    "issue": "title", "location(s)": "location", "occurrence id": "occurrence_id",
}

_PATH_RE = re.compile(r"[\w$@.+\-/\\]+\.[A-Za-z0-9]{1,6}")
_CWE_RE = re.compile(r"CWE[-\s]?(\d+)", re.IGNORECASE)
_PDF_HEADING_RE = re.compile(
    r"FINDING\s+(\d+)\s+(?P<title>.+?)\s+(?:CRITICAL|HIGH|MEDIUM|LOW|INFO)\s+"
    r"(?P<label>True Positive|False Positive|Needs Review|Unverified \(regex\)|Unresolved(?: \([^)]*\))?|Unknown \([^)]*\))"
)

_DEV_SLUGS = {
    "false positive": "false_positive",
    "real security issue": "real_security_issue",
    "real, but low risk": "real_low_risk",
    "real but low risk": "real_low_risk",
    "code quality": "code_quality",
}


# --- normalisation ------------------------------------------------------------

def norm_title(text: Any) -> str:
    """Lower-case alphanumerics only, a trailing truncation mark dropped."""
    text = re.sub(r"(…|\.\.\.)\s*$", "", str(text or ""))
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def norm_cwe(text: Any) -> str | None:
    match = _CWE_RE.search(str(text or ""))
    if match:
        return f"CWE-{int(match.group(1))}"
    digits = re.fullmatch(r"\s*(\d+)\s*", str(text or ""))
    return f"CWE-{int(digits.group(1))}" if digits else None


def norm_path(path: str) -> str:
    return re.sub(r"^\./", "", path.strip().replace("\\", "/")).lower()


def norm_verdict(text: Any) -> str | None:
    """A scanner verdict label, in stable words. "Needs Review" is kept apart from
    "false_positive": it is what the legacy export *displayed* for a false positive."""
    t = str(text or "").strip().lower().replace("_", " ").replace("-", " ")
    if not t:
        return None
    if "true positive" in t or t == "tp":
        return "true_positive"
    if "false positive" in t or t == "fp":
        return "false_positive"
    if "needs review" in t:
        return "needs_review"
    if "unverified" in t or "unresolved" in t:
        return "unverified"
    return re.sub(r"\s+", "_", t)


def dev_label(text: Any) -> str | None:
    t = str(text or "").strip().lower()
    if not t:
        return None
    return _DEV_SLUGS.get(t) or re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def parse_locations(text: Any) -> list[tuple[str, int | None]]:
    """(path, line) pairs from a location cell: the first paragraph, one path per
    line (the scanner export appends a code excerpt after a blank line)."""
    paragraph = re.split(r"\n\s*\n", str(text or "").strip(), maxsplit=1)[0]
    out: list[tuple[str, int | None]] = []
    for raw in paragraph.splitlines():
        for token in _PATH_RE.findall(raw):
            path, _, line = token.partition(":")
            out.append((path, int(line) if line.isdigit() else None))
        # "path:12" — the regex above stops at the extension, so look for the line separately.
        for match in re.finditer(r"([\w$@.+\-/\\]+\.[A-Za-z0-9]{1,6}):(\d+)", raw):
            out = [(p, ln) for p, ln in out if p != match.group(1)]
            out.append((match.group(1), int(match.group(2))))
    seen: set[tuple[str, int | None]] = set()
    unique = []
    for item in out:
        if item not in seen:
            seen.add(item)
            unique.append(item)
    return unique


# --- readers ------------------------------------------------------------------

@dataclass
class Row:
    """One finding as one source described it."""

    source: str
    ordinal: int | None
    title: str
    cwe: str | None
    paths: list[tuple[str, int | None]] = field(default_factory=list)
    verdict: str | None = None          # S / P: the displayed verdict; R: the original SB verdict
    dev: str | None = None              # R only: the developer's result
    extra: dict[str, Any] = field(default_factory=dict)

    @property
    def path_set(self) -> set[str]:
        return {norm_path(p) for p, _ in self.paths}


def _column_index(letter: str) -> int:
    letter = letter.strip().upper()
    index = 0
    for ch in letter:
        index = index * 26 + (ord(ch) - ord("A") + 1)
    return index - 1


def parse_columns(spec: str | None) -> dict[str, str]:
    columns = dict(DEFAULT_REVIEW_COLUMNS)
    for part in (spec or "").split(","):
        if part.strip():
            key, _, letter = part.partition("=")
            columns[key.strip()] = letter.strip()
    return columns


def read_review(path: Path, columns: dict[str, str] | None = None, sheet: str | None = None) -> list[Row]:
    from openpyxl import load_workbook

    columns = columns or dict(DEFAULT_REVIEW_COLUMNS)
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet or DEFAULT_REVIEW_SHEET] if (sheet or DEFAULT_REVIEW_SHEET) in workbook.sheetnames \
        else workbook[workbook.sheetnames[0]]
    index = {name: _column_index(letter) for name, letter in columns.items()}
    rows: list[Row] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        def cell(name: str) -> Any:
            i = index.get(name)
            return values[i] if i is not None and i < len(values) else None

        title = str(cell("title") or "").strip()
        if not title:
            continue
        number = cell("number")
        rows.append(Row(
            source="R", ordinal=int(number) if str(number).strip().isdigit() else None, title=title,
            cwe=norm_cwe(cell("cwe")), paths=parse_locations(cell("location")),
            verdict=norm_verdict(cell("original")), dev=dev_label(cell("dev")),
        ))
    return rows


def read_scanner(path: Path, sheet: str | None = None) -> list[Row]:
    """The scanner Excel export. Header-driven (so the legacy and the current
    layout both work); a headerless legacy sheet falls back to its fixed columns."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    name = sheet or (ISSUE_GROUPS_SHEET if ISSUE_GROUPS_SHEET in workbook.sheetnames else DEFAULT_SCANNER_SHEET)
    worksheet = workbook[name] if name in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows_iter = worksheet.iter_rows(values_only=True)
    first = next(rows_iter, None)
    if first is None:
        return []
    mapping = {i: _SCANNER_HEADERS[str(v).strip().lower()]
               for i, v in enumerate(first) if v is not None and str(v).strip().lower() in _SCANNER_HEADERS}
    data: Iterable[tuple] = rows_iter
    if "title" not in mapping.values():  # no header row: the legacy fixed layout
        mapping = {i: field_name for i, field_name in enumerate(_LEGACY_SCANNER_COLUMNS)}
        data = [first, *rows_iter]
    out: list[Row] = []
    for values in data:
        record = {field_name: values[i] for i, field_name in mapping.items() if i < len(values)}
        title = str(record.get("title") or "").strip()
        if not title:
            continue
        number = record.get("number")
        out.append(Row(
            source="S", ordinal=int(number) if str(number).strip().isdigit() else None, title=title,
            cwe=norm_cwe(record.get("cwe")), paths=parse_locations(record.get("location")),
            verdict=norm_verdict(record.get("verdict_raw") or record.get("verdict")),
            extra={"occurrence_id": record.get("occurrence_id")},
        ))
    return out


def read_pdf(path: Path) -> list[Row]:
    """The scanner PDF's numbered finding headings, with the verdict pill, the CWE
    and the paths listed in each finding's block."""
    from pypdf import PdfReader

    text = " ".join(" ".join((page.extract_text() or "") for page in PdfReader(str(path)).pages).split())
    matches = list(_PDF_HEADING_RE.finditer(text))
    rows: list[Row] = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        block = text[match.end():end]
        cwe_match = _CWE_RE.search(block)
        rows.append(Row(
            source="P", ordinal=int(match.group(1)), title=match.group("title").strip(),
            cwe=f"CWE-{int(cwe_match.group(1))}" if cwe_match else None,
            paths=[(p, None) for p in dict.fromkeys(_PATH_RE.findall(block)) if "/" in p],
            verdict=norm_verdict(match.group("label")),
        ))
    return rows


# --- the join -----------------------------------------------------------------

EXACT, FUZZY, NONE = "exact", "fuzzy", "none"


def _paths_agree(a: Row, b: Row) -> bool:
    """True when either side lists no path, or any path of one is a suffix of one of the other."""
    if not a.paths or not b.paths:
        return True
    for x in a.path_set:
        for y in b.path_set:
            if x == y or x.endswith("/" + y) or y.endswith("/" + x):
                return True
    return False


def _classify(hub: Row, other: Row) -> str | None:
    """`exact`, `fuzzy` or None for a hub row against a candidate."""
    if hub.cwe and other.cwe and hub.cwe != other.cwe:
        return None
    a, b = norm_title(hub.title), norm_title(other.title)
    if not a or not b:
        return None
    if a == b:
        return EXACT if _paths_agree(hub, other) else None
    shorter, longer = sorted((a, b), key=len)
    if len(shorter) >= 25 and longer.startswith(shorter):
        return FUZZY if _paths_agree(hub, other) else None
    if SequenceMatcher(None, a, b).ratio() >= 0.85 and _paths_agree(hub, other):
        return FUZZY
    return None


def _claim(hub: Row, pool: list[Row], claimed: set[int]) -> tuple[Row | None, str]:
    """The best unclaimed candidate for `hub`. A tie between candidates that look
    identical is broken by order but never reported as exact."""
    exact: list[Row] = []
    fuzzy: list[Row] = []
    for candidate in pool:
        if id(candidate) in claimed:
            continue
        kind = _classify(hub, candidate)
        if kind == EXACT:
            exact.append(candidate)
        elif kind == FUZZY:
            fuzzy.append(candidate)
    if len(exact) == 1:
        return exact[0], EXACT
    if exact:
        exact.sort(key=lambda c: -len(hub.path_set & c.path_set))
        return exact[0], FUZZY  # ambiguous: several rows are indistinguishable
    if fuzzy:
        return fuzzy[0], FUZZY
    return None, NONE


@dataclass
class Joined:
    r: Row | None = None
    s: Row | None = None
    p: Row | None = None
    confidence: str = NONE  # the weakest link among the joins made

    @property
    def sources(self) -> list[str]:
        return [name for name, row in (("R", self.r), ("S", self.s), ("P", self.p)) if row]


def join(review: list[Row], scanner: list[Row], pdf: list[Row] | None = None) -> list[Joined]:
    """Join R and P onto S (the richest source). Rows present in a source only
    stay alone. Order of the output follows S, then R-only, then P-only."""
    claimed: set[int] = set()
    joined = [Joined(s=row) for row in scanner]
    for source, rows, attr in (("R", review, "r"), ("P", pdf or [], "p")):
        for entry in joined:
            match, confidence = _claim(entry.s, rows, claimed)
            if match is None:
                continue
            claimed.add(id(match))
            setattr(entry, attr, match)
            # The record is only as sure as its weakest join.
            entry.confidence = FUZZY if FUZZY in (confidence, entry.confidence) else EXACT
    for rows, attr in ((review, "r"), (pdf or [], "p")):
        for row in rows:
            if id(row) not in claimed:
                joined.append(Joined(**{attr: row}))
    return joined


# --- output -------------------------------------------------------------------

def _group_ref(entry: Joined) -> str:
    for name, row in (("S", entry.s), ("R", entry.r), ("P", entry.p)):
        if row is not None:
            return f"{name}:{row.ordinal if row.ordinal is not None else norm_title(row.title)[:40]}"
    return "?"


def to_records(joined: list[Joined]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for entry in joined:
        hub = entry.s or entry.r or entry.p
        paths = (entry.s.paths if entry.s else None) or (entry.r.paths if entry.r else None) \
            or (entry.p.paths if entry.p else None) or [(None, None)]
        ref = _group_ref(entry)
        multi_source = len(entry.sources) > 1
        confidence = entry.confidence if multi_source else NONE
        for index, (file_path, line) in enumerate(paths):
            records.append({
                "id": f"{ref}#{index}",
                "group_ref": ref,
                "file_path": file_path,
                "line": line,
                "title": hub.title,
                "cwe": hub.cwe,
                "rule_id": None,  # the historical exports carry no rule id
                "sources": entry.sources,
                "join_confidence": confidence,
                "needs_manual_review": confidence == FUZZY,
                "ordinal_R": entry.r.ordinal if entry.r else None,
                "ordinal_S": entry.s.ordinal if entry.s else None,
                "ordinal_P": entry.p.ordinal if entry.p else None,
                "scanner_verdict_R": entry.r.verdict if entry.r else None,
                "scanner_verdict_S": entry.s.verdict if entry.s else None,
                "scanner_label_P": entry.p.verdict if entry.p else None,
                "dev_label_R": entry.r.dev if entry.r else None,
                "occurrence_id": (entry.s.extra.get("occurrence_id") if entry.s else None) or None,
                "adjudication": ADJUDICATION,
            })
    return records


def summarize(review: list[Row], scanner: list[Row], pdf: list[Row] | None, joined: list[Joined]) -> dict[str, Any]:
    have = [name for name, rows in (("R", review), ("S", scanner), ("P", pdf)) if rows]
    complete = [e for e in joined if len(e.sources) == len(have)]
    only: dict[str, list[str]] = {"R": [], "S": [], "P": []}
    for entry in joined:
        if len(entry.sources) == 1:
            row = entry.r or entry.s or entry.p
            only[entry.sources[0]].append(f"#{row.ordinal} {row.title[:70]}")
    differences = sum(
        1 for e in joined
        if e.r is not None and e.s is not None and e.r.verdict != e.s.verdict
    )
    return {
        "rows": {"R": len(review), "S": len(scanner), "P": len(pdf or [])},
        "joined_across_all_sources": len(complete),
        "join_confidence": dict(Counter(e.confidence for e in joined if len(e.sources) > 1)),
        "only_in_one_source": only,
        "r_s_verdict_differences": differences,
    }


def format_summary(summary: dict[str, Any]) -> str:
    lines = ["Join summary", "============"]
    rows = summary["rows"]
    lines.append(f"rows read:                     R={rows['R']}  S={rows['S']}  P={rows['P']}")
    lines.append(f"joined across every source:    {summary['joined_across_all_sources']}")
    lines.append(f"join confidence (multi-source): {summary['join_confidence'] or '-'}")
    lines.append(f"R/S verdict differences:       {summary['r_s_verdict_differences']}")
    for source, items in summary["only_in_one_source"].items():
        lines.append(f"present only in {source}:             {len(items)}")
        for item in items[:20]:
            lines.append(f"    {item}")
        if len(items) > 20:
            lines.append(f"    ... and {len(items) - 20} more")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--review", required=True, type=Path, help="developer review workbook (R)")
    parser.add_argument("--scanner", required=True, type=Path, help="scanner Excel export (S)")
    parser.add_argument("--pdf", type=Path, help="scanner PDF (P), optional")
    parser.add_argument("--out", required=True, type=Path, help="output .jsonl")
    parser.add_argument("--review-sheet", help=f"sheet of R (default {DEFAULT_REVIEW_SHEET!r})")
    parser.add_argument(
        "--review-columns",
        help="column letters in R, e.g. number=A,title=B,location=C,cwe=D,original=F,dev=G",
    )
    args = parser.parse_args(argv)

    review = read_review(args.review, parse_columns(args.review_columns), args.review_sheet)
    scanner = read_scanner(args.scanner)
    pdf = read_pdf(args.pdf) if args.pdf else None
    joined = join(review, scanner, pdf)
    records = to_records(joined)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, sort_keys=True) + "\n")
    print(format_summary(summarize(review, scanner, pdf, joined)))
    print(f"\nwrote {len(records)} record(s) to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
