"""The scanner's verdict and the human review state, side by side and by their
real names, in the API and in both exports (SAST-00-030)."""

import logging
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app import db as db_module
from app.config import settings
from app.db import init_db
from app.main import app
from app.models import Finding, Scan
from app.services import feedback, scan_runner
from app.services.scan_export import _display_verdict

from tests.auth_helpers import ADMIN_HEADERS, make_user

pytestmark = pytest.mark.integration

HDR = ADMIN_HEADERS


# --- labels -----------------------------------------------------------------

@pytest.mark.parametrize("value,reason,expected", [
    ("true_positive", None, "True Positive"),
    ("TRUE_POSITIVE", None, "True Positive"),
    ("false_positive", None, "False Positive"),
    ("FALSE-POSITIVE", None, "False Positive"),
    ("unverified", None, "Unresolved"),
    ("unverified", "not_reviewed_regex", "Unresolved (regex, not reviewed)"),
    ("unverified", "llm_missing_verdict", "Unresolved (model gave no verdict)"),
    ("unverified", "something_new", "Unresolved (something_new)"),
    (None, None, "—"),
    ("", None, "—"),
    ("weird", None, "Unknown (weird)"),
    ("needs_review", None, "Unknown (needs_review)"),
])
def test_display_verdict_labels(value, reason, expected):
    assert _display_verdict(value, reason) == expected


def test_a_verdict_outside_the_contract_is_logged(caplog):
    with caplog.at_level(logging.WARNING, logger="app.services.scan_export"):
        _display_verdict("weird")
    assert "outside the contract" in caplog.text


def test_needs_review_is_not_a_verdict_label_for_anything():
    for verdict in ("true_positive", "false_positive", "unverified", "", None, "weird"):
        assert "Needs Review" not in _display_verdict(verdict)


# --- review state -----------------------------------------------------------

class _Review:
    def __init__(self, reviewer_verdict=None, is_override=False, resolution_status="open"):
        self.reviewer_verdict = reviewer_verdict
        self.is_override = is_override
        self.resolution_status = resolution_status


@pytest.mark.parametrize("review,scanner,expected", [
    (None, "false_positive", "pending"),
    (_Review("needs_more_info"), "false_positive", "pending"),
    (_Review("false_positive"), "false_positive", "agreed"),
    (_Review("confirmed_issue"), "true_positive", "agreed"),
    (_Review("false_positive", is_override=True), "true_positive", "overridden"),
    (_Review("confirmed_issue", is_override=True), "false_positive", "overridden"),
    (_Review("confirmed_issue", resolution_status="resolved"), "true_positive", "resolved"),
    (_Review("confirmed_issue", resolution_status="risk_accepted"), "true_positive", "resolved"),
    (_Review("false_positive", is_override=True, resolution_status="resolved"), "true_positive", "resolved"),
])
def test_review_state(review, scanner, expected):
    assert feedback.review_state(review, scanner) == expected


# --- the API and the exports ------------------------------------------------

@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("USE_MOCK_ENGINE", "true")
    settings.use_mock_engine = True
    settings.scan_logs_dir = str(tmp_path / "scan_logs")
    init_db(f"sqlite:///{(tmp_path / 'display.db').as_posix()}")
    return TestClient(app)


@pytest.fixture()
def seeded(client, monkeypatch):
    """One finding per verdict case: confirmed, scanner-dismissed, unresolved (regex),
    unresolved (the model gave no verdict)."""
    monkeypatch.setattr(scan_runner, "execute_scan", lambda scan_id: None)
    project_id = client.post(
        "/projects", headers=HDR,
        json={"name": "verdict-display", "default_git_url": "https://github.com/org/repo.git"},
    ).json()["id"]
    scan_id = client.post(
        "/scans", headers=HDR, json={"project_id": project_id, "commit_sha": "vd"}
    ).json()["scan_id"]
    with db_module.SessionLocal() as db:
        db.get(Scan, scan_id).status = "success"
        rows = {
            "tp": Finding(scan_id=scan_id, severity="HIGH", title="SQL injection", cwe="CWE-89",
                          rule_id="sqli", file_path="A.java", line=10, verdict="true_positive",
                          kind="llm", issue_class="security", confidence=0.9),
            "fp": Finding(scan_id=scan_id, severity="MEDIUM", title="Weak hash", cwe="CWE-328",
                          rule_id="weak-hash", file_path="B.java", line=20, verdict="false_positive",
                          kind="llm", issue_class="security", confidence=0.3),
            "regex": Finding(scan_id=scan_id, severity="LOW", title="Serializable", cwe="CWE-502",
                             rule_id="svuid", group_by_rule=True, file_path="C.java", line=1,
                             verdict="unverified", unresolved_reason="not_reviewed_regex",
                             kind="regex", issue_class="quality"),
            "missing": Finding(scan_id=scan_id, severity="HIGH", title="Command injection",
                               cwe="CWE-78", rule_id="llm-extra", file_path="D.java", line=5,
                               verdict="unverified", unresolved_reason="llm_missing_verdict",
                               kind="llm", issue_class="security"),
        }
        db.add_all(rows.values())
        db.commit()
        ids = {k: v.id for k, v in rows.items()}
    return scan_id, ids


def _findings(client, scan_id):
    body = client.get(f"/scans/{scan_id}/findings", headers=HDR).json()
    return {f["title"]: f for f in body["findings"]}


def _override(client, scan_id, finding_id):
    headers = make_user("reviewer@test.local", role="admin")[1]
    return client.post(
        f"/scans/{scan_id}/findings/{finding_id}/review", headers=headers,
        json={"decision": "false_positive", "comment": "input is validated upstream"},
    )


def test_every_finding_starts_with_review_pending(client, seeded):
    scan_id, _ids = seeded
    assert {f["review_state"] for f in _findings(client, scan_id).values()} == {"pending"}


def test_an_override_reads_overridden_and_keeps_the_scanners_verdict(client, seeded):
    scan_id, ids = seeded
    assert _override(client, scan_id, ids["tp"]).status_code == 200
    sqli = _findings(client, scan_id)["SQL injection"]
    assert sqli["review_state"] == "overridden"
    assert sqli["verdict"] == "false_positive"      # the effective verdict
    assert sqli["llm_verdict"] == "true_positive"   # what the scanner decided


def test_the_xlsx_has_the_five_verdict_columns_and_raw_values_equal_the_db(client, seeded):
    scan_id, ids = seeded
    _override(client, scan_id, ids["tp"])
    resp = client.get(f"/scans/{scan_id}/export.xlsx", headers=HDR)
    sheet = load_workbook(BytesIO(resp.content))["Findings"]
    headers = [c.value for c in sheet[1]]
    for column in ("Scanner Verdict", "Review State", "Issue Class", "Unresolved Reason", "Verdict (raw)"):
        assert column in headers
    assert "Verdict" not in headers

    rows = {
        r[headers.index("Issue")].value: {h: r[i].value for i, h in enumerate(headers)}
        for r in sheet.iter_rows(min_row=2)
    }
    with db_module.SessionLocal() as db:
        stored = {f.title: f.verdict for f in db.query(Finding).filter(Finding.scan_id == scan_id)}
    for title, row in rows.items():
        assert row["Verdict (raw)"] == stored[title], title

    assert rows["SQL injection"]["Scanner Verdict"] == "True Positive"
    assert rows["SQL injection"]["Review State"] == "overridden → False Positive"
    assert rows["Weak hash"]["Scanner Verdict"] == "False Positive"
    assert rows["Weak hash"]["Review State"] == "pending"
    assert rows["Serializable"]["Scanner Verdict"] == "Unresolved (regex, not reviewed)"
    assert rows["Serializable"]["Issue Class"] == "quality"
    assert rows["Serializable"]["Unresolved Reason"] == "not_reviewed_regex"
    assert rows["Command injection"]["Scanner Verdict"] == "Unresolved (model gave no verdict)"
    assert rows["Command injection"]["Unresolved Reason"] == "llm_missing_verdict"


def _pdf_text(client, scan_id):
    resp = client.get(f"/scans/{scan_id}/export.pdf", headers=HDR)
    assert resp.status_code == 200
    return " ".join(" ".join(p.extract_text() for p in PdfReader(BytesIO(resp.content)).pages).split())


def test_the_pdf_never_says_needs_review_and_names_the_buckets(client, seeded):
    scan_id, _ids = seeded
    text = _pdf_text(client, scan_id)
    assert "Needs Review" not in text and "Dismissed" not in text
    assert "False Positive" in text
    for bucket in ("Confirmed:", "Unresolved:", "False Positive:"):
        assert bucket in text
    assert "Scanner: False Positive · Review: pending" in text


def test_the_pdf_shows_the_scanner_verdict_for_an_overridden_finding(client, seeded):
    scan_id, ids = seeded
    _override(client, scan_id, ids["tp"])
    text = _pdf_text(client, scan_id)
    assert "Scanner: True Positive · Review: overridden → False Positive" in text


def test_the_json_keeps_the_raw_enum_in_verdict(client, seeded):
    scan_id, _ids = seeded
    found = _findings(client, scan_id)
    assert {f["verdict"] for f in found.values()} == {"true_positive", "false_positive", "unverified"}
