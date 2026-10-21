"""Occurrence ids and the version stamp as the API and exports show them
(SAST-00-040)."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app import db as db_module
from app.db import init_db
from app.main import app
from app.models import Finding, Scan
from app.services import scan_runner

from tests.auth_helpers import ADMIN_HEADERS

pytestmark = pytest.mark.integration

HDR = ADMIN_HEADERS
ID_A, ID_B, ID_C = "a" * 64, "b" * 64, "c" * 64


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("USE_MOCK_ENGINE", "true")
    settings.use_mock_engine = True
    settings.scan_logs_dir = str(tmp_path / "scan_logs")
    init_db(f"sqlite:///{(tmp_path / 'surf.db').as_posix()}")
    return TestClient(app)


def _project(client, name):
    return client.post(
        "/projects", headers=HDR,
        json={"name": name, "default_git_url": "https://github.com/org/repo.git"},
    ).json()["id"]


def test_a_mock_scan_records_every_version_key(client):
    project_id = _project(client, "mock-versions")
    scan_id = client.post(
        "/scans", headers=HDR, json={"project_id": project_id, "commit_sha": "m1"}
    ).json()["scan_id"]
    with db_module.SessionLocal() as db:
        versions = db.get(Scan, scan_id).metadata_json["versions"]
    assert set(versions) == {
        "engine_commit", "rule_catalog_sha256", "result_cache_version",
        "system_prompt_sha256", "llm_provider", "llm_model",
    }
    assert versions["engine_commit"] == "mock"


@pytest.fixture()
def seeded(client, monkeypatch):
    monkeypatch.setattr(scan_runner, "execute_scan", lambda scan_id: None)
    project_id = _project(client, "occurrence-surfacing")
    scan_id = client.post(
        "/scans", headers=HDR, json={"project_id": project_id, "commit_sha": "o1"}
    ).json()["scan_id"]
    with db_module.SessionLocal() as db:
        scan = db.get(Scan, scan_id)
        scan.status = "success"
        db.add_all([
            Finding(scan_id=scan_id, severity="HIGH", title="SQL injection", cwe="CWE-89",
                    rule_id="sqli", file_path="A.java", line=10, verdict="true_positive",
                    kind="llm", occurrence_id=ID_A, identity_quality="exact"),
            # Two members of one grouped construct, in two files.
            Finding(scan_id=scan_id, severity="LOW", title="Serializable", cwe="CWE-502",
                    rule_id="svuid", group_by_rule=True, file_path="B.java", line=1,
                    verdict="unverified", unresolved_reason="not_reviewed_regex", kind="regex",
                    occurrence_id=ID_B, identity_quality="exact"),
            Finding(scan_id=scan_id, severity="LOW", title="Serializable", cwe="CWE-502",
                    rule_id="svuid", group_by_rule=True, file_path="C.java", line=1,
                    verdict="unverified", unresolved_reason="not_reviewed_regex", kind="regex",
                    occurrence_id=ID_C, identity_quality="file_level"),
        ])
        db.commit()
    return scan_id


def test_the_findings_api_returns_ids_and_a_groups_member_ids(client, seeded):
    body = client.get(f"/scans/{seeded}/findings", headers=HDR).json()
    by_title = {f["title"]: f for f in body["findings"]}
    assert by_title["SQL injection"]["occurrence_id"] == ID_A
    assert by_title["SQL injection"]["member_occurrence_ids"] == [ID_A]
    group = by_title["Serializable"]
    assert group["occurrence_count"] == 2
    assert group["member_occurrence_ids"] == [ID_B, ID_C]


def test_the_xlsx_occurrence_id_column_equals_the_api_value(client, seeded):
    api = {f["title"]: f for f in client.get(f"/scans/{seeded}/findings", headers=HDR).json()["findings"]}
    resp = client.get(f"/scans/{seeded}/export.xlsx", headers=HDR)
    assert resp.status_code == 200
    sheet = load_workbook(BytesIO(resp.content))["Findings"]
    headers = [c.value for c in sheet[1]]
    column = headers.index("Occurrence ID")
    by_issue = {row[headers.index("Issue")].value: row[column].value for row in sheet.iter_rows(min_row=2)}
    for title, finding in api.items():
        assert by_issue[title].split("\n") == finding["member_occurrence_ids"]


def test_the_pdf_shows_each_findings_occurrence_id(client, seeded):
    from pypdf import PdfReader

    resp = client.get(f"/scans/{seeded}/export.pdf", headers=HDR)
    assert resp.status_code == 200
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(resp.content)).pages)
    flattened = text.replace("\n", "").replace(" ", "")
    assert ID_A in flattened
    assert ID_B in flattened and ID_C in flattened
