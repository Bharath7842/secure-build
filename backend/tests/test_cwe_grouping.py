"""One row per CWE over the issue groups (`scan_export.group_by_cwe`), and the
exports that show it."""

from io import BytesIO

from openpyxl import load_workbook

from app.models import Finding, Project, Scan
from app.services.scan_export import (
    build_excel_report,
    count_findings,
    group_by_cwe,
    group_findings_by_issue,
)

_ids = iter(range(10_000))


def _finding(rule_id, cwe, verdict, path, *, severity="HIGH", matched="x", kind="regex", **extra):
    return Finding(
        id=f"f{next(_ids)}", scan_id="s1", kind=kind, rule_id=rule_id, cwe=cwe, verdict=verdict,
        severity=severity, title=extra.pop("title", rule_id), file_path=path, line=extra.pop("line", 3),
        matched_text=matched, issue_class="security", owasp_category="A03:2021-Injection",
        unresolved_reason=extra.pop("unresolved_reason", "not_reviewed_regex" if verdict == "unverified" else None),
        **extra,
    )


def _rows(findings):
    return {row["key"]: row for row in group_by_cwe(group_findings_by_issue(findings))}


def test_rules_sharing_a_cwe_roll_up_into_one_row_with_every_file():
    rows = _rows([
        _finding("java-sql-string-concat", "CWE-89", "true_positive", "a/A.java", matched='"..." + id'),
        _finding("java-sql-string-concat", "CWE-89", "true_positive", "a/B.java", matched='"..." + name'),
        _finding("java-jdbc-prepare-concat", "cwe 089", "true_positive", "a/A.java", severity="CRITICAL"),
        _finding("jpa-native-query", "89", "true_positive", "a/C.java"),
    ])
    assert list(rows) == ["CWE-89"]
    row = rows["CWE-89"]
    assert row["groups"] == 4 and row["occurrences"] == 4
    assert row["files"] == ["a/A.java", "a/B.java", "a/C.java"] and row["file_count"] == 3
    assert row["severity"] == "critical"
    assert row["by_bucket"]["true_positive"]["files"] == 3


def test_each_verdict_keeps_its_own_cell_and_its_own_groups():
    findings = [
        _finding("sqli", "CWE-89", "true_positive", "A.java"),
        _finding("sqli", "CWE-89", "unverified", "B.java", matched="y"),
        _finding("sqli", "CWE-89", "unverified", "C.java", matched="y"),
        _finding("sqli", "CWE-89", "false_positive", "D.java", matched="z"),
    ]
    grouped = group_findings_by_issue(findings)
    (row,) = group_by_cwe(grouped)
    cells = row["by_bucket"]
    assert {b: (c["groups"], c["occurrences"]) for b, c in cells.items()} == {
        "true_positive": (1, 1), "unverified": (1, 2), "false_positive": (1, 1),
    }
    by_id = {g.id: g for g in grouped}
    for bucket, cell in cells.items():
        assert {by_id[i].verdict for i in cell["finding_ids"]} == {bucket}
    # The members kept their own verdicts: nothing was copied across.
    assert [f.verdict for f in findings] == ["true_positive", "unverified", "unverified", "false_positive"]


def test_findings_without_a_cwe_stay_apart_per_rule():
    rows = _rows([
        _finding("rule-a", None, "true_positive", "A.java"),
        _finding("rule-b", "", "true_positive", "B.java"),
        _finding("llm-extra", None, "true_positive", "C.java", kind="llm", title="Open redirect in login"),
    ])
    assert set(rows) == {"rule:rule-a", "rule:rule-b", "topic:open redirect in login"}
    assert all(row["cwe"] is None for row in rows.values())


def test_rows_are_ordered_by_severity_then_confirmed_occurrences():
    rows = group_by_cwe(group_findings_by_issue([
        _finding("xss", "CWE-79", "unverified", "A.ts", severity="HIGH"),
        _finding("xss", "CWE-79", "unverified", "B.ts", severity="HIGH"),
        _finding("log", "CWE-117", "true_positive", "C.java", severity="HIGH"),
        _finding("cmd", "CWE-78", "true_positive", "D.java", severity="CRITICAL"),
    ]))
    assert [r["key"] for r in rows] == ["CWE-78", "CWE-117", "CWE-79"]


def test_unresolved_reasons_are_counted_per_occurrence():
    counts = count_findings(group_findings_by_issue([
        _finding("secret", "CWE-798", "unverified", "A.java", unresolved_reason="llm_budget_exhausted"),
        _finding("secret", "CWE-798", "unverified", "B.java", unresolved_reason="llm_budget_exhausted"),
        _finding("secret", "CWE-798", "unverified", "C.java", unresolved_reason="not_in_sample"),
        _finding("svuid", "CWE-502", "unverified", "D.java", unresolved_reason=None),
        _finding("sqli", "CWE-89", "true_positive", "E.java"),
    ]))
    assert counts["by_unresolved_reason"] == {"llm_budget_exhausted": 2, "not_in_sample": 1, "unknown": 1}


def test_the_excel_export_has_a_by_cwe_sheet_matching_the_counts():
    grouped = group_findings_by_issue([
        _finding("sqli", "CWE-89", "true_positive", "A.java"),
        _finding("sqli2", "CWE-89", "unverified", "B.java"),
        _finding("xss", "CWE-79", "false_positive", "C.ts"),
    ])
    project = Project(id="p1", name="demo")
    scan = Scan(id="s1", project_id="p1", version=1, scan_type="sast", status="success")
    workbook = load_workbook(BytesIO(build_excel_report(project, scan, grouped)))
    assert workbook.sheetnames[:3] == ["Summary", "By CWE", "Findings"]
    sheet = workbook["By CWE"]
    headers = [c.value for c in sheet[1]]
    rows = [dict(zip(headers, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
    assert [r["CWE"] for r in rows] == ["CWE-89", "CWE-79"]
    assert rows[0]["Confirmed"] == "1 (1)" and rows[0]["Unresolved"] == "1 (1)"
    assert sum(r["Occurrences"] for r in rows) == sum(
        cell["occurrences"] for cell in count_findings(grouped)["totals"].values()
    )
