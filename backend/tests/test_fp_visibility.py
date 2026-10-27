"""False positives the scanner is >= 90 % sure of are left out of every reading
of a scan (`app/services/fp_visibility.py`), and only those."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db as db_module
from app.config import settings
from app.db import init_db
from app.main import app
from app.models import Finding, FindingDecision, Scan
from app.services import scan_runner
from app.services.fp_visibility import fp_certainty
from app.services.scan_export import check_report_consistency

from tests.auth_helpers import ADMIN_HEADERS, make_user

pytestmark = pytest.mark.integration

HDR = ADMIN_HEADERS
V5 = {"result_cache_version": "5"}
V4 = {"result_cache_version": "4"}


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("USE_MOCK_ENGINE", "true")
    settings.use_mock_engine = True
    settings.scan_logs_dir = str(tmp_path / "scan_logs")
    init_db(f"sqlite:///{(tmp_path / 'fp_visibility.db').as_posix()}")
    return TestClient(app)


# name -> (verdict, confidence, deciding stage or None, evidence version, reason_tag)
CASES = {
    "deep_sure": ("false_positive", 0.05, "deep_scan", V4, None),
    "deep_unsure": ("false_positive", 0.3, "deep_scan", V4, None),
    "group_sure": ("false_positive", 0.95, "group_confirm", V5, None),
    "group_unsure": ("false_positive", 0.6, "group_confirm", V5, None),
    "group_old_scale": ("false_positive", 0.95, "group_confirm", V4, None),
    "proof": ("false_positive", None, "validator", V5, "match_is_declaration"),
    "no_history": ("false_positive", 0.01, None, None, None),
    "true_positive": ("true_positive", 0.99, "group_confirm", V5, None),
    "unresolved": ("unverified", None, "group_confirm", V5, None),
}
HIDDEN = {"deep_sure", "group_sure", "proof"}


@pytest.fixture()
def seeded(client, monkeypatch):
    monkeypatch.setattr(scan_runner, "execute_scan", lambda scan_id: None)
    project_id = client.post(
        "/projects", headers=HDR,
        json={"name": "fp-visibility", "default_git_url": "https://github.com/org/repo.git"},
    ).json()["id"]
    scan_id = client.post(
        "/scans", headers=HDR, json={"project_id": project_id, "commit_sha": "cc"}
    ).json()["scan_id"]
    ids = {}
    with db_module.SessionLocal() as db:
        db.get(Scan, scan_id).status = "success"
        for i, (name, (verdict, confidence, stage, version, tag)) in enumerate(CASES.items()):
            # Distinct CWE and rule per case: every case is its own group.
            finding = Finding(
                scan_id=scan_id, kind="llm", severity="HIGH", title=name, rule_id=name,
                cwe=f"CWE-{100 + i}", file_path=f"src/{name}.java", line=3, verdict=verdict,
                confidence=confidence, reason_tag=tag, issue_class="security",
                unresolved_reason="llm_missing_verdict" if verdict == "unverified" else None,
            )
            db.add(finding)
            db.flush()
            ids[name] = finding.id
            if stage:
                db.add(FindingDecision(
                    finding_id=finding.id, scan_id=scan_id, stage=stage, verdict=verdict,
                    confidence=confidence, evidence_version=version,
                ))
        db.commit()
    return scan_id, ids


def _listed(client, scan_id):
    body = client.get(f"/scans/{scan_id}/findings", headers=HDR).json()
    return body, {f["title"] for f in body["findings"]}


def test_only_confident_false_positives_are_left_out_of_the_findings_list(client, seeded):
    scan_id, _ids = seeded
    body, titles = _listed(client, scan_id)
    assert titles == set(CASES) - HIDDEN
    assert body["hidden_fp_count"] == len(HIDDEN)
    assert body["fp_count"] == 4  # deep_unsure, group_unsure, group_old_scale, no_history
    assert sum(row["occurrences"] for row in body["cwe_groups"]) == len(CASES) - len(HIDDEN)


def test_exports_and_review_queue_leave_them_out_too(client, seeded):
    scan_id, _ids = seeded
    xlsx = client.get(f"/scans/{scan_id}/export.xlsx", headers=HDR)
    assert xlsx.status_code == 200
    sheet = load_workbook(BytesIO(xlsx.content))["Findings"]
    issue = [c.value for c in sheet[1]].index("Issue")
    exported = {row[issue] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported == set(CASES) - HIDDEN

    queue = client.get("/feedback", headers=HDR, params={"scan_id": scan_id, "limit": 500})
    assert queue.status_code == 200, queue.text
    assert {r["title"] for r in queue.json()["items"]} == set(CASES) - HIDDEN


def test_a_reviewed_false_positive_is_shown_whatever_the_certainty(client, seeded):
    scan_id, ids = seeded
    headers = make_user("alice@test.local", role="admin")[1]
    r = client.post(f"/scans/{scan_id}/findings/{ids['group_sure']}/review", headers=headers,
                    json={"decision": "false_positive", "comment": "Checked: the value is a constant."})
    assert r.status_code == 200, r.text
    body, titles = _listed(client, scan_id)
    assert "group_sure" in titles
    assert body["hidden_fp_count"] == len(HIDDEN) - 1


def test_reports_stay_consistent_with_rows_left_out(client, seeded):
    scan_id, _ids = seeded
    with db_module.SessionLocal() as db:
        assert check_report_consistency(db, db.get(Scan, scan_id)) == []


@pytest.mark.parametrize("stage, version, confidence, expected", [
    ("deep_scan", V4, 0.1, pytest.approx(0.9)),
    ("flow_verify", None, 0.2, pytest.approx(0.8)),
    ("triage_llm", V4, 0.0, 1.0),
    ("group_confirm", V5, 0.9, 0.9),
    ("joern_confirm", {"result_cache_version": "12"}, 0.7, 0.7),
    ("reconcile", V5, 1.4, 1.0),            # clamped
    ("group_confirm", V4, 0.99, None),      # the old prompt never said which scale
    ("group_confirm", {"result_cache_version": "mock"}, 0.99, None),
    ("dedupe", V5, 0.99, None),             # not a deciding stage
])
def test_fp_certainty_reads_the_scale_of_the_stage_that_decided(stage, version, confidence, expected):
    finding = Finding(verdict="false_positive", reason_tag=None)
    decision = FindingDecision(stage=stage, verdict="false_positive", confidence=confidence,
                               evidence_version=version)
    assert fp_certainty(finding, decision) == expected


def test_fp_certainty_needs_a_dismissal():
    decision = FindingDecision(stage="group_confirm", verdict="false_positive", confidence=0.99,
                               evidence_version=V5)
    assert fp_certainty(Finding(verdict="true_positive"), decision) is None
    assert fp_certainty(Finding(verdict="unverified"), decision) is None
    # The finding is a dismissal now, but the deciding stage confirmed it: no certainty.
    confirmed = FindingDecision(stage="group_confirm", verdict="true_positive", confidence=0.99,
                                evidence_version=V5)
    assert fp_certainty(Finding(verdict="false_positive"), confirmed) is None
    assert fp_certainty(Finding(verdict="false_positive", reason_tag="stub_no_sink"), None) == 1.0
