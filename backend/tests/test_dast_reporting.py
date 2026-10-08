"""Reporting (DAST-11-010 .. -040): the report data set, the scan-to-scan diff and first/last seen,
and the PDF and XLSX exports with their redaction, separation and formula guards."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app import db as app_db
from app.config import settings
from app.dast.domain.diff import DiffItem, DiffStatus, compute_diff, item_key
from app.main import app
from app.models import DastEvidence, Finding, Project, Scan
from app.services.dast_report import coverage_narrative

pytestmark = pytest.mark.integration

HDR = {"Authorization": f"Bearer {settings.sast_api_token}"}


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("USE_MOCK_ENGINE", "true")
    monkeypatch.setattr(settings, "use_mock_engine", True)
    monkeypatch.setattr(settings, "scan_logs_dir", str(tmp_path / "scan_logs"))
    app_db.init_db(f"sqlite:///{(tmp_path / 'r.db').as_posix()}")
    return TestClient(app)


@pytest.fixture()
def project_id(client):
    return client.post("/projects", json={"name": "Report Demo", "default_git_url": ""}, headers=HDR).json()["id"]


def create_scan(client, project_id, **dast):
    body = {"project_id": project_id, "scan_type": "dast", "dast": {"target_url": "http://localhost:3000", "authorized": True, **dast}}
    response = client.post("/scans", json=body, headers=HDR)
    assert response.status_code == 202, response.text
    return response.json()["scan_id"]


def seed(scan_id, *, rule_id="dast-sqli", severity="HIGH", verdict="true_positive", target="GET /api/x", param="q", title=None, test_excerpt="HTTP 500", baseline="HTTP 200", signals=None, roles=None, source="", reason="because", remediation_ref="rem-sqli"):
    with app_db.SessionLocal() as db:
        row = Finding(
            scan_id=scan_id, severity=severity, rule_id=rule_id, title=title or f"{rule_id} on {target}", file_path=target, kind="dast", verdict=verdict,
            target=target, param=param, location="query", confidence_label="HIGH", standards=["OWASP A03:2021", "CWE-89"], remediation_ref=remediation_ref,
            roles=roles, source=source, verdict_reason=reason, first_seen_scan_id=scan_id, last_seen_scan_id=scan_id,
        )
        db.add(row)
        db.flush()
        db.add(DastEvidence(finding_id=row.id, baseline_excerpt=baseline, test_excerpt=test_excerpt, signals=signals if signals is not None else ([{"name": "db_error_surfaced", "detail": "SQL syntax"}] if verdict == "true_positive" else []), audit_ref="3"))
        db.commit()
        return row.id


def wipe_findings(scan_id):
    with app_db.SessionLocal() as db:
        ids = [f.id for f in db.query(Finding).filter(Finding.scan_id == scan_id)]
        db.query(DastEvidence).filter(DastEvidence.finding_id.in_(ids)).delete(synchronize_session=False)
        db.query(Finding).filter(Finding.scan_id == scan_id).delete(synchronize_session=False)
        db.commit()


def blank_scan(client, project_id):
    scan_id = create_scan(client, project_id)
    wipe_findings(scan_id)
    return scan_id


def report(client, scan_id, **params):
    r = client.get(f"/scans/{scan_id}/report", params=params, headers=HDR)
    assert r.status_code == 200, r.text
    return r.json()


# ---- the report data set (DAST-11-010) ------------------------------------------------------------------


def test_coverage_narrative_is_complete_only_when_nothing_cut_the_run_short():
    assert coverage_narrative({}) == {"complete": True, "reasons": []}
    cancelled = coverage_narrative({"stop_reason": "cancelled"})
    assert not cancelled["complete"] and "cancelled" in cancelled["reasons"][0]
    lost_creds = coverage_narrative({"credentials_lost": True})
    assert not lost_creds["complete"] and "credentials" in lost_creds["reasons"][0]
    capped = coverage_narrative({"active_testing": {"targets_capped": 3}})
    assert not capped["complete"] and "3 endpoint(s)" in capped["reasons"][0]
    session_lost = coverage_narrative({"active_testing": {"session_lost_reason": "401/403"}})
    assert not session_lost["complete"] and "session was lost" in session_lost["reasons"][0]
    degraded = coverage_narrative({"authorization": {"degraded": True, "degraded_reason": "no second identity"}})
    assert not degraded["complete"] and degraded["reasons"] == ["no second identity"]
    cap_reached = coverage_narrative({"authorization": {"cap_reached": True}})
    assert not cap_reached["complete"] and "replay cap" in cap_reached["reasons"][0]
    # several reasons at once are all kept, not just the first
    multi = coverage_narrative({"credentials_lost": True, "active_testing": {"targets_capped": 1}})
    assert len(multi["reasons"]) == 2


def test_the_report_counts_equal_the_stored_findings_and_unverified_is_separate(client, project_id):
    scan_id = create_scan(client, project_id)
    body = report(client, scan_id)
    with app_db.SessionLocal() as db:
        rows = db.query(Finding).filter(Finding.scan_id == scan_id).all()
    tp = [r for r in rows if r.verdict == "true_positive"]
    unverified = [r for r in rows if r.verdict == "unverified"]
    assert body["summary"]["confirmed"] == len(tp) and body["summary"]["unverified"] == len(unverified) and body["summary"]["total"] == len(rows)
    for sev in ("critical", "high", "medium", "low", "info"):
        assert body["summary"]["severity_counts"][sev] == len([r for r in tp if r.severity.lower() == sev])
    assert sum(body["summary"]["severity_counts"].values()) == len(tp)  # unverified never enters the severity totals
    assert body["unverified"] == []
    assert all(f["verdict"] == "true_positive" for f in body["confirmed"])


def test_findings_carry_evidence_remediation_and_standards(client, project_id):
    scan_id = blank_scan(client, project_id)
    seed(scan_id)
    finding = report(client, scan_id)["confirmed"][0]
    assert finding["evidence"]["test"] == "HTTP 500" and finding["evidence"]["signals"][0]["name"] == "db_error_surfaced" and finding["evidence"]["audit_ref"] == "3"
    assert "GET /api/x" in finding["evidence"]["request_text"] and finding["evidence"]["request"]["parameter"] == "q"
    assert "parameterized" in finding["remediation"].lower() or finding["remediation"]
    assert finding["standards"] == ["OWASP A03:2021", "CWE-89"] and finding["target"] == "GET /api/x" and finding["param"] == "q"


def test_skipped_checks_packs_and_settings_are_in_the_report(client, project_id):
    scan_id = create_scan(client, project_id, intensity="light")
    with app_db.SessionLocal() as db:
        scan = db.get(Scan, scan_id)
        meta = dict(scan.metadata_json)
        meta["dast_skipped_degraded"] = [{"check_id": "dast-ssrf", "title": "SSRF", "status": "degraded", "dependency": "oast", "effective_mode": "assist-only", "reason": "SSRF is assist-only: no OAST callback address is configured"}]
        scan.metadata_json = meta
        db.commit()
    body = report(client, scan_id)
    assert body["skipped_degraded"][0]["check_id"] == "dast-ssrf" and "OAST" in body["skipped_degraded"][0]["reason"]
    assert {p["id"] for p in body["rule_packs"]} >= {"owasp-api-top10-2023"} and body["catalog_version"]
    assert body["settings"]["intensity"] == "light" and body["settings"]["safe_mode"] is True and body["settings"]["rule_checks_total"] > 10


def test_the_report_coverage_block_carries_the_narrative(client, project_id):
    scan_id = create_scan(client, project_id)
    body = report(client, scan_id)
    assert body["coverage"]["complete"] is True and body["coverage"]["reasons"] == []
    with app_db.SessionLocal() as db:
        scan = db.get(Scan, scan_id)
        meta = dict(scan.metadata_json)
        meta["discovery"] = {**meta.get("discovery", {}), "credentials_lost": True}
        scan.metadata_json = meta
        db.commit()
    degraded = report(client, scan_id)
    assert degraded["coverage"]["complete"] is False and degraded["coverage"]["reasons"]


def test_the_report_is_dast_only_and_sast_scans_are_unchanged(client):
    sast_project = client.post("/projects", json={"name": "S", "default_git_url": "https://github.com/org/repo.git"}, headers=HDR).json()["id"]
    sast = client.post("/scans", json={"project_id": sast_project}, headers=HDR).json()["scan_id"]
    assert client.get(f"/scans/{sast}/report", headers=HDR).status_code == 400
    assert client.get(f"/scans/{sast}/export.pdf", headers=HDR).headers["content-type"] == "application/pdf"
    assert client.get("/scans/missing/report", headers=HDR).status_code == 404


# ---- the diff (DAST-11-020) -----------------------------------------------------------------------------


def di(check, target, param=None, severity="HIGH", verdict="true_positive", roles=()):
    return DiffItem(check, target, param, severity, verdict, tuple(roles))


def statuses(result):
    return {e.key: e.status for e in result.entries}


def test_new_fixed_and_still_open_are_classified_on_check_template_and_parameter():
    previous = [di("dast-sqli", "GET /api/users/1", "q"), di("dast-xss", "GET /a", "n"), di("dast-cmdi", "GET /c", "x")]
    current = [di("dast-sqli", "GET /api/users/2", "q"), di("dast-xss", "GET /b", "n")]  # /users/1 and /users/2 are one template
    result = compute_diff(current, previous)
    assert result.baseline
    by = statuses(result)
    assert by[item_key(di("dast-sqli", "GET /api/users/9", "q"))] is DiffStatus.STILL_OPEN
    assert by[item_key(di("dast-xss", "GET /b", "n"))] is DiffStatus.NEW and by[item_key(di("dast-xss", "GET /a", "n"))] is DiffStatus.FIXED
    assert by[item_key(di("dast-cmdi", "GET /c", "x"))] is DiffStatus.FIXED
    assert result.counts() == {"new": 1, "fixed": 2, "still_open": 1}


def test_no_baseline_treats_everything_as_new():
    result = compute_diff([di("dast-sqli", "GET /a", "q")], None)
    assert not result.baseline and [e.status for e in result.entries] == [DiffStatus.NEW] and "no earlier scan" in result.entries[0].note


def test_a_check_that_did_not_run_is_never_reported_as_fixed():
    previous = [di("dast-sqli", "GET /a", "q"), di("dast-ssrf", "GET /f", "url")]
    result = compute_diff([], previous, current_checks_run={"dast-sqli"}, previous_checks_run={"dast-sqli", "dast-ssrf"})
    assert [e.item.check_id for e in result.entries] == ["dast-sqli"] and result.entries[0].status is DiffStatus.FIXED
    assert [n.item.check_id for n in result.not_comparable] == ["dast-ssrf"] and "cannot be called fixed" in result.not_comparable[0].reason
    assert result.checks_not_run_current == ["dast-ssrf"]
    later = compute_diff([di("dast-xss", "GET /x", "n")], [], current_checks_run={"dast-xss"}, previous_checks_run=set())
    assert later.entries[0].status is DiffStatus.NEW and "did not run this check" in later.entries[0].note


def test_unverified_findings_are_diffed_with_their_verdict_and_role_pairs_stay_separate():
    previous = [di("dast-bfla", "GET /admin", None, verdict="unverified", roles=("Root", "Alice"))]
    current = [di("dast-bfla", "GET /admin", None, verdict="true_positive", roles=("Root", "Alice")), di("dast-bfla", "GET /admin", None, roles=("Root", "Bob"))]
    result = compute_diff(current, previous)
    assert result.counts() == {"new": 1, "fixed": 0, "still_open": 1}
    still = next(e for e in result.entries if e.status is DiffStatus.STILL_OPEN)
    assert still.item.verdict == "true_positive" and still.item.roles == ("Root", "Alice")


def test_the_diff_endpoint_compares_with_the_previous_scan_and_an_override(client, project_id):
    first = create_scan(client, project_id)
    second = create_scan(client, project_id)
    third = create_scan(client, project_id)
    diff = client.get(f"/scans/{third}/diff", headers=HDR).json()
    assert diff["baseline"] is True and diff["previous_scan_id"] == second and diff["counts"]["new"] == 0 and diff["counts"]["fixed"] == 0 and diff["counts"]["still_open"] > 0
    seed(third, rule_id="dast-cmdi", target="GET /api/new", param="host", severity="CRITICAL")
    wipe_one = next(f for f in client.get(f"/scans/{third}/findings", headers=HDR).json()["findings"] if f["rule_id"] == "dast-sqli")
    with app_db.SessionLocal() as db:
        victim = db.get(Finding, wipe_one["id"])
        key = (victim.rule_id, victim.target)
        db.query(DastEvidence).filter(DastEvidence.finding_id == victim.id).delete()
        db.delete(victim)
        db.commit()
    against_first = client.get(f"/scans/{third}/diff", params={"against": first}, headers=HDR).json()
    assert against_first["previous_scan_id"] == first
    fresh = [e for e in against_first["entries"] if e["status"] == "new"]
    assert [(e["check_id"], e["target"]) for e in fresh] == [("dast-cmdi", "GET /api/new")]
    assert (key[0], key[1]) in {(e["check_id"], e["target"]) for e in against_first["entries"] if e["status"] == "fixed"}
    assert client.get(f"/scans/{third}/diff", params={"against": third}, headers=HDR).status_code == 400
    assert client.get(f"/scans/{first}/diff", headers=HDR).json()["baseline"] is False  # nothing earlier
    assert client.get("/scans/zzz/diff", headers=HDR).status_code == 404


def test_the_first_scan_has_no_baseline_and_a_different_target_is_never_a_baseline(client, project_id):
    one = create_scan(client, project_id)
    assert client.get(f"/scans/{one}/diff", headers=HDR).json()["baseline"] is False
    other = create_scan(client, project_id, target_url="http://localhost:4000")
    assert client.get(f"/scans/{other}/diff", headers=HDR).json()["baseline"] is False
    again = create_scan(client, project_id)
    assert client.get(f"/scans/{again}/diff", headers=HDR).json()["previous_scan_id"] == one


def test_first_and_last_seen_follow_a_finding_across_scans(client, project_id):
    first = create_scan(client, project_id)
    second = create_scan(client, project_id)
    third = create_scan(client, project_id)
    with app_db.SessionLocal() as db:
        for scan_id in (first, second, third):
            rows = db.query(Finding).filter(Finding.scan_id == scan_id, Finding.kind == "dast").all()
            assert rows and all(r.last_seen_scan_id == scan_id for r in rows)
            assert all(r.first_seen_scan_id == first for r in rows), scan_id  # open since the very first scan


def test_the_diff_respects_the_rule_set_each_scan_ran_with(client, project_id):
    first = create_scan(client, project_id, rule_overrides={"dast-xss": {"enabled": True}})
    second = create_scan(client, project_id, rule_overrides={"dast-xss": {"enabled": False}})
    seed(first, rule_id="dast-xss", target="GET /only-before", param="q")
    diff = client.get(f"/scans/{second}/diff", headers=HDR).json()
    assert not [e for e in diff["entries"] if e["check_id"] == "dast-xss" and e["target"] == "GET /only-before" and e["status"] == "fixed"]
    assert any(n["check_id"] == "dast-xss" and "cannot be called fixed" in n["reason"] for n in diff["not_comparable"])
    assert any("dast-xss" in note for note in diff["notes"])


# ---- PDF (DAST-11-030) ----------------------------------------------------------------------------------


def pdf_text(client, scan_id):
    response = client.get(f"/scans/{scan_id}/export.pdf", headers=HDR)
    assert response.status_code == 200 and response.headers["content-type"] == "application/pdf"
    reader = PdfReader(io.BytesIO(response.content))
    return "\n".join(page.extract_text() for page in reader.pages), len(reader.pages)


def test_the_pdf_has_every_required_section_and_labels_unverified_findings(client, project_id):
    create_scan(client, project_id)
    scan_id = create_scan(client, project_id)
    text, pages = pdf_text(client, scan_id)
    for section in ("DAST Security Report", "SEVERITY SUMMARY", "CONFIRMED FINDINGS", "SKIPPED AND DEGRADED CHECKS", "CHANGES SINCE THE PREVIOUS SCAN", "STILL OPEN", "REMEDIATION", "PARAMETER", "REQUEST SENT", "BASELINE RESPONSE", "RESPONSE RECEIVED"):
        assert section in text, section
    assert "Report Demo" in text and "localhost:3000" in text
    assert pages >= 3


def test_no_credential_shape_or_secret_reaches_the_pdf_or_the_workbook(client, project_id):
    scan_id = blank_scan(client, project_id)
    secrets = ["sk-ant-abcdefghijklmnop1234", "AKIAIOSFODNN7EXAMPLE", "Bearer eyJhbGciOiJIUzI1NiJ9.payload.signature123456"]
    seed(scan_id, test_excerpt=" ".join(secrets), baseline="key " + secrets[1], reason="leaked " + secrets[0])
    text, _ = pdf_text(client, scan_id)
    xlsx = client.get(f"/scans/{scan_id}/export.xlsx", headers=HDR).content
    sheet_text = " ".join(str(c.value) for ws in load_workbook(io.BytesIO(xlsx)).worksheets for row in ws.iter_rows() for c in row if c.value)
    for secret in ("sk-ant-abcdefghijklmnop1234", "AKIAIOSFODNN7EXAMPLE", "eyJhbGciOiJIUzI1NiJ9.payload"):
        assert secret not in text and secret not in sheet_text
    assert "[REDACTED]" in text


def test_a_large_scan_paginates_without_failing_or_truncating_rows(client, project_id):
    scan_id = blank_scan(client, project_id)
    for i in range(160):
        seed(scan_id, rule_id="dast-sqli" if i % 2 else "dast-xss", target=f"GET /api/item{i}", param=f"p{i}", verdict="true_positive" if i % 4 else "unverified", test_excerpt="line\n" * 40)
    text, pages = pdf_text(client, scan_id)
    assert pages > 20 and "/api/item159" in text and "/api/item0" in text
    sheet = load_workbook(io.BytesIO(client.get(f"/scans/{scan_id}/export.xlsx", headers=HDR).content))
    assert sheet["Findings"].max_row - 1 + sheet["Unverified"].max_row - 1 == 160


# ---- XLSX (DAST-11-040) ---------------------------------------------------------------------------------


def test_the_workbook_has_the_required_sheets_headers_and_matching_counts(client, project_id):
    create_scan(client, project_id)
    scan_id = create_scan(client, project_id)
    wb = load_workbook(io.BytesIO(client.get(f"/scans/{scan_id}/export.xlsx", headers=HDR).content))
    assert wb.sheetnames == ["Summary", "Findings", "Unverified", "Diff", "Skipped Checks"]
    headers = [c.value for c in wb["Findings"][1]]
    for column in ("Check id", "Severity", "Confidence", "Verdict", "Target", "Parameter", "Location", "Standards", "Remediation", "First seen", "Last seen"):
        assert column in headers
    body = report(client, scan_id)
    assert wb["Findings"].max_row - 1 == body["summary"]["confirmed"] and wb["Unverified"].max_row - 1 == body["summary"]["unverified"]
    verdicts = {r[headers.index("Verdict")].value for r in wb["Findings"].iter_rows(min_row=2)}
    assert verdicts == {"Confirmed"}
    assert wb["Unverified"].max_row - 1 == body["summary"]["unverified"]
    first_seen = {r[headers.index("First seen")].value for r in wb["Findings"].iter_rows(min_row=2)}
    assert first_seen == {"v1"}  # carried from the first scan
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows()}
    assert summary["Confirmed findings"] == body["summary"]["confirmed"] and summary["Still open"] == body["diff"]["counts"]["still_open"]
    assert len(list(wb["Diff"].iter_rows(min_row=2))) == len(body["diff"]["entries"])


@pytest.mark.parametrize("payload", ["=HYPERLINK(\"http://evil.test\",\"x\")", "+1+1", "-2+3", "@SUM(A1:A2)", "\t=1+1"])
def test_cells_that_would_start_a_formula_are_stored_as_text(client, project_id, payload):
    scan_id = blank_scan(client, project_id)
    seed(scan_id, test_excerpt=payload, title=payload, reason=payload)
    wb = load_workbook(io.BytesIO(client.get(f"/scans/{scan_id}/export.xlsx", headers=HDR).content))
    sheet = wb["Findings"]
    headers = [c.value for c in sheet[1]]
    for column in ("Title", "Evidence (test response)", "Why"):
        cell = sheet.cell(2, headers.index(column) + 1)
        assert cell.data_type == "s" and str(cell.value).startswith(payload[0]), (column, cell.data_type)
    raw = io.BytesIO(client.get(f"/scans/{scan_id}/export.xlsx", headers=HDR).content)
    import zipfile

    assert b"<f>" not in zipfile.ZipFile(raw).read("xl/worksheets/sheet2.xml")  # no formula element was written


def test_sast_exports_are_unchanged_by_the_dast_branch():
    from app.services.scan_export import export_filenames

    project = Project(name="P")
    assert export_filenames(project, Scan(version=2, scan_type="sast"))["pdf"] == "P-v2.pdf"
    assert export_filenames(project, Scan(version=2, scan_type="dast"))["xlsx"] == "P-dast-v2.xlsx"


# ---- the scan list carries finding counts (DAST-13-010) ---------------------------------------------------------


def test_the_scan_list_carries_confirmed_counts_by_severity_and_the_unverified_count(client, project_id):
    scan_id = blank_scan(client, project_id)
    seed(scan_id, rule_id="dast-sqli", severity="HIGH", target="GET /a", param="x")
    seed(scan_id, rule_id="dast-cmdi", severity="CRITICAL", target="GET /b", param="y")
    seed(scan_id, rule_id="dast-xss", severity="MEDIUM", verdict="unverified", target="GET /c", param="z")
    other = create_scan(client, project_id)
    listed = {s["id"]: s for s in client.get("/scans", params={"scan_type": "dast"}, headers=HDR).json()}
    assert listed[scan_id]["dast_counts"] == {"critical": 1, "high": 1, "medium": 0, "low": 0, "info": 0, "unverified": 1}  # unverified is never a severity
    assert listed[other]["dast_counts"]["unverified"] == 0
    sast_project = client.post("/projects", json={"name": "s2", "default_git_url": "https://github.com/org/repo.git"}, headers=HDR).json()["id"]
    sast = client.post("/scans", json={"project_id": sast_project}, headers=HDR).json()["scan_id"]
    assert next(s for s in client.get("/scans", headers=HDR).json() if s["id"] == sast)["dast_counts"] is None
    assert client.get(f"/scans/{scan_id}/status", headers=HDR).status_code == 200
