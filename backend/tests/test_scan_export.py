"""PDF / Excel export smoke tests."""

import pytest
from io import BytesIO

from openpyxl import load_workbook

from app.models import Finding, Project, Scan
from app.services.scan_export import (
    build_excel_report,
    build_pdf_report,
    count_findings,
    export_filenames,
    group_findings_by_issue,
)


def _sample(project_name: str = "Demo App"):
    project = Project(
        id="p1",
        name=project_name,
        default_git_url="https://example.com/repo.git",
        default_branch="main",
    )
    scan = Scan(
        id="s1",
        project_id="p1",
        git_url=project.default_git_url,
        branch="main",
        commit_sha="abcdef1234567890",
        status="success",
        version=3,
        llm_provider="bedrock",
        llm_model="haiku",
    )
    findings = [
        Finding(
            id="f1",
            scan_id="s1",
            severity="high",
            title="SQL injection via concatenated user input",
            cwe="CWE-89",
            owasp_category="A03:2021-Injection",
            file_path="app/db.py",
            line=42,
            snippet='query = "SELECT * FROM users WHERE id=" + user_id',
            verdict="TRUE_POSITIVE",
            confidence=0.91,
            risk="Attackers can read or modify arbitrary database rows.",
            ai_fix="Use parameterized queries.\n\nUse bind parameters instead of string concat.",
            example_fix=(
                'cursor.execute(\n'
                '    "SELECT * FROM users WHERE id = ?",\n'
                "    (user_id,),\n"
                ")"
            ),
            verdict_reason="User input concatenated into SQL.",
            source='request.getParameter("id")',
            sink="stmt.executeQuery(sql)",
            data_flow='id from request.getParameter("id") is concatenated into sql, then passed to executeQuery',
        ),
        Finding(
            id="f2",
            scan_id="s1",
            severity="low",
            title="Hardcoded token in test",
            file_path="tests/test_auth.py",
            line=10,
            snippet='TOKEN = "demo"',
            verdict="FALSE_POSITIVE",
            confidence=0.8,
            verdict_reason="Test fixture only.",
        ),
    ]
    return project, scan, findings


def test_group_findings_by_issue_merges_same_rule_id_across_files():
    """Two findings with same rule_id and different file paths should merge."""
    f1 = Finding(
        id="f1", scan_id="s1", severity="high", title="SQL injection",
        cwe="CWE-89", rule_id="sql-concat", file_path="app/db.py", line=10,
        verdict="true_positive",
    )
    f2 = Finding(
        id="f2", scan_id="s1", severity="high", title="SQL injection",
        cwe="CWE-89", rule_id="sql-concat", file_path="app/query.py", line=20,
        verdict="true_positive",
    )
    grouped = group_findings_by_issue([f1, f2])
    assert len(grouped) == 1
    assert grouped[0].affected_files == ["app/db.py", "app/query.py"]
    assert grouped[0].file_path in ["app/db.py", "app/query.py"]


def test_group_findings_by_issue_merges_llm_extra_with_same_cwe_title():
    """Two 'llm-extra' findings with same CWE and similar title should merge."""
    f1 = Finding(
        id="f1", scan_id="s1", severity="medium", title="XSS vulnerability in DOM",
        cwe="CWE-79", rule_id="llm-extra", file_path="frontend/app.tsx", line=50,
        verdict="true_positive",
    )
    f2 = Finding(
        id="f2", scan_id="s1", severity="medium", title="XSS vulnerability in template",
        cwe="CWE-79", rule_id="llm-extra", file_path="frontend/form.tsx", line=30,
        verdict="true_positive",
    )
    grouped = group_findings_by_issue([f1, f2])
    assert len(grouped) == 1
    assert sorted(grouped[0].affected_files or []) == ["frontend/app.tsx", "frontend/form.tsx"]


def test_group_findings_by_issue_keeps_different_cwe_separate():
    """Two 'llm-extra' findings with different CWEs should stay separate."""
    f1 = Finding(
        id="f1", scan_id="s1", severity="high", title="SQL issue",
        cwe="CWE-89", rule_id="llm-extra", file_path="app/db.py", line=10,
        verdict="true_positive",
    )
    f2 = Finding(
        id="f2", scan_id="s1", severity="high", title="XSS issue",
        cwe="CWE-79", rule_id="llm-extra", file_path="app/view.py", line=20,
        verdict="true_positive",
    )
    grouped = group_findings_by_issue([f1, f2])
    assert len(grouped) == 2
    assert all(f.affected_files is None for f in grouped)


def test_group_findings_by_issue_keeps_different_verdicts_separate():
    """Two findings with same rule_id but different verdicts should stay separate."""
    f1 = Finding(
        id="f1", scan_id="s1", severity="high", title="SQL injection",
        cwe="CWE-89", rule_id="sql-concat", file_path="app/db.py", line=10,
        verdict="true_positive",
    )
    f2 = Finding(
        id="f2", scan_id="s1", severity="high", title="SQL injection",
        cwe="CWE-89", rule_id="sql-concat", file_path="app/query.py", line=20,
        verdict="false_positive",
    )
    grouped = group_findings_by_issue([f1, f2])
    assert len(grouped) == 2
    assert all(f.affected_files is None for f in grouped)


def test_group_findings_by_issue_merges_identical_matched_text_across_files():
    """The generated-code case: one rule tripping on byte-identical code across
    many files collapses to a single entry listing every location."""
    findings = [
        Finding(
            id=f"f{i}", scan_id="s1", severity="high", title="Hardcoded password",
            cwe="CWE-798", rule_id="secret-password-assign",
            file_path=f"castorgen/Type{i}Descriptor.java", line=6,
            matched_text='password = "hunter2secret"', verdict="unverified",
        )
        for i in range(5)
    ]
    grouped = group_findings_by_issue(findings)
    assert len(grouped) == 1
    assert len(grouped[0].affected_files or []) == 5


def test_group_findings_by_issue_keeps_different_matched_text_separate():
    """Grouping by rule alone would merge a real credential with a test
    placeholder. Keying on the matched text keeps them distinct."""
    real = Finding(
        id="f1", scan_id="s1", severity="high", title="Hardcoded password",
        cwe="CWE-798", rule_id="secret-password-assign", file_path="app/Db.java", line=10,
        matched_text='password = "Pr0dS3cret!"', verdict="unverified",
    )
    placeholder = Finding(
        id="f2", scan_id="s1", severity="high", title="Hardcoded password",
        cwe="CWE-798", rule_id="secret-password-assign", file_path="test/Fixture.java", line=4,
        matched_text='password = "changeme"', verdict="unverified",
    )
    grouped = group_findings_by_issue([real, placeholder])
    assert len(grouped) == 2


def test_group_findings_by_issue_merges_different_matched_text_when_group_by_rule():
    """A rule opted into group_by_rule (e.g. serializable-missing-svuid, whose
    matched text is the class declaration and so always differs by class
    name) collapses to one row per rule regardless of matched_text."""
    findings = [
        Finding(
            id=f"f{i}", scan_id="s1", severity="info", title="Serializable class without serialVersionUID",
            cwe="CWE-502", rule_id="serializable-missing-svuid",
            file_path=f"app/Class{i}.java", line=6,
            matched_text=f"class Class{i} implements Serializable",
            group_by_rule=True, verdict="unverified",
        )
        for i in range(3)
    ]
    grouped = group_findings_by_issue(findings)
    assert len(grouped) == 1
    assert len(grouped[0].affected_files or []) == 3


def test_group_findings_by_issue_normalizes_whitespace_in_matched_text():
    """Trivially different formatting of the same construct still groups."""
    f1 = Finding(
        id="f1", scan_id="s1", severity="high", title="Hardcoded password",
        rule_id="secret-password-assign", file_path="a.java", line=1,
        matched_text='password  =  "x"', verdict="unverified",
    )
    f2 = Finding(
        id="f2", scan_id="s1", severity="high", title="Hardcoded password",
        rule_id="secret-password-assign", file_path="b.java", line=1,
        matched_text='password = "x"', verdict="unverified",
    )
    assert len(group_findings_by_issue([f1, f2])) == 1


def test_group_findings_by_issue_llm_findings_unaffected_by_matched_text():
    """LLM findings carry no matched_text, so adding it to the key must not
    change how they group — this is the regression guard for the existing
    rule-only behaviour."""
    f1 = Finding(
        id="f1", scan_id="s1", severity="high", title="SQL injection",
        cwe="CWE-89", rule_id="sql-concat", file_path="app/db.py", line=10,
        verdict="true_positive",
    )
    f2 = Finding(
        id="f2", scan_id="s1", severity="high", title="SQL injection",
        cwe="CWE-89", rule_id="sql-concat", file_path="app/query.py", line=20,
        verdict="true_positive",
    )
    assert len(group_findings_by_issue([f1, f2])) == 1


def test_display_verdict_labels_unverified_distinctly():
    """An unresolved finding must not read as confirmed or as dismissed, and a
    dismissed one is a False Positive — never "Needs Review"."""
    from app.services.scan_export import _display_verdict

    assert _display_verdict("unverified", "not_reviewed_regex") == "Unresolved (regex, not reviewed)"
    assert _display_verdict("unverified") == "Unresolved"
    assert _display_verdict("true_positive") == "True Positive"
    assert _display_verdict("false_positive") == "False Positive"


def test_group_findings_by_issue_singleton_unchanged():
    """A single finding should pass through with affected_files unset."""
    f1 = Finding(
        id="f1", scan_id="s1", severity="high", title="SQL injection",
        cwe="CWE-89", rule_id="sql-concat", file_path="app/db.py", line=10,
        verdict="true_positive",
    )
    grouped = group_findings_by_issue([f1])
    assert len(grouped) == 1
    assert grouped[0] is f1
    assert grouped[0].affected_files is None


def test_export_filenames_are_safe():
    project, scan, _ = _sample("Acme / Secure!")
    names = export_filenames(project, scan)
    assert names["pdf"].endswith(".pdf")
    assert names["xlsx"].endswith(".xlsx")
    assert " " not in names["pdf"]
    assert "/" not in names["pdf"]


def test_build_pdf_report_returns_pdf_bytes():
    project, scan, findings = _sample()
    data = build_pdf_report(project, scan, findings)
    assert data[:4] == b"%PDF"
    assert len(data) > 500


def test_pdf_findings_include_section_labels_and_summary_furniture():
    """Landscape PDF: summary furniture + uppercase section labels in detail cards."""
    from pypdf import PdfReader

    project, scan, findings = _sample()
    data = build_pdf_report(project, scan, findings)
    reader = PdfReader(BytesIO(data))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    for label in (
        "DESCRIPTION",
        "DATA FLOW",
        "AFFECTED LOCATION(S)",
        "CODE CONTEXT",
        "RISK",
        "RECOMMENDATION",
        "EXAMPLE FIX",
    ):
        assert label in text, f"missing PDF section label: {label}"
    assert "QUALITY GATE PASSED" in text
    assert "FINDINGS INDEX" in text
    assert "DETAILED FINDINGS" in text
    assert "OWASP DISTRIBUTION" in text
    assert "Page 1 of" in text
    assert "Injection" in text
    assert "A03:2021 - Injection" not in text
    assert "SQL injection via concatenated user input" in text
    assert "request.getParameter" in text
    assert "executeQuery" in text
    assert "Attackers can read" in text
    assert "cursor.execute" in text
    assert "Why true positive" not in text
    assert "How to fix" not in text


def test_build_excel_report_returns_xlsx_bytes():
    project, scan, findings = _sample()
    data = build_excel_report(project, scan, findings)
    # ZIP/OOXML signature
    assert data[:2] == b"PK"
    assert len(data) > 500


def headers_row(sheet):
    return [cell.value for cell in sheet[1]]


def test_excel_findings_use_expected_columns():
    project, scan, findings = _sample()
    data = build_excel_report(project, scan, findings)
    wb = load_workbook(BytesIO(data))
    sheet = wb["Findings"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == [
        "#",
        "Severity",
        "Scanner Verdict",
        "Review State",
        "Issue Class",
        "Unresolved Reason",
        "Verdict (raw)",
        "Counted As",
        "OWASP Category",
        "CWE",
        "CVSS",
        "Confidence",
        "Files",
        "Issue",
        "Location(s)",
        "Data Flow",
        "Risk",
        "Recommendation",
        "Example Fix",
        "Occurrence ID",
    ]
    by_header = dict(zip(headers_row(sheet), [cell.value for cell in sheet[2]]))
    assert by_header["Severity"] == "HIGH"
    assert by_header["Scanner Verdict"] == "True Positive"
    assert by_header["Review State"] == "pending"
    assert by_header["Verdict (raw)"] == "TRUE_POSITIVE"
    assert by_header["Counted As"] == "Confirmed"
    assert by_header["OWASP Category"] == "A03:2021 – Injection"
    assert by_header["CWE"] == "CWE-89"
    assert by_header["Confidence"] == "91%"
    assert by_header["Issue"] == "SQL injection via concatenated user input"
    assert "app/db.py" in (by_header["Location(s)"] or "")
    assert "SELECT * FROM users" in (by_header["Location(s)"] or "")
    assert "request.getParameter" in (by_header["Data Flow"] or "")
    assert "executeQuery" in (by_header["Data Flow"] or "")
    assert "Attackers can read" in (by_header["Risk"] or "")
    assert by_header["Recommendation"] == "Use parameterized queries."
    assert "cursor.execute" in (by_header["Example Fix"] or "")


def test_excel_summary_includes_gate_and_owasp_distribution():
    project, scan, findings = _sample()
    data = build_excel_report(project, scan, findings)
    wb = load_workbook(BytesIO(data))
    summary = wb["Summary"]
    assert summary["A2"].value == "QUALITY GATE PASSED"
    values = [summary.cell(r, 1).value for r in range(1, summary.max_row + 1)]
    assert "OWASP distribution" in values


def test_pdf_and_excel_redact_git_credentials_from_repository_field():
    from pypdf import PdfReader

    project, scan, findings = _sample()
    scan.git_url = "https://oauth2:glpat-secret-token-12345@gitlab.com/org/repo.git"
    pdf_text = "\n".join(
        (page.extract_text() or "")
        for page in PdfReader(BytesIO(build_pdf_report(project, scan, findings))).pages
    )
    assert "glpat-secret-token-12345" not in pdf_text
    assert "https://gitlab.com/org/repo.git" in pdf_text

    wb = load_workbook(BytesIO(build_excel_report(project, scan, findings)))
    summary = wb["Summary"]
    summary_values = [summary.cell(r, 2).value for r in range(1, summary.max_row + 1)]
    assert "glpat-secret-token-12345" not in str(summary_values)
    assert "https://gitlab.com/org/repo.git" in summary_values


def test_excel_strips_illegal_control_chars_from_binary_snippets():
    """Binary/non-text snippets (e.g. .class paths) must not crash openpyxl."""
    project, scan, findings = _sample()
    findings[0].file_path = "NPG/NPGEjb/com/telcordia/inpac/ejb/security/SecurityMgrEJB.class"
    findings[0].line = 1
    findings[0].snippet = (
        "1: \x00\x01\x02" + bytes(range(3, 32)).decode("latin-1") + "visible text"
    )
    data = build_excel_report(project, scan, findings)
    assert data[:2] == b"PK"
    wb = load_workbook(BytesIO(data))
    sheet = wb["Findings"]
    location_cell = sheet.cell(2, headers_row(sheet).index("Location(s)") + 1).value or ""
    assert "\x00" not in location_cell
    assert "visible text" in location_cell


def test_pdf_strips_illegal_control_chars_from_binary_snippets():
    project, scan, findings = _sample()
    findings[0].snippet = "1: \x00\x01binary\xff\xfe and text"
    data = build_pdf_report(project, scan, findings)
    assert data[:4] == b"%PDF"


def test_empty_findings_still_exports():
    project, scan, _ = _sample()
    assert build_pdf_report(project, scan, [])[:4] == b"%PDF"
    assert build_excel_report(project, scan, [])[:2] == b"PK"


def test_pdf_handles_oversized_example_fix_without_layout_error():
    """Huge example_fix must page-break / truncate instead of LayoutError."""
    project, scan, findings = _sample()
    findings[0].example_fix = "\n".join(f"line_{i} = value_{i}()" for i in range(400))
    findings[0].affected_files = [f"src/path/to/File{i}.java" for i in range(80)]
    findings[0].snippet = "\n".join(f"token_{i}" for i in range(120))
    data = build_pdf_report(project, scan, findings)
    assert data[:4] == b"%PDF"
    assert len(data) > 500


def test_excel_truncates_oversized_cells():
    from app.services.scan_export import _EXCEL_CELL_MAX

    project, scan, findings = _sample()
    findings[0].example_fix = "X" * (_EXCEL_CELL_MAX + 5000)
    data = build_excel_report(project, scan, findings)
    wb = load_workbook(BytesIO(data))
    sheet = wb["Findings"]
    cell = sheet.cell(2, headers_row(sheet).index("Example Fix") + 1).value or ""
    assert len(cell) <= _EXCEL_CELL_MAX
    assert "truncated" in cell.lower()


def test_group_findings_by_issue_keeps_distinct_flow_paths_separate():
    """kind=flow + llm-extra + same CWE/title must not merge distinct hop chains."""
    hops_a = [
        {"file": "AController.java", "method": "AController.create", "line": 10},
        {"file": "ADao.java", "method": "ADao.query", "line": 20},
    ]
    hops_b = [
        {"file": "BController.java", "method": "BController.create", "line": 10},
        {"file": "BDao.java", "method": "BDao.query", "line": 20},
    ]
    f1 = Finding(
        id="f1", scan_id="s1", severity="high", title="SQL injection in query",
        cwe="CWE-89", rule_id="llm-extra", file_path="ADao.java", line=20,
        verdict="true_positive", kind="flow",
        affected_files=["AController.java", "ADao.java"],
        flow_path=hops_a,
    )
    f2 = Finding(
        id="f2", scan_id="s1", severity="high", title="SQL injection in query",
        cwe="CWE-89", rule_id="llm-extra", file_path="BDao.java", line=20,
        verdict="true_positive", kind="flow",
        affected_files=["BController.java", "BDao.java"],
        flow_path=hops_b,
    )
    grouped = group_findings_by_issue([f1, f2])
    assert len(grouped) == 2
    by_sink = {g.file_path: g for g in grouped}
    assert by_sink["ADao.java"].affected_files == ["AController.java", "ADao.java"]
    assert by_sink["BDao.java"].affected_files == ["BController.java", "BDao.java"]
    assert by_sink["ADao.java"].flow_path == hops_a
    assert by_sink["BDao.java"].flow_path == hops_b


def _sca_sample():
    project, scan, _ = _sample()
    scan.scan_type = "sca"
    scan.llm_provider = None
    scan.llm_model = None
    scan.metadata_json = {
        "discovery": {
            "dependency_summary": {
                "total_dependencies": 3,
                "vulnerable_dependencies": 1,
                "manifests_parsed": ["requirements.txt", "package.json"],
                "ecosystems": ["pypi", "npm"],
                "severity_counts": {"CRITICAL": 1, "HIGH": 0, "MEDIUM": 0, "LOW": 0, "UNKNOWN": 0},
            }
        }
    }
    findings = [
        Finding(
            id="sca1",
            scan_id="s1",
            severity="CRITICAL",
            cvss=9.8,
            title="requests@2.25.0 — CVE-2021-1234 (CRITICAL): A test vulnerability",
            rule_id="CVE-2021-1234",
            file_path="requirements.txt",
            matched_text="requests@2.25.0",
            snippet="A test vulnerability in requests allowing request smuggling.",
            kind="sca",
            verdict="unverified",
            ai_fix="Upgrade requests to 2.26.0 or later.",
        ),
    ]
    return project, scan, findings


def test_sca_export_filenames_include_kind():
    project, scan, _ = _sca_sample()
    names = export_filenames(project, scan)
    assert names["pdf"].endswith("-sca-v3.pdf")
    assert names["xlsx"].endswith("-sca-v3.xlsx")


def test_sca_pdf_uses_dependency_layout_not_sast_furniture():
    from pypdf import PdfReader

    project, scan, findings = _sca_sample()
    data = build_pdf_report(project, scan, findings)
    assert data[:4] == b"%PDF"
    text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(data)).pages)
    assert "Dependency Vulnerability (SCA) Report" in text
    assert "OSV.dev" in text
    assert "ADVISORIES INDEX" in text
    assert "DETAILED ADVISORIES" in text
    assert "SUGGESTED UPGRADE" in text
    assert "ADVISORY SUMMARY" in text
    assert "requests@2.25.0" in text
    assert "CVE-2021-1234" in text
    assert "2.26.0" in text
    assert "requirements.txt" in text
    for sast_only in (
        "SAST Scan Report",
        "QUALITY GATE",
        "OWASP DISTRIBUTION",
        "CODE CONTEXT",
        "DATA FLOW",
        "EXAMPLE FIX",
        "Unverified (regex)",
        "Analysis engine",
    ):
        assert sast_only not in text, f"SCA PDF still contains SAST furniture: {sast_only}"


def test_sca_excel_uses_dependency_columns():
    project, scan, findings = _sca_sample()
    data = build_excel_report(project, scan, findings)
    assert data[:2] == b"PK"
    wb = load_workbook(BytesIO(data))

    summary = wb["Summary"]
    assert summary["A1"].value == "Secure Build — Dependency Vulnerability (SCA) Report"
    labels = [summary.cell(r, 1).value for r in range(1, summary.max_row + 1)]
    assert "Lookup source" in labels
    assert "Vulnerable packages" in labels
    assert "OWASP distribution" not in labels
    assert "QUALITY GATE PASSED" not in labels
    assert summary["A2"].value != "QUALITY GATE PASSED"

    sheet = wb["Findings"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == [
        "#",
        "Severity",
        "Package",
        "Advisory",
        "CVSS",
        "Manifest",
        "Summary",
        "Suggested fix",
    ]
    row = [cell.value for cell in sheet[2]]
    assert row[1] == "CRITICAL"
    assert row[2] == "requests@2.25.0"
    assert row[3] == "CVE-2021-1234"
    assert row[4] == 9.8
    assert row[5] == "requirements.txt"
    assert "request smuggling" in (row[6] or "")
    assert "2.26.0" in (row[7] or "")


def test_sca_empty_findings_still_exports():
    project, scan, _ = _sca_sample()
    assert build_pdf_report(project, scan, [])[:4] == b"%PDF"
    assert build_excel_report(project, scan, [])[:2] == b"PK"


def _construct_rows():
    """3 true-positive rows of one construct + 2 dismissed rows of another."""
    tp = [
        Finding(id=f"tp{i}", scan_id="s1", severity="high", title="SQL concat", rule_id="java-sql-string-concat",
                kind="regex", matched_text='"SELECT * FROM t WHERE id = " + id', verdict="true_positive",
                file_path=f"src/Dao{i}.java", line=10, owasp_category="A03:2021-Injection")
        for i in range(3)
    ]
    fp = [
        Finding(id=f"fp{i}", scan_id="s1", severity="medium", title="Password literal", rule_id="secret-password-assign",
                kind="regex", matched_text='password = "test"', verdict="false_positive",
                file_path=f"src/Fixture{i}.java", line=3, owasp_category="A07:2021-Identification-and-Authentication-Failures")
        for i in range(2)
    ]
    return tp + fp


def test_headline_counts_actionable_groups_with_occurrences():
    grouped = group_findings_by_issue(_construct_rows())
    assert sorted(f.occurrence_count for f in grouped) == [2, 3]

    counts = count_findings(grouped)
    assert counts["actionable"]["high"] == {"groups": 1, "occurrences": 3}
    assert counts["actionable"]["medium"] == {"groups": 0, "occurrences": 0}
    assert counts["dismissed"] == {"groups": 1, "occurrences": 2}


def test_singleton_groups_count_one_occurrence():
    project, scan, findings = _sample()
    grouped = group_findings_by_issue(findings)
    assert [f.occurrence_count for f in grouped] == [1, 1]
    counts = count_findings(grouped)
    assert counts["actionable"]["high"] == {"groups": 1, "occurrences": 1}
    assert counts["dismissed"] == {"groups": 1, "occurrences": 1}


def test_excel_summary_reports_groups_occurrences_and_dismissed():
    project, scan, _ = _sample()
    data = build_excel_report(project, scan, group_findings_by_issue(_construct_rows()))
    summary = load_workbook(BytesIO(data))["Summary"]
    rows = {summary.cell(r, 1).value: [summary.cell(r, c).value for c in (2, 3, 4)]
            for r in range(1, summary.max_row + 1)}
    assert rows["Actionable issues (groups)"][0] == 1
    assert rows["Occurrences"][0] == 3
    assert rows["HIGH"][:2] == [1, 3]
    assert rows["False positive"][:2] == [1, 2]
    # The OWASP table leaves the dismissed category out entirely.
    from app.services.scan_export import OWASP_NAMES

    assert rows[OWASP_NAMES["A03"]] == [1, 1, 3]
    assert OWASP_NAMES["A07"] not in rows


def test_pdf_headline_names_groups_and_dismissed():
    from pypdf import PdfReader

    project, scan, _ = _sample()
    data = build_pdf_report(project, scan, group_findings_by_issue(_construct_rows()))
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages[:1])
    assert "group · 3 occurrences" in text
    flat = " ".join(text.split())  # the note wraps across lines
    assert "Confirmed: 1 group (3 occurrences)" in flat
    assert "Unresolved: 0 groups (0 occurrences)" in flat
    assert "False Positive: 1 group (2 occurrences)" in flat
    assert "Dismissed" not in text


def _llm(fid, cwe, title, path, line, rule_id="llm-extra", verdict="true_positive", kind="llm"):
    return Finding(
        id=fid, scan_id="s1", severity="high", title=title, cwe=cwe, rule_id=rule_id,
        file_path=path, line=line, verdict=verdict, kind=kind,
    )


@pytest.mark.parametrize("cwe_a,cwe_b,title_a,title_b", [
    ("CWE-89", "CWE-089", "SQL Injection via SOAP Message Fields", "SQL Injection via String Concatenation in getActionCode"),
    ("CWE-79", "cwe 79", "Reflected XSS in response body", "Stored cross-site scripting in template"),
    ("CWE-22", "22", "Path traversal in upload handler", "Unsanitised filename used for file read"),
])
def test_consolidates_same_file_same_cwe_whatever_rule_or_title(cwe_a, cwe_b, title_a, title_b):
    findings = [
        _llm("a", cwe_a, title_a, "src/Mgr.java", 10, rule_id="rule-one"),
        _llm("b", cwe_b, title_b, "src/Mgr.java", 90, rule_id="llm-extra"),
        _llm("c", cwe_a, title_a, "src/Mgr.java", 200, rule_id="rule-two"),
    ]
    grouped = group_findings_by_issue(findings)
    assert len(grouped) == 1
    assert grouped[0].occurrence_count == 3
    assert [(x["file_path"], x["line"]) for x in grouped[0].affected_lines] == [
        ("src/Mgr.java", 10), ("src/Mgr.java", 90), ("src/Mgr.java", 200),
    ]


def test_consolidation_does_not_merge_different_files_with_different_topics():
    findings = [
        _llm("a", "CWE-89", "SQL Injection via SOAP Message Fields", "a/One.java", 1, rule_id="r1"),
        _llm("b", "CWE-89", "Dynamic query built from request parameter", "b/Two.java", 1, rule_id="r2"),
    ]
    assert len(group_findings_by_issue(findings)) == 2


def test_consolidation_merges_identifier_only_title_differences_across_files():
    findings = [
        _llm("a", "CWE-89", "SQL Injection via String Concatenation in getActionCode", "a/One.java", 1, rule_id="r1"),
        _llm("b", "CWE-89", "SQL Injection via String Concatenation in getPortMsgType", "b/Two.java", 5, rule_id="r2"),
    ]
    grouped = group_findings_by_issue(findings)
    assert len(grouped) == 1
    assert sorted(grouped[0].affected_files) == ["a/One.java", "b/Two.java"]


def test_consolidation_never_crosses_verdicts_or_missing_cwe():
    findings = [
        _llm("a", "CWE-89", "SQL injection one", "f.java", 1, rule_id="r1"),
        _llm("b", "CWE-89", "SQL injection two", "f.java", 2, rule_id="r2", verdict="false_positive"),
        _llm("c", None, "No cwe one", "f.java", 3, rule_id="r3"),
        _llm("d", None, "No cwe two", "f.java", 4, rule_id="r4"),
    ]
    assert len(group_findings_by_issue(findings)) == 4


@pytest.mark.parametrize("kind", ["regex", "flow", "dast"])
def test_consolidation_leaves_other_kinds_alone(kind):
    findings = [
        _llm("a", "CWE-89", "Thing one", "f.java", 1, rule_id="r1", kind=kind),
        _llm("b", "CWE-89", "Thing two", "f.java", 2, rule_id="r2", kind=kind),
    ]
    assert len(group_findings_by_issue(findings)) == 2
