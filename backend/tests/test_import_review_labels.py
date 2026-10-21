"""The review-label import (SAST-00-080): three exports joined by finding
identity, never by row number. Everything here is a small synthetic triplet;
the real workbooks are supplied by the project owner and never committed."""

import importlib.util
import json
import random
import sys
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "accuracy" / "import_review_labels.py"
spec = importlib.util.spec_from_file_location("import_review_labels", SCRIPT)
irl = importlib.util.module_from_spec(spec)
sys.modules["import_review_labels"] = irl  # dataclasses resolve their module by name
spec.loader.exec_module(irl)

LEGACY_S_HEADERS = [
    "#", "Severity", "Verdict", "OWASP Category", "CWE", "CVSS", "Confidence", "Files", "Issue",
    "Location(s)", "Data Flow", "Risk", "Recommendation", "Example Fix",
]

# number, title, cwe, locations, R original, DEV result, S displayed
FINDINGS = [
    (1, "SQL injection through string concatenation", "CWE-89", ["src/main/java/com/acme/dao/OrderDao.java:12"],
     "True Positive", "False positive", "True Positive"),
    (2, "Sensitive data written to the log", "CWE-532", ["src/main/java/com/acme/AuthService.java:30"],
     "True Positive", "Real security issue", "Needs Review"),
    (3, "Hardcoded password", "CWE-798", ["src/main/java/com/acme/Config.java:5"],
     "Needs Review", "False positive", "Needs Review"),
    (4, "Unvalidated error message assigned to a component property and rendered", "CWE-79",
     ["src/app/shared/error.component.ts:44"], "True Positive", "Code quality", "True Positive"),
    # One group listing three files.
    (5, "Exception message returned in the response body", "CWE-209",
     ["src/main/java/com/acme/web/A.java", "src/main/java/com/acme/web/B.java", "src/main/java/com/acme/web/C.java"],
     "True Positive", "Real, but low risk", "True Positive"),
    (6, "Missing serialVersionUID", "CWE-502", ["src/main/java/com/acme/Dto.java:1"],
     "Unverified (regex)", "Code quality", "Unverified (regex)"),
]
# Present in the PDF only, between #3 and #4 — it shifts the PDF's numbering from there on.
PDF_ONLY = ("Unsanitized API response assigned to component property", "CWE-79", ["src/app/api/client.ts:9"])


def _review_workbook(path: Path, findings=FINDINGS, truncate_title_of=4):
    wb = Workbook()
    ws = wb.active
    ws.title = "Secure Build Findings"
    ws.append(["Finding", "Title", "Location", "CWE", "Severity", "Original SB verdict", "DEV result", "H", "DEV explanation", "J"])
    for number, title, cwe, locations, original, dev, _shown in findings:
        shown_title = title[:40] + "…" if number == truncate_title_of else title
        ws.append([number, shown_title, "\n".join(locations), cwe, "HIGH", original, dev, None,
                   "free text that might quote code: stmt.executeQuery(sql)", None])
    wb.save(path)


def _scanner_workbook(path: Path, findings=FINDINGS, order=None):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Secure Build — SAST Scan Report"])
    ws = wb.create_sheet("Findings")
    ws.append(LEGACY_S_HEADERS)
    rows = list(findings)
    if order is not None:
        rows = [rows[i] for i in order]
    for number, title, cwe, locations, _original, _dev, shown in rows:
        ws.append([number, "HIGH", shown, "A03:2021", cwe, None, "90%", len(locations), title,
                   "\n".join(locations) + "\n\n1: stmt.executeQuery(sql);", "", "", "", ""])
    wb.save(path)


def _pdf(path: Path, findings=FINDINGS, extra=PDF_ONLY, extra_after=3):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    entries = [(t, c, locs, shown) for _n, t, c, locs, _o, _d, shown in findings]
    if extra is not None:
        entries.insert(extra_after, (extra[0], extra[1], extra[2], "True Positive"))
    for number, (title, cwe, locations, shown) in enumerate(entries, start=1):
        pdf.drawString(30, y, f"FINDING {number} {title} HIGH {shown}")
        pdf.drawString(30, y - 14, f"OWASP A03:2021 CWE {cwe} CVSS - Confidence 90%")
        pdf.drawString(30, y - 28, "Affected location(s) " + " ".join(locations))
        y -= 60
        if y < 80:
            pdf.showPage()
            y = 800
    pdf.save()


@pytest.fixture()
def files(tmp_path):
    r, s, p = tmp_path / "R.xlsx", tmp_path / "S.xlsx", tmp_path / "P.pdf"
    _review_workbook(r)
    _scanner_workbook(s)
    _pdf(p)
    return r, s, p


def _load(files, **kwargs):
    r, s, p = files
    return irl.read_review(r), irl.read_scanner(s), irl.read_pdf(p)


def _by_title(joined):
    out = {}
    for entry in joined:
        row = entry.s or entry.r or entry.p
        out[irl.norm_title(row.title)[:30]] = entry
    return out


# --- the readers --------------------------------------------------------------

def test_the_readers_normalise_what_they_find(files):
    review, scanner, pdf = _load(files)
    assert [(r.ordinal, r.cwe, r.verdict, r.dev) for r in review][:3] == [
        (1, "CWE-89", "true_positive", "false_positive"),
        (2, "CWE-532", "true_positive", "real_security_issue"),
        (3, "CWE-798", "needs_review", "false_positive"),
    ]
    assert [s.verdict for s in scanner][:3] == ["true_positive", "needs_review", "needs_review"]
    assert scanner[4].path_set == {
        "src/main/java/com/acme/web/a.java", "src/main/java/com/acme/web/b.java",
        "src/main/java/com/acme/web/c.java",
    }
    assert [(p.ordinal, p.verdict) for p in pdf][:2] == [(1, "true_positive"), (2, "needs_review")]
    assert len(review) == len(scanner) == 6 and len(pdf) == 7


def test_a_scanner_sheet_in_the_current_layout_is_read_by_its_headers(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    ws.append(["#", "Severity", "Scanner Verdict", "Review State", "Issue Class", "Unresolved Reason",
               "Verdict (raw)", "Counted As", "OWASP Category", "CWE", "Issue", "Location(s)", "Occurrence ID"])
    ws.append([1, "HIGH", "False Positive", "pending", "security", "", "false_positive", "False Positive",
               "A03", "CWE-89", "SQL injection", "src/A.java:3", "ab" * 32])
    path = tmp_path / "S.xlsx"
    wb.save(path)
    (row,) = irl.read_scanner(path)
    assert (row.title, row.cwe, row.verdict) == ("SQL injection", "CWE-89", "false_positive")
    assert row.extra["occurrence_id"] == "ab" * 32


def test_the_pdf_reader_understands_the_current_exporter(tmp_path):
    from tests.test_scan_export import _sample
    from app.services.scan_export import build_pdf_report, group_findings_by_issue

    project, scan, findings = _sample()
    path = tmp_path / "current.pdf"
    path.write_bytes(build_pdf_report(project, scan, group_findings_by_issue(findings)))
    rows = irl.read_pdf(path)
    assert [r.ordinal for r in rows] == [1, 2]
    assert rows[0].title.startswith("SQL injection")
    assert rows[0].cwe == "CWE-89"
    assert [r.verdict for r in rows] == ["true_positive", "false_positive"]


# --- the join -----------------------------------------------------------------

def test_an_extra_pdf_finding_does_not_shift_the_join(files):
    review, scanner, pdf = _load(files)
    joined = irl.join(review, scanner, pdf)
    by_title = _by_title(joined)

    for number, title, *_rest in FINDINGS:
        entry = by_title[irl.norm_title(title)[:30]]
        assert entry.sources == ["R", "S", "P"], title
        # The PDF numbered it differently after the inserted finding; the join did not care.
        assert entry.p.ordinal == (number if number <= 3 else number + 1)
        assert entry.s.ordinal == number and entry.r.ordinal == number

    only_p = [e for e in joined if e.sources == ["P"]]
    assert len(only_p) == 1 and only_p[0].p.title == PDF_ONLY[0]
    assert sum(1 for e in joined if e.sources == ["R", "S", "P"]) == 6


def test_every_unambiguous_row_joins_exactly_except_the_truncated_title(files):
    review, scanner, pdf = _load(files)
    by_title = _by_title(irl.join(review, scanner, pdf))
    confidences = {title[:25]: by_title[irl.norm_title(title)[:30]].confidence for _n, title, *_ in FINDINGS}
    # R #4 carries a title cut short with an ellipsis: close, but not provably the same.
    assert confidences["Unvalidated error message"] == irl.FUZZY
    assert {v for k, v in confidences.items() if k != "Unvalidated error message"} == {irl.EXACT}


def test_the_order_of_the_scanner_rows_has_no_effect(tmp_path, files):
    r, _s, p = files
    shuffled = tmp_path / "S_shuffled.xlsx"
    order = list(range(len(FINDINGS)))
    random.Random(7).shuffle(order)
    _scanner_workbook(shuffled, order=order)
    base = irl.join(irl.read_review(r), irl.read_scanner(_s), irl.read_pdf(p))
    other = irl.join(irl.read_review(r), irl.read_scanner(shuffled), irl.read_pdf(p))

    def shape(joined):
        return sorted(
            (irl.norm_title((e.s or e.r or e.p).title), tuple(e.sources), e.confidence) for e in joined
        )

    assert shape(base) == shape(other)


def test_a_review_row_with_no_scanner_row_stays_alone(tmp_path):
    r, s = tmp_path / "R.xlsx", tmp_path / "S.xlsx"
    _review_workbook(r)
    _scanner_workbook(s, findings=FINDINGS[:5])
    review, scanner = irl.read_review(r), irl.read_scanner(s)
    joined = irl.join(review, scanner)
    only_r = [e for e in joined if e.sources == ["R"]]
    assert [e.r.ordinal for e in only_r] == [6]
    assert only_r[0].confidence == irl.NONE


def test_a_cwe_mismatch_never_joins(tmp_path):
    r, s = tmp_path / "R.xlsx", tmp_path / "S.xlsx"
    mismatched = [(*FINDINGS[0][:2], "CWE-79", *FINDINGS[0][3:])]
    _review_workbook(r, findings=mismatched)
    _scanner_workbook(s, findings=FINDINGS[:1])
    joined = irl.join(irl.read_review(r), irl.read_scanner(s))
    assert sorted(e.sources for e in joined) == [["R"], ["S"]]


# --- the records --------------------------------------------------------------

def test_a_grouped_row_expands_to_one_record_per_listed_file(files):
    review, scanner, pdf = _load(files)
    records = irl.to_records(irl.join(review, scanner, pdf))
    group = [r for r in records if r["title"].startswith("Exception message returned")]
    assert len(group) == 3
    assert len({r["group_ref"] for r in group}) == 1
    assert sorted(r["file_path"] for r in group) == [
        "src/main/java/com/acme/web/A.java", "src/main/java/com/acme/web/B.java",
        "src/main/java/com/acme/web/C.java",
    ]
    assert len({r["id"] for r in group}) == 3


def test_a_single_location_keeps_its_line(files):
    review, scanner, pdf = _load(files)
    record = next(r for r in irl.to_records(irl.join(review, scanner, pdf)) if r["cwe"] == "CWE-89")
    assert (record["file_path"], record["line"]) == ("src/main/java/com/acme/dao/OrderDao.java", 12)


def test_the_three_labels_stay_separate_and_nothing_is_adjudicated(files):
    review, scanner, pdf = _load(files)
    records = irl.to_records(irl.join(review, scanner, pdf))
    assert {r["adjudication"] for r in records} == {"dev_label_only"}
    log = next(r for r in records if r["cwe"] == "CWE-532")
    assert (log["scanner_verdict_R"], log["scanner_verdict_S"], log["scanner_label_P"], log["dev_label_R"]) == (
        "true_positive", "needs_review", "needs_review", "real_security_issue")


def test_fuzzy_records_are_flagged_for_a_human_and_pdf_only_rows_are_unjoined(files):
    review, scanner, pdf = _load(files)
    records = irl.to_records(irl.join(review, scanner, pdf))
    fuzzy = [r for r in records if r["join_confidence"] == "fuzzy"]
    assert fuzzy and all(r["needs_manual_review"] for r in fuzzy)
    pdf_only = next(r for r in records if r["sources"] == ["P"])
    assert pdf_only["join_confidence"] == "none" and pdf_only["dev_label_R"] is None


def test_the_labels_hold_no_source_code_and_no_explanations(files):
    review, scanner, pdf = _load(files)
    text = "\n".join(json.dumps(r) for r in irl.to_records(irl.join(review, scanner, pdf)))
    assert "executeQuery" not in text and "free text" not in text
    assert set(irl.to_records(irl.join(review, scanner, pdf))[0]) == {
        "id", "group_ref", "file_path", "line", "title", "cwe", "rule_id", "sources", "join_confidence",
        "needs_manual_review", "ordinal_R", "ordinal_S", "ordinal_P", "scanner_verdict_R",
        "scanner_verdict_S", "scanner_label_P", "dev_label_R", "occurrence_id", "adjudication",
    }


def test_the_summary_counts_the_joins_the_leftovers_and_the_verdict_differences(files):
    review, scanner, pdf = _load(files)
    joined = irl.join(review, scanner, pdf)
    summary = irl.summarize(review, scanner, pdf, joined)
    assert summary["rows"] == {"R": 6, "S": 6, "P": 7}
    assert summary["joined_across_all_sources"] == 6
    assert len(summary["only_in_one_source"]["P"]) == 1 and not summary["only_in_one_source"]["R"]
    # R "True Positive" displayed as "Needs Review" in S: finding 2 only. (#3 is Needs Review in both.)
    assert summary["r_s_verdict_differences"] == 1
    text = irl.format_summary(summary)
    assert "joined across every source:    6" in text and "present only in P" in text


def test_the_command_line_writes_the_jsonl_and_prints_the_summary(files, tmp_path, capsys):
    r, s, p = files
    out = tmp_path / "labels" / "npg.jsonl"
    code = irl.main(["--review", str(r), "--scanner", str(s), "--pdf", str(p), "--out", str(out)])
    assert code == 0
    lines = [json.loads(line) for line in out.read_text().splitlines()]
    # Eight scanner records (the grouped finding lists three files) plus the PDF-only finding.
    assert len(lines) == 9
    assert {line["adjudication"] for line in lines} == {"dev_label_only"}
    captured = capsys.readouterr().out
    assert "Join summary" in captured and "wrote" in captured


def test_a_review_workbook_with_other_columns_can_be_mapped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["n", "x", "y", "dev", "orig", "title", "cwe", "loc"])
    ws.append([7, None, None, "Code quality", "True Positive", "A finding", "CWE-1", "src/A.java:1"])
    path = tmp_path / "R.xlsx"
    wb.save(path)
    columns = irl.parse_columns("number=A,dev=D,original=E,title=F,cwe=G,location=H")
    (row,) = irl.read_review(path, columns)
    assert (row.ordinal, row.title, row.cwe, row.verdict, row.dev) == (7, "A finding", "CWE-1", "true_positive", "code_quality")
