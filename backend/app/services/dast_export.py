"""PDF and XLSX exports for DAST scans (DAST-11-030, -040): the DAST branch of `scan_export`.

Built from the report data set (`dast_report.build_dast_report`), so the files can never disagree
with the report view. They reuse `scan_export`'s page furniture, styles, severity colours and
spreadsheet helpers; SAST, SCA and SBOM exports are untouched.

* Unverified findings are always a separate, labelled section with the amber treatment, never mixed
  into the confirmed list or the severity totals.
* Evidence is redacted again on the way out (key shapes and bearer tokens) as a last line of defence;
  the stored excerpts are already redacted, and sensitive passive matches were never quoted.
* A cell that would start a spreadsheet formula (`=`, `+`, `-`, `@`, and a leading tab or carriage
  return) is stored as text: evidence and titles carry target-controlled text, and opening the file
  must not be able to run anything.
* Tables use wrapping paragraph cells with repeated header rows, so long scans paginate without
  truncated rows.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import CondPageBreak, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session, object_session

from app.dast.domain.redaction import redact_text
from app.models import Finding, Project, Scan
from app.services import scan_export as X
from app.services.dast_report import SEVERITIES, build_dast_report

AMBER = colors.HexColor("#b45309")
AMBER_TINT = colors.HexColor("#fffbeb")
GREEN = colors.HexColor("#047857")
EXCEL_AMBER = PatternFill("solid", fgColor="B45309")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe(text: Any) -> str:
    """Report text with credential shapes masked and illegal characters removed."""
    return X._sanitize_text(redact_text("" if text is None else str(text)))


# ---- the data the exports read -------------------------------------------------------------------------


def report_for(project: Project, scan: Scan) -> dict[str, Any]:
    db: Session | None = object_session(scan)
    if db is None:
        raise RuntimeError("a DAST export needs the scan to be attached to a database session")
    report = build_dast_report(db, scan, project)
    versions = {s.id: s.version for s in db.query(Scan.id, Scan.version).filter(Scan.project_id == scan.project_id)}
    seen = {f.id: (f.first_seen_scan_id, f.last_seen_scan_id) for f in db.query(Finding.id, Finding.first_seen_scan_id, Finding.last_seen_scan_id).filter(Finding.scan_id == scan.id)}
    for item in report["confirmed"] + report["unverified"] + report["dismissed"]:
        first, last = seen.get(item["id"], (None, None))
        item["first_seen"] = f"v{versions[first]}" if first in versions else ""
        item["last_seen"] = f"v{versions[last]}" if last in versions else ""
    return report


# ---- PDF ---------------------------------------------------------------------------------------------------


def _table(rows: list[list[Any]], widths: list[float], styles: Any, *, header_bg: colors.Color = X.INK) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1, splitInRow=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), header_bg),
                ("BOX", (0, 0), (-1, -1), 0.5, X.RULE),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, X.HAIRLINE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 1.6 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6 * mm),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ]
        )
    )
    return table


def _cell(text: Any, styles: Any, style: str = "TableCell") -> Paragraph:
    return Paragraph(X._esc(_safe(text)).replace("\n", "<br/>"), styles[style])


def _head(label: str, styles: Any, *, center: bool = False) -> Paragraph:
    return Paragraph(f"<b>{X._esc(label.upper())}</b>", styles["TableHeadCenter" if center else "TableHead"])


def _meta_grid(report: dict[str, Any], styles: Any, width: float) -> Table:
    scan, settings = report["scan"], report["settings"]
    packs = ", ".join(p["title"] for p in report["rule_packs"]) or "—"
    fields = [
        ("Target", scan["target"] or "—"),
        ("Project", scan["project"] or "—"),
        ("Scan", f"v{scan['version']}"),
        ("Status", scan["status"] or "—"),
        ("Intensity", settings.get("intensity", "—")),
        ("Method policy", "safe (read-only methods)" if settings.get("safe_mode", True) else "state-changing methods allowed"),
        ("Identities", ", ".join(f"{i['label']} ({i['role']})" for i in settings.get("identities", [])) or "anonymous only"),
        ("Rule packs", packs),
    ]
    cols = 4
    rows = [
        [Paragraph(f'<font color="#64748b"><b>{X._esc(label.upper())}</b></font><br/>{X._esc(_safe(value))}', styles["FindingMeta"]) for label, value in fields[i : i + cols]]
        for i in range(0, len(fields), cols)
    ]
    grid = Table(rows, colWidths=[width / cols] * cols)
    grid.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
                ("BOX", (0, 0), (-1, -1), 0.5, X.RULE),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, X.HAIRLINE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 2 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
            ]
        )
    )
    return grid


def _severity_table(report: dict[str, Any], styles: Any, width: float) -> Table:
    counts = report["summary"]["severity_counts"]
    rows = [[_head("Severity", styles), _head("Confirmed findings", styles, center=True)]]
    for sev in SEVERITIES:
        rows.append([_cell(sev.upper(), styles), _cell(counts.get(sev, 0), styles, "TableCellCenter")])
    rows.append([Paragraph("<b>TOTAL CONFIRMED</b>", styles["TableCell"]), Paragraph(f"<b>{report['summary']['confirmed']}</b>", styles["TableCellCenter"])])
    table = _table(rows, [width * 0.5, width * 0.5], styles)
    table.hAlign = "LEFT"
    return table


def _finding_block(idx: int, f: dict[str, Any], *, styles: Any, content_width: float, unverified: bool) -> list[Any]:
    pad = 5 * mm
    inner = content_width - 2 * pad
    sev = X._severity_bucket(f["severity"])
    rail = AMBER if unverified else X.SEVERITY_COLORS[sev]
    tint = AMBER_TINT if unverified else X.SEVERITY_TINTS[sev]
    verdict_label = "UNVERIFIED" if unverified else "CONFIRMED"
    sev_w, verdict_w = 26 * mm, 34 * mm
    title_w = max(40 * mm, inner - sev_w - verdict_w - 4 * mm)
    header = Table(
        [[
            Paragraph(f'<font color="#64748b"><b>{"REVIEW ITEM" if f.get("review_item") else "FINDING"} {idx}</b></font>&nbsp;&nbsp;<b>{X._esc(_safe(f["title"]))}</b>', styles["FindingTitle"]),
            X._pill(sev.upper(), X.SEVERITY_COLORS[sev], styles, width=sev_w),
            X._pill(verdict_label, AMBER if unverified else colors.HexColor("#b91c1c"), styles, width=verdict_w),
        ]],
        colWidths=[title_w, sev_w + 2 * mm, verdict_w + 2 * mm],
    )
    header.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (1, 0), (-1, 0), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (0, 0), 4 * mm), ("RIGHTPADDING", (1, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5 * mm),
            ]
        )
    )
    meta_fields = [
        ("Check", f["check_id"] or "—"),
        ("Target", f["target"] or "—"),
        ("Parameter", f"{f['param']} ({f['location']})" if f["param"] else (f["location"] or "—")),
        ("Confidence", f["confidence"] or "—"),
        ("Standards", ", ".join(f["standards"]) or "—"),
    ]
    weights = [0.18, 0.22, 0.2, 0.1, 0.3]
    meta = Table(
        [[Paragraph(f'<font color="#64748b"><b>{X._esc(label.upper())}</b></font><br/>{X._esc(_safe(value))}', styles["FindingMeta"]) for label, value in meta_fields]],
        colWidths=[inner * w for w in weights],
    )
    meta.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.white), ("BOX", (0, 0), (-1, -1), 0.5, X.HAIRLINE), ("INNERGRID", (0, 0), (-1, -1), 0.4, X.HAIRLINE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm), ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 1.8 * mm), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8 * mm),
            ]
        )
    )
    head_inner = Table([[header], [meta]], colWidths=[inner])
    head_inner.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    strips: list[Any] = [head_inner]
    prose = min(inner, X._PROSE_MAX_WIDTH)
    reason = _safe(f["reason"]).strip()
    if unverified:
        reason = (reason + "\n\n" if reason else "") + "Human review is required: this is not a confirmed vulnerability."
    if reason:
        strips.append(X._labelled_block("Why this was reported" if not unverified else "Why this needs review", reason, styles=styles, width=prose, label_color="#1e3a5f", accent="#fffbeb" if unverified else None))
    if f["roles"]:
        strips.append(X._labelled_block("Role pair (owner → acting)", " → ".join(f["roles"]), styles=styles, width=prose, label_color="#334155"))
    if len(f["affected"]) > 1:
        strips.append(X._labelled_block("Also observed on", "\n".join(_safe(a) for a in f["affected"][:12]) + (f"\n… and {len(f['affected']) - 12} more" if len(f["affected"]) > 12 else ""), styles=styles, width=inner, label_color="#334155"))
    evidence = f.get("evidence")
    if evidence:
        request_text = evidence.get("request_text") or ""
        if request_text:
            strips.extend(X._code_boxes("Request sent", _safe(request_text), styles=styles, width=inner))
        if evidence["baseline"]:
            strips.extend(X._code_boxes("Baseline response", _safe(evidence["baseline"]), styles=styles, width=inner))
        if evidence["test"]:
            strips.extend(X._code_boxes("Response received" if evidence.get("request_text") or evidence["baseline"] else "Observed", _safe(evidence["test"]), styles=styles, width=inner))
        if evidence["signals"]:
            strips.append(X._labelled_block("Detector signals", "\n".join(f"{s.get('name')}: {s.get('detail', '')}".strip(": ") for s in evidence["signals"]), styles=styles, width=prose, label_color="#0369a1", accent="#f0f9ff"))
    if f["remediation"]:
        strips.append(X._labelled_block("Remediation", _safe(f["remediation"]), styles=styles, width=prose, label_color="#065f46", accent="#ecfdf5"))
    out: list[Any] = [X._Bookmark(f"{idx}. [{sev.upper()}] {f['title']}"[:120], f"{'u' if unverified else 'f'}-{idx}"), CondPageBreak(45 * mm)]
    for pos, strip in enumerate(strips):
        out.append(X._rail(strip, content_width=content_width, pad=pad, sev_color=rail, bg=tint, first=(pos == 0), last=(pos == len(strips) - 1)))
    out.append(Spacer(1, 5 * mm))
    return out


def _skipped_table(report: dict[str, Any], styles: Any, width: float) -> Any:
    rows = [[_head(h, styles) for h in ("Check", "Status", "Needs", "Ran as", "Reason")]]
    for item in report["skipped_degraded"]:
        rows.append([_cell(item.get("title") or item.get("check_id"), styles), _cell(item.get("status"), styles), _cell(item.get("dependency"), styles), _cell(item.get("effective_mode"), styles), _cell(item.get("reason"), styles)])
    return _table(rows, [width * w for w in (0.22, 0.1, 0.12, 0.12, 0.44)], styles)


def _diff_section(report: dict[str, Any], styles: Any, width: float) -> list[Any]:
    diff = report["diff"]
    out: list[Any] = [Paragraph("CHANGES SINCE THE PREVIOUS SCAN", styles["SectionHeading"])]
    if not diff["baseline"]:
        out.append(Paragraph("No baseline: there is no earlier successful scan of this target, so every finding is new.", styles["FindingBody"]))
    else:
        out.append(Paragraph(f"Compared with scan v{diff['previous_version']}.", styles["FindingMeta"]))
    counts = diff["counts"]
    out.append(Paragraph(f"<b>NEW</b> {counts['new']} &nbsp;·&nbsp; <b>FIXED</b> {counts['fixed']} &nbsp;·&nbsp; <b>STILL OPEN</b> {counts['still_open']}", styles["FindingBody"]))
    out.append(Spacer(1, 2 * mm))
    for note in diff["notes"]:
        if diff["baseline"] or "No baseline" not in note:
            out.append(Paragraph(X._esc(_safe(note)), styles["FindingMeta"]))
    for status, label in (("new", "NEW"), ("fixed", "FIXED"), ("still_open", "STILL OPEN")):
        entries = [e for e in diff["entries"] if e["status"] == status]
        if not entries:
            continue
        rows = [[_head(h, styles) for h in (label, "Check", "Target", "Parameter", "Severity", "Verdict")]]
        for e in entries:
            rows.append([_cell(label, styles), _cell(e["check_id"], styles), _cell(e["target"], styles), _cell(e["param"] or "—", styles), _cell(e["severity"], styles), _cell("Unverified" if e["verdict"] == "unverified" else "Confirmed" if e["verdict"] == "true_positive" else e["verdict"], styles)])
        out.append(Spacer(1, 2 * mm))
        out.append(_table(rows, [width * w for w in (0.1, 0.2, 0.28, 0.14, 0.12, 0.16)], styles))
    if diff["not_comparable"]:
        rows = [[_head(h, styles) for h in ("Not comparable", "Check", "Target", "Reason")]]
        for n in diff["not_comparable"]:
            rows.append([_cell("—", styles), _cell(n["check_id"], styles), _cell(n["target"], styles), _cell(n["reason"], styles)])
        out.append(Spacer(1, 2 * mm))
        out.append(_table(rows, [width * w for w in (0.12, 0.2, 0.28, 0.4)], styles))
    return out


def build_dast_pdf(project: Project, scan: Scan, findings: list[Finding] | None = None) -> bytes:
    report = report_for(project, scan)
    buf = io.BytesIO()
    page = landscape(A4)
    left = right = 13 * mm
    content_width = page[0] - left - right
    doc = SimpleDocTemplate(
        buf, pagesize=page, leftMargin=left, rightMargin=right, topMargin=12 * mm, bottomMargin=18 * mm,
        title=f"{project.name} — Secure Build DAST Report v{scan.version}", author="Secure Build", subject="Dynamic application security testing report",
    )
    styles = X._build_styles()
    summary = report["summary"]
    story: list[Any] = [
        X._Bookmark("Summary", "summary"),
        Paragraph(X._esc(_safe(project.name or "Application")), styles["ReportTitle"]),
        Paragraph(f"Secure Build — DAST Security Report&nbsp;&nbsp;·&nbsp;&nbsp;v{scan.version}&nbsp;&nbsp;·&nbsp;&nbsp;{X._esc(X._fmt_dt(scan.created_at))}", styles["ReportSub"]),
    ]
    note = X._scan_status_banner(scan)
    if note:
        story.append(Paragraph(f'<font color="#b45309"><b>{X._esc(note)}</b></font>', styles["ReportSub"]))
    story += [_meta_grid(report, styles, content_width), Spacer(1, 5 * mm)]
    story.append(X._severity_strip(summary["severity_counts"], styles, content_width))
    story.append(Spacer(1, 3 * mm))
    overall = (
        "Issues found"
        if summary["confirmed"]
        else "Needs review"
        if summary["unverified"]
        else "No confirmed issues"
    )
    story.append(Paragraph(
        f"<b>Overall:</b> {overall}. <b>{summary['confirmed']}</b> confirmed finding(s) are counted above. <font color=\"#b45309\"><b>{summary['unverified']} unverified</b></font> "
        f"finding(s) ({summary['review_items']} listed for human review) are reported separately and are not part of the severity totals.",
        styles["FindingBody"],
    ))
    story += [Spacer(1, 5 * mm), Paragraph("SEVERITY SUMMARY", styles["SectionHeading"]), _severity_table(report, styles, content_width * 0.5)]

    story += [PageBreak(), X._Bookmark("Confirmed findings", "confirmed"), Paragraph("CONFIRMED FINDINGS", styles["SectionHeading"])]
    if report["confirmed"]:
        for idx, f in enumerate(report["confirmed"], start=1):
            story += _finding_block(idx, f, styles=styles, content_width=content_width, unverified=False)
    else:
        story.append(Paragraph("No confirmed findings.", styles["FindingBody"]))

    story += [PageBreak(), X._Bookmark("Unverified findings", "unverified"), Paragraph('<font color="#b45309">UNVERIFIED — HUMAN REVIEW REQUIRED</font>', styles["SectionHeading"])]
    story.append(Paragraph("These are leads, not confirmed vulnerabilities: a check that could not be proven either way, or an area a scanner cannot test (business flows, unsafe API consumption, SSRF without an out-of-band address).", styles["FindingMeta"]))
    story.append(Spacer(1, 3 * mm))
    if report["unverified"]:
        for idx, f in enumerate(report["unverified"], start=1):
            story += _finding_block(idx, f, styles=styles, content_width=content_width, unverified=True)
    else:
        story.append(Paragraph("No unverified findings.", styles["FindingBody"]))

    story += [PageBreak(), X._Bookmark("Skipped and degraded checks", "skipped"), Paragraph("SKIPPED AND DEGRADED CHECKS", styles["SectionHeading"])]
    if report["skipped_degraded"]:
        story.append(_skipped_table(report, styles, content_width))
    else:
        story.append(Paragraph("Every enabled check ran as designed.", styles["FindingBody"]))

    story += [Spacer(1, 6 * mm), Paragraph("RULE PACKS AND SETTINGS", styles["SectionHeading"])]
    packs = ", ".join(f"{p['title']} ({p['id']})" for p in report["rule_packs"]) or "—"
    custom = ", ".join(f"{c['id']} v{c['version']}" for c in report["custom_rules"]) or "none"
    settings = report["settings"]
    story.append(Paragraph(
        f"<b>Rule packs:</b> {X._esc(_safe(packs))}<br/><b>Catalog version:</b> {X._esc(report['catalog_version'] or '—')}<br/><b>Custom rules:</b> {X._esc(custom)}<br/>"
        f"<b>Checks enabled:</b> {settings.get('rule_checks_enabled', '—')} of {settings.get('rule_checks_total', '—')}<br/>"
        f"<b>Budgets:</b> up to {settings.get('max_targets', '—')} endpoints, {settings.get('max_tests_per_target', '—')} tests each; rate limit {settings.get('rate_limit_per_sec', '—')} requests/s",
        styles["FindingBody"],
    ))

    story += [PageBreak(), X._Bookmark("Changes since the previous scan", "diff")] + _diff_section(report, styles, content_width)
    doc.build(story, canvasmaker=X._make_canvas(project.name or "Application", scan.version, left, right, page, footer_prefix="Secure Build DAST"))
    return buf.getvalue()


# ---- XLSX --------------------------------------------------------------------------------------------------


def _xl_text(value: Any) -> str:
    return X._excel_cell(_safe(value))


def _put(sheet: Any, row: int, col: int, value: Any) -> Any:
    """Write a cell. Text that would start a formula is stored as text and quote-prefixed, so opening
    the file cannot run it."""
    cell = sheet.cell(row, col)
    if isinstance(value, str):
        text = _xl_text(value)
        cell.value = text
        if text.startswith(_FORMULA_PREFIXES):
            cell.data_type = "s"
            cell.quotePrefix = True
    else:
        cell.value = value
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = X._THIN
    return cell


def _finding_rows(items: list[dict[str, Any]]) -> list[list[Any]]:
    out = []
    for f in items:
        evidence = f.get("evidence") or {}
        out.append([
            f["check_id"], f["title"], f["severity"], f["confidence"], "Unverified" if f["verdict"] == "unverified" else "Confirmed" if f["verdict"] == "true_positive" else f["verdict"],
            f["target"], f["param"] or "", f["location"] or "", ", ".join(f["standards"]), f["remediation"], f.get("first_seen", ""), f.get("last_seen", ""),
            " → ".join(f["roles"]), f.get("diff_status", ""), evidence.get("request_text", ""), evidence.get("baseline", ""), evidence.get("test", ""), f["reason"], f["source"],
        ])
    return out


FINDING_HEADERS = [
    "Check id", "Title", "Severity", "Confidence", "Verdict", "Target", "Parameter", "Location", "Standards", "Remediation", "First seen", "Last seen",
    "Role pair", "Since previous scan", "Request sent", "Baseline response", "Evidence (test response)", "Why", "Source",
]
FINDING_WIDTHS = [22, 40, 11, 11, 13, 30, 14, 11, 28, 60, 10, 10, 16, 14, 50, 40, 60, 50, 18]


def _sheet_table(sheet: Any, headers: list[str], rows: list[list[Any]], widths: list[int], *, header_fill: PatternFill | None = None) -> None:
    X._xl_header_row(sheet, 1, headers, widths=widths)
    if header_fill is not None:
        for col in range(1, len(headers) + 1):
            sheet.cell(1, col).fill = header_fill
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            _put(sheet, r, c, value)
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{X.get_column_letter(len(headers))}{len(rows) + 1}"


def build_dast_xlsx(project: Project, scan: Scan, findings: list[Finding] | None = None) -> bytes:
    report = report_for(project, scan)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    s = report["summary"]
    settings, diff = report["settings"], report["diff"]
    lines: list[tuple[str, Any]] = [
        ("Project", report["scan"]["project"]), ("Target", report["scan"]["target"]), ("Scan", f"v{scan.version}"), ("Status", scan.status),
        ("Created", X._fmt_dt(scan.created_at)), ("Intensity", settings.get("intensity", "")),
        ("Method policy", "safe (read-only methods)" if settings.get("safe_mode", True) else "state-changing methods allowed"),
        ("Rule packs", ", ".join(p["title"] for p in report["rule_packs"])), ("Catalog version", report["catalog_version"]),
        ("Checks enabled", f"{settings.get('rule_checks_enabled', '')} of {settings.get('rule_checks_total', '')}"),
        ("", ""), ("Confirmed findings", s["confirmed"]),
    ]
    lines += [(sev.upper(), s["severity_counts"][sev]) for sev in SEVERITIES]
    lines += [("Unverified findings (reported separately)", s["unverified"]), ("  of which listed for human review", s["review_items"]), ("Dismissed", s["dismissed"]),
              ("", ""), ("New since previous scan", diff["counts"]["new"]), ("Fixed", diff["counts"]["fixed"]), ("Still open", diff["counts"]["still_open"]),
              ("Comparison", "no baseline: every finding is new" if not diff["baseline"] else f"against v{diff['previous_version']}"),
              ("Skipped or degraded checks", len(report["skipped_degraded"]))]
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 70
    for r, (label, value) in enumerate(lines, start=1):
        a, b = _put(summary, r, 1, label), _put(summary, r, 2, value)
        a.font = Font(bold=True)
        sev = label.lower()
        if sev in X.EXCEL_SEVERITY_FILLS:
            a.fill, a.font = X.EXCEL_SEVERITY_FILLS[sev], Font(bold=True, color="FFFFFF")
        if label.startswith("Unverified"):
            a.fill, a.font = EXCEL_AMBER, Font(bold=True, color="FFFFFF")

    _sheet_table(wb.create_sheet("Findings"), FINDING_HEADERS, _finding_rows(report["confirmed"]), FINDING_WIDTHS)
    _sheet_table(wb.create_sheet("Unverified"), FINDING_HEADERS, _finding_rows(report["unverified"]), FINDING_WIDTHS, header_fill=EXCEL_AMBER)
    diff_rows = [[e["status"].replace("_", " ").title(), e["check_id"], e["target"], e["param"] or "", e["severity"], e["verdict"], " → ".join(e["roles"]), e["note"]] for e in diff["entries"]]
    diff_rows += [["Not comparable", n["check_id"], n["target"], n["param"] or "", n["severity"], "", "", n["reason"]] for n in diff["not_comparable"]]
    _sheet_table(wb.create_sheet("Diff"), ["Status", "Check id", "Target", "Parameter", "Severity", "Verdict", "Role pair", "Note"], diff_rows, [16, 24, 34, 14, 11, 14, 18, 60])
    skipped_rows = [[i.get("check_id"), i.get("title"), i.get("status"), i.get("dependency"), i.get("effective_mode"), i.get("reason")] for i in report["skipped_degraded"]]
    _sheet_table(wb.create_sheet("Skipped Checks"), ["Check id", "Title", "Status", "Needs", "Ran as", "Reason"], skipped_rows, [26, 36, 12, 14, 14, 70])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
