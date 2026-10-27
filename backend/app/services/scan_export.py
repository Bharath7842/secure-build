"""Generate downloadable PDF and Excel reports for a completed scan.

Layout notes
------------
PDF is landscape A4. The width goes where it earns its keep — code context,
location lists, and the findings index — while prose blocks stay capped at a
readable measure (`_PROSE_MAX_WIDTH`) rather than stretching to 270mm.

Each finding is ONE card: a continuous severity rail down the left edge, with
sections separated by hairline rules. The previous revision wrapped every
section in its own bordered shell which then contained another bordered box,
so a single finding rendered as seven double-bordered fragments. Sections
still break across pages independently; the rail is what makes them read as
one unit.

Page furniture: quality-gate banner, severity strip, OWASP distribution, and a
findings index on the summary page; PDF outline bookmarks per finding; and
"Page X of Y" via a two-pass canvas.

SCA scans (`scan.scan_type == "sca"`) use a separate layout: dependency
inventory + advisory cards (package / CVE / CVSS / suggested upgrade) with no
OWASP, CWE, quality-gate, data-flow, or code-context sections.

SBOM scans (`scan.scan_type == "sbom"`) use a third layout: component
inventory + license/cost-classification breakdown, sourced from
`scan.metadata_json["discovery"]["components"]` rather than `Finding` rows —
an SBOM scan writes no findings, so the `findings` argument to
`build_pdf_report`/`build_excel_report` is ignored on this path. PDF/Excel are
the only export formats for SBOM — the CycloneDX/SPDX JSON documents in
`discovery` are internal building blocks, not a separate download surface.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.platypus import (
    CondPageBreak,
    Flowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.models import ISSUE_CLASSES, UNVERIFIED_VERDICT, Finding, Project, Scan
from app.services.gates import build_gate_summary
from app.services.git_workspace import redact_git_url
from app.services.module_path import to_module_path
from app.services.occurrence import semantic_signature
from app.services.severity import severity_basis_text

logger = logging.getLogger(__name__)

SEVERITY_ORDER = ("critical", "high", "medium", "low", "info")

SEVERITY_COLORS = {
    "critical": colors.HexColor("#7f1d1d"),
    "high": colors.HexColor("#b91c1c"),
    "medium": colors.HexColor("#c2410c"),
    "low": colors.HexColor("#a16207"),
    "info": colors.HexColor("#475569"),
}

# Soft background tints matching the severity rail, used for the finding card
# and the index rows so severity is readable without colour-coding every cell.
SEVERITY_TINTS = {
    "critical": colors.HexColor("#fef2f2"),
    "high": colors.HexColor("#fef2f2"),
    "medium": colors.HexColor("#fff7ed"),
    "low": colors.HexColor("#fefce8"),
    "info": colors.HexColor("#f8fafc"),
}

EXCEL_SEVERITY_FILLS = {
    "critical": PatternFill("solid", fgColor="7F1D1D"),
    "high": PatternFill("solid", fgColor="B91C1C"),
    "medium": PatternFill("solid", fgColor="C2410C"),
    "low": PatternFill("solid", fgColor="A16207"),
    "info": PatternFill("solid", fgColor="475569"),
}

# SBOM license/cost-classification buckets — mirrors sbom_scan._COST_ORDER.
COST_ORDER = ("fully_free", "partially_free", "fully_paid", "unknown")
COST_LABELS = {
    "fully_free": "Fully Free",
    "partially_free": "Partially Free",
    "fully_paid": "Fully Paid",
    "unknown": "Unknown",
}
COST_COLORS = {
    "fully_free": colors.HexColor("#047857"),
    "partially_free": colors.HexColor("#c2410c"),
    "fully_paid": colors.HexColor("#b91c1c"),
    "unknown": colors.HexColor("#475569"),
}
EXCEL_COST_FILLS = {
    "fully_free": PatternFill("solid", fgColor="047857"),
    "partially_free": PatternFill("solid", fgColor="C2410C"),
    "fully_paid": PatternFill("solid", fgColor="B91C1C"),
    "unknown": PatternFill("solid", fgColor="475569"),
}

# SBOM component "Type" column — sbom_scan.build_components's origin_type.
TYPE_COLORS = {
    "Open Source": colors.HexColor("#047857"),
    "Proprietary": colors.HexColor("#b91c1c"),
    "Internal": colors.HexColor("#475569"),
}
EXCEL_TYPE_FILLS = {
    "Open Source": PatternFill("solid", fgColor="047857"),
    "Proprietary": PatternFill("solid", fgColor="B91C1C"),
    "Internal": PatternFill("solid", fgColor="475569"),
}

OWASP_NAMES = {
    "A01": "Broken Access Control",
    "A02": "Cryptographic Failures",
    "A03": "Injection",
    "A04": "Insecure Design",
    "A05": "Security Misconfiguration",
    "A06": "Vulnerable and Outdated Components",
    "A07": "Identification and Authentication Failures",
    "A08": "Software and Data Integrity Failures",
    "A09": "Security Logging and Monitoring Failures",
    "A10": "Server-Side Request Forgery (SSRF)",
}

INK = colors.HexColor("#152238")
MUTED = colors.HexColor("#64748b")
HAIRLINE = colors.HexColor("#e2e8f0")
RULE = colors.HexColor("#cbd5e1")

# PDF: keep individual section blocks flowable so ReportLab can page them.
_PDF_CODE_MAX = 6000
_PDF_SECTION_MAX = 2500
_PDF_CODE_LINES_PER_CHUNK = 22
# Body prose stops widening past this; landscape width goes to code and tables.
_PROSE_MAX_WIDTH = 170 * mm
_CODE_FONT_SIZE = 8.5
# Courier advance width is 0.6em, so chars-per-line = width_pt / (size * 0.6).
_COURIER_ADVANCE = 0.6
# Excel cell hard limit is 32767; leave headroom and mark truncation.
_EXCEL_CELL_MAX = 32000
_INDEX_MAX_ROWS = 60


def _llm_display(scan) -> str:
    """`provider / model` for the report header; the CLI providers read as the plain model name."""
    provider = (scan.llm_provider or "").strip().lower()
    if provider in ("claude_cli", "openai_cli"):
        from app.services.scan_runner import _friendly_model_label

        return _friendly_model_label(provider, scan.llm_model) or ("OpenAI" if provider == "openai_cli" else "Claude")
    return f"{scan.llm_provider or '—'} / {scan.llm_model or '—'}"


def _sanitize_text(text: str) -> str:
    """Strip XML/Excel-illegal control chars (e.g. binary .class snippets)."""
    return ILLEGAL_CHARACTERS_RE.sub("", text)


def _esc(text: Any) -> str:
    raw = _sanitize_text("" if text is None else str(text))
    return (
        raw.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _hex(color: colors.Color) -> str:
    """'#rrggbb' for embedding a Color in Paragraph markup."""
    return "#" + color.hexval()[2:].rjust(6, "0")[-6:]


def _severity_key(value: str | None) -> str:
    return (value or "info").strip().lower()


def _severity_bucket(value: str | None) -> str:
    """Normalise to a known severity. Unknown labels fall into `info` so they
    stay visible in the breakdown — the previous version created an orphan key
    that the display loop never iterated, so those findings silently vanished
    from the severity table while still inflating the total."""
    key = _severity_key(value)
    return key if key in SEVERITY_ORDER else "info"


# How an `unresolved_reason` reads after "Unresolved (" in a report.
_UNRESOLVED_LABELS = {
    "not_reviewed_regex": "regex, not reviewed",
    "not_reviewed_joern": "code-graph query, not reviewed",
    "not_reviewed_sca": "dependency match, not reviewed",
    "llm_missing_verdict": "model gave no verdict",
    "llm_unrecognised_verdict": "model verdict not recognised",
    "llm_call_failed": "model call failed",
    "llm_budget_exhausted": "LLM budget reached",
    "cap_reached": "review cap reached",
    "evidence_missing": "evidence missing",
}


def _display_verdict(value: str | None, unresolved_reason: str | None = None) -> str:
    """The scanner's verdict by its real name. FALSE_POSITIVE is "False Positive" —
    never "Needs Review", which is a review *state*, not a verdict (see
    `_review_state_label`). An `unverified` verdict reads "Unresolved" and, when
    the reason is known, says why. A value outside the contract reads
    "Unknown (<raw>)" and is logged: it means a non-enum verdict reached the
    export layer."""
    raw = (value or "").strip()
    if not raw:
        return "—"
    key = raw.lower().replace("-", "_").replace(" ", "_")
    if key == "true_positive":
        return "True Positive"
    if key == "false_positive":
        return "False Positive"
    if key == UNVERIFIED_VERDICT:
        if unresolved_reason:
            return f"Unresolved ({_UNRESOLVED_LABELS.get(unresolved_reason, unresolved_reason)})"
        return "Unresolved"
    logger.warning("export saw a verdict outside the contract: %r", value)
    return f"Unknown ({raw})"


def _scanner_verdict(finding: Finding) -> str | None:
    """The verdict the scanner stored. Reading a finding through
    `apply_effective_verdicts` overwrites `verdict` with the reviewer's and keeps
    the scanner's in `llm_verdict`; exports must show the scanner's."""
    return getattr(finding, "llm_verdict", None) or finding.verdict


_REVIEWER_VERDICT_LABELS = {"confirmed_issue": "True Positive", "false_positive": "False Positive"}


def _review_state(finding: Finding) -> str:
    return getattr(finding, "review_state", None) or "pending"


def _review_state_label(finding: Finding) -> str:
    """`pending`, `agreed`, `overridden → <reviewer verdict>` or `resolved`."""
    state = _review_state(finding)
    if state == "overridden":
        reviewer = (getattr(finding, "review", None) or {}).get("reviewer_verdict")
        if reviewer in _REVIEWER_VERDICT_LABELS:
            return f"overridden → {_REVIEWER_VERDICT_LABELS[reviewer]}"
    return state


def _fmt_dt(value: datetime | None) -> str:
    if value is None:
        return "—"
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _is_sca_scan(scan: Scan) -> bool:
    return (getattr(scan, "scan_type", None) or "sast").strip().lower() == "sca"


def _is_dast_scan(scan: Scan) -> bool:
    return (getattr(scan, "scan_type", None) or "").strip().lower() == "dast"


def _is_sbom_scan(scan: Scan) -> bool:
    return (getattr(scan, "scan_type", None) or "sast").strip().lower() == "sbom"


def _safe_filename(project_name: str, version: int, ext: str, *, kind: str | None = None) -> str:
    slug = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in (project_name or "scan"))
    slug = "-".join(part for part in slug.split("-") if part)[:60] or "scan"
    if kind:
        return f"{slug}-{kind}-v{version}.{ext}"
    return f"{slug}-v{version}.{ext}"


def export_filenames(project: Project, scan: Scan) -> dict[str, str]:
    kind = "sca" if _is_sca_scan(scan) else "sbom" if _is_sbom_scan(scan) else "dast" if _is_dast_scan(scan) else None
    return {
        "pdf": _safe_filename(project.name, scan.version, "pdf", kind=kind),
        "xlsx": _safe_filename(project.name, scan.version, "xlsx", kind=kind),
    }


def _scan_status_banner(scan: Scan) -> str | None:
    if scan.status == "cancelled":
        return "Partial report — scan was stopped before completion."
    if scan.status == "failed":
        return "Scan failed — report includes findings saved before the failure."
    return None


def _count_by_severity(findings: list[Finding]) -> dict[str, int]:
    counts = {sev: 0 for sev in SEVERITY_ORDER}
    for finding in findings:
        counts[_severity_bucket(finding.severity)] += 1
    return counts


def _occurrences(finding: Finding) -> int:
    """Raw rows behind one grouped finding (1 for an ungrouped row)."""
    return getattr(finding, "occurrence_count", None) or 1


# Machine-readable anchors in the exports that `check_report_consistency` reads back.
BUCKET_ROW_PREFIX = "Groups — "          # XLSX summary rows: "Groups — confirmed" | groups | occurrences
COUNTED_AS_PREFIX = "Counted as: "       # each PDF finding block: "... · Counted as: Confirmed"

# The verdict buckets every headline count is split into, in report order. A
# group's bucket is its effective verdict (the reviewer's, once there is one);
# `_verdict_bucket` is the one place that decides it, and the seam where mixed
# groups (EPIC-01) will extend counting.
VERDICT_BUCKETS = ("true_positive", UNVERIFIED_VERDICT, "false_positive")
UNKNOWN_BUCKET = "unknown"
UNCLASSIFIED = "unclassified"

# Report names for the verdict buckets.
_BUCKET_LABELS = {
    "true_positive": "Confirmed",
    UNVERIFIED_VERDICT: "Unresolved",
    "false_positive": "False Positive",
    UNKNOWN_BUCKET: "Unknown",
}


def _verdict_bucket(finding: Finding) -> str:
    key = (finding.verdict or "").strip().lower()
    return key if key in VERDICT_BUCKETS else UNKNOWN_BUCKET


def _counted_as(finding: Finding) -> str:
    """The report bucket this group is counted in, by name."""
    return _BUCKET_LABELS[_verdict_bucket(finding)]


def _empty_cell() -> dict[str, int]:
    return {"groups": 0, "occurrences": 0}


def count_findings(grouped: list[Finding]) -> dict[str, Any]:
    """THE counting function. Every headline number in the API, the PDF, the
    Excel export, the quality gate and the UI chips comes from here, over
    `group_findings_by_issue` output (group once per request and pass the result
    around — grouping mutates ORM instances).

    Returns, all in groups and occurrences (rows behind the groups):

    * `by_verdict`  — {bucket: {severity: cell}} for `true_positive`,
      `unverified` and `false_positive`;
    * `totals`      — {bucket: cell} over all severities (plus `unknown` when a
      value outside the contract is present, so it cannot vanish);
    * `by_issue_class` — {security|quality|inactive|unclassified: cell};
    * `actionable`  — {severity: cell} for true positive + unverified groups
      (the headline: what still needs attention), and `dismissed` — the
      false-positive cell. Both keep the shape older clients read;
    * `by_owasp`    — [{category, groups, severe, occurrences}] over actionable
      groups, ordered by size;
    * `by_unresolved_reason` — {reason: occurrences} over the unresolved bucket,
      largest first (`unknown` for a row that names none).
    """
    by_verdict = {
        bucket: {sev: _empty_cell() for sev in SEVERITY_ORDER} for bucket in VERDICT_BUCKETS
    }
    unknown = {sev: _empty_cell() for sev in SEVERITY_ORDER}
    by_issue_class = {name: _empty_cell() for name in (*sorted(ISSUE_CLASSES), UNCLASSIFIED)}
    by_reason: dict[str, int] = {}
    for finding in grouped:
        bucket = _verdict_bucket(finding)
        if bucket == UNVERIFIED_VERDICT:
            reasons = getattr(finding, "unresolved_reason_counts", None) or {
                finding.unresolved_reason or UNKNOWN_BUCKET: _occurrences(finding)
            }
            for reason, n in reasons.items():
                by_reason[reason] = by_reason.get(reason, 0) + n
        severity = _severity_bucket(finding.severity)
        occurrences = _occurrences(finding)
        issue_class = getattr(finding, "issue_class", None)
        issue_class = issue_class if issue_class in ISSUE_CLASSES else UNCLASSIFIED
        for cell in ((by_verdict.get(bucket) or unknown)[severity], by_issue_class[issue_class]):
            cell["groups"] += 1
            cell["occurrences"] += occurrences

    def total(per_severity: dict[str, dict[str, int]]) -> dict[str, int]:
        return {
            "groups": sum(c["groups"] for c in per_severity.values()),
            "occurrences": sum(c["occurrences"] for c in per_severity.values()),
        }

    totals = {bucket: total(by_verdict[bucket]) for bucket in VERDICT_BUCKETS}
    if total(unknown)["groups"]:
        totals[UNKNOWN_BUCKET] = total(unknown)
    actionable = {
        sev: {
            "groups": by_verdict["true_positive"][sev]["groups"] + by_verdict[UNVERIFIED_VERDICT][sev]["groups"]
            + unknown[sev]["groups"],
            "occurrences": by_verdict["true_positive"][sev]["occurrences"]
            + by_verdict[UNVERIFIED_VERDICT][sev]["occurrences"] + unknown[sev]["occurrences"],
        }
        for sev in SEVERITY_ORDER
    }
    return {
        "by_verdict": by_verdict,
        "totals": totals,
        "by_issue_class": by_issue_class,
        "actionable": actionable,
        "dismissed": dict(totals["false_positive"]),
        "by_owasp": [
            {"category": label, "groups": groups, "severe": severe, "occurrences": occurrences}
            for label, groups, severe, occurrences in _count_by_owasp(grouped)
        ],
        "by_unresolved_reason": dict(sorted(by_reason.items(), key=lambda kv: (-kv[1], kv[0]))),
    }


def verdict_bucket_totals(counts: dict[str, Any]) -> dict[str, dict[str, int]]:
    """{bucket label: cell} from `count_findings`, in report order — the names the
    PDF cover and Excel summary use (Confirmed / Unresolved / False Positive)."""
    return {
        _BUCKET_LABELS[bucket]: counts["totals"][bucket]
        for bucket in (*VERDICT_BUCKETS, UNKNOWN_BUCKET)
        if bucket in counts["totals"]
    }


def _count_by_owasp(findings: list[Finding]) -> list[tuple[str, int, int, int]]:
    """(category label, actionable groups, critical+high groups, occurrences)
    over grouped findings, dismissed groups left out; ordered by risk then size."""
    buckets: dict[str, list[int]] = {}
    for finding in findings:
        if _verdict_bucket(finding) == "false_positive":
            continue
        label = _owasp_header(finding)
        entry = buckets.setdefault(label, [0, 0, 0])
        entry[0] += 1
        if _severity_bucket(finding.severity) in ("critical", "high"):
            entry[1] += 1
        entry[2] += _occurrences(finding)
    rows = [(label, total, severe, occ) for label, (total, severe, occ) in buckets.items()]
    rows.sort(key=lambda r: (-r[2], -r[1], r[0]))
    return rows


def _owasp_header(finding: Finding) -> str:
    cat = (finding.owasp_category or "").strip().upper()
    if not cat:
        return "Uncategorized"
    key = cat.replace(":2021", "")[:3]
    name = OWASP_NAMES.get(key)
    if name:
        return name
    raw = (finding.owasp_category or "").strip()
    stripped = re.sub(r"^A\d{2}(?::\d{4})?\s*[-–—:]?\s*", "", raw, flags=re.I).strip()
    return stripped or raw or "Uncategorized"


def _affected_file_count(finding: Finding) -> int:
    if finding.affected_files:
        return len(finding.affected_files)
    return 1 if finding.file_path else 0


def _risk_text(finding: Finding) -> str:
    return (getattr(finding, "risk", None) or finding.verdict_reason or "").strip()


def _recommendation_text(finding: Finding) -> str:
    example = (getattr(finding, "example_fix", None) or "").strip()
    ai = (finding.ai_fix or "").strip()
    if not ai:
        return ""
    if not example:
        return ai
    # ai_fix is often recommendation + example_fix joined by blank lines.
    return ai.split("\n\n", 1)[0].strip()


def _example_fix_text(finding: Finding) -> str:
    return (getattr(finding, "example_fix", None) or "").strip()


def _owasp_display(finding: Finding) -> str:
    """Full OWASP label for reports, e.g. A03:2021 – Injection.

    The separator match allows surrounding whitespace; the previous pattern
    required the dash to sit flush against the year, so the common
    "A03:2021 - Injection" form fell through to the raw string."""
    raw = (finding.owasp_category or "").strip()
    if not raw:
        return "Uncategorized"
    match = re.match(r"^(A\d{2}:\d{4})\s*[-–—:]\s*(.+)$", raw, re.I)
    if match:
        return f"{match.group(1)} – {match.group(2).strip()}"
    return raw


def _cwe_display(finding: Finding) -> str:
    return (finding.cwe or "").strip() or "—"


def _severity_text(finding: Finding) -> str:
    """"HIGH · CVSS 3.1 7.5 (CVSS:3.1/AV:N/…) · CWE-89 baseline: …; was MEDIUM"."""
    if not finding.cvss_vector and not finding.severity_basis:
        return ""
    parts = [_severity_bucket(finding.severity).upper()]
    if finding.cvss is not None:
        parts.append(f"CVSS 3.1 {finding.cvss:.1f}" + (f" ({finding.cvss_vector})" if finding.cvss_vector else ""))
    basis = severity_basis_text(finding.severity_basis)
    if basis:
        parts.append(basis)
    return " · ".join(parts)


def _description_text(finding: Finding) -> str:
    return (finding.verdict_reason or _risk_text(finding) or finding.title or "").strip()


def _confidence_display(finding: Finding) -> str:
    value = finding.confidence
    if value is None:
        return ""
    try:
        pct = float(value)
    except (TypeError, ValueError):
        return ""
    if pct > 1.0:
        pct /= 100.0
    pct = max(0.0, min(1.0, pct))
    return f"{pct:.0%}"


def _taint_flow_text(finding: Finding) -> str:
    parts: list[str] = []
    source = (finding.source or "").strip()
    sink = (finding.sink or "").strip()
    flow = (finding.data_flow or "").strip()
    if source:
        parts.append(f"Source: {source}")
    if sink:
        parts.append(f"Sink: {sink}")
    if flow:
        parts.append(f"Flow: {flow}")
    evidence = _reachability_text(finding)
    if evidence:
        parts.append(evidence)
    return "\n".join(parts)


# Mirrors sast_engine pipeline/reachability.py _STATUS_RANK (the app can't
# import the engine).
_REACHABILITY_RANK = {"path": 0, "sanitized_path": 1, "entry_reachable": 2}


def _reachability_text(finding: Finding) -> str:
    """The Joern source->sink path attached to this finding, one step per line."""
    evidence = finding.reachability or {}
    steps = evidence.get("taint_path") or []
    if not steps:
        return ""
    status = evidence.get("status")
    if status == "entry_reachable":
        lines = ["Call path from an entry point (call graph, method level):"]
        lines += [
            f"  {i}. {s.get('method') or ''} -> {str(s.get('code') or '').strip()}  "
            f"({to_module_path(s.get('file')) or s.get('file')}:{s.get('line')})"
            for i, s in enumerate(steps, start=1)
        ]
    else:
        sanitized = status == "sanitized_path"
        lines = [f"Dataflow evidence (code property graph{', through a sanitizer' if sanitized else ''}):"]
        lines += [
            f"  {i}. {str(s.get('code') or '').strip()}  ({to_module_path(s.get('file')) or s.get('file')}:{s.get('line')})"
            for i, s in enumerate(steps, start=1)
        ]
    elided = (evidence.get("steps_total") or len(steps)) - len(steps)
    if elided > 0:
        lines.append(f"  ... {elided} intermediate step(s) omitted")
    return "\n".join(lines)


def _location_paths(finding: Finding) -> list[str]:
    if finding.affected_files:
        return [to_module_path(p) or p for p in finding.affected_files if p]
    path = to_module_path(finding.file_path) or ""
    if not path:
        return []
    if finding.line is not None:
        return [f"{path}:{finding.line}"]
    return [path]


def _primary_location(finding: Finding) -> str:
    paths = _location_paths(finding)
    if not paths:
        return "—"
    if len(paths) == 1:
        return paths[0]
    return f"{paths[0]}  (+{len(paths) - 1})"


def _truncate_text(text: str, limit: int) -> str:
    raw = (text or "").strip()
    if len(raw) <= limit:
        return raw
    return raw[: max(0, limit - 20)].rstrip() + "\n… [truncated]"


def _excel_scalar(value: Any) -> Any:
    """Coerce cell values for openpyxl — strings must not carry control chars."""
    if isinstance(value, str):
        return _sanitize_text(value)
    return value


def _excel_cell(text: Any) -> str:
    raw = _sanitize_text("" if text is None else str(text))
    if len(raw) <= _EXCEL_CELL_MAX:
        return raw
    return raw[: _EXCEL_CELL_MAX - 20].rstrip() + "\n… [truncated]"


def _location_paths_display(finding: Finding, *, max_paths: int = 8) -> str:
    paths = _location_paths(finding)
    if not paths:
        return ""
    if len(paths) <= max_paths:
        return "\n".join(paths)
    extra = len(paths) - max_paths
    return "\n".join(paths[:max_paths]) + f"\n… and {extra} more file(s)"


def _location_code(finding: Finding) -> str:
    return (finding.snippet or "").strip()


def _location_block(finding: Finding) -> str:
    paths = _location_paths(finding)
    parts: list[str] = []
    if paths:
        parts.append("\n".join(paths))
    code = _location_code(finding)
    if code:
        parts.append(code)
    return "\n\n".join(parts)


def _package_label(finding: Finding) -> str:
    """`name@version` from SCA matched_text, falling back to the title."""
    return (finding.matched_text or "").strip() or (finding.title or "").strip() or "—"


def _advisory_id(finding: Finding) -> str:
    return (finding.rule_id or "").strip() or "—"


def _advisory_summary(finding: Finding) -> str:
    """OSV advisory prose stored on `snippet` — not source code."""
    return (finding.snippet or "").strip()


def _sca_inventory(scan: Scan) -> dict[str, Any]:
    discovery = (getattr(scan, "metadata_json", None) or {}).get("discovery") or {}
    summary = discovery.get("dependency_summary") or {}
    manifests = summary.get("manifests_parsed") or []
    ecosystems = summary.get("ecosystems") or []
    return {
        "total_dependencies": int(summary.get("total_dependencies") or 0),
        "vulnerable_dependencies": int(summary.get("vulnerable_dependencies") or 0),
        "manifests_parsed": [str(m) for m in manifests if m],
        "ecosystems": [str(e) for e in ecosystems if e],
        "severity_counts": summary.get("severity_counts") or {},
    }


def _sbom_data(scan: Scan) -> dict[str, Any]:
    discovery = (getattr(scan, "metadata_json", None) or {}).get("discovery") or {}
    summary = discovery.get("component_summary") or {}
    components = discovery.get("components") or []
    license_summary = summary.get("license_summary") or {}
    return {
        "total_components": int(summary.get("total_components") or len(components)),
        "ecosystems": [str(e) for e in (summary.get("ecosystems") or []) if e],
        "manifests_parsed": [str(m) for m in (summary.get("manifests_parsed") or []) if m],
        "license_summary": {level: int(license_summary.get(level) or 0) for level in COST_ORDER},
        "components": components,
    }


def _component_label(comp: dict[str, Any]) -> str:
    group = (comp.get("group") or "").strip()
    name = (comp.get("name") or "").strip() or "—"
    return f"{group}:{name}" if group else name


def _component_license_label(comp: dict[str, Any]) -> str:
    licenses = comp.get("licenses") or []
    if licenses:
        first = licenses[0] or {}
        return first.get("spdx_id") or first.get("name") or "Unknown"
    return "Unknown"


def _component_type_label(comp: dict[str, Any]) -> str:
    """Normalises `origin_type` (deterministic fallback, refined by LLM
    enrichment — see `sbom_scan.py`) to one of the three canonical labels,
    defaulting anything unrecognized to "Open Source"."""
    raw = (comp.get("origin_type") or "").strip().lower()
    if raw == "proprietary":
        return "Proprietary"
    if raw == "internal":
        return "Internal"
    return "Open Source"


def _component_relationship(comp: dict[str, Any]) -> str:
    return (comp.get("dependency_relationship") or "Direct").strip() or "Direct"


def _component_supplier(comp: dict[str, Any]) -> str:
    return (comp.get("vendor") or "").strip() or "Unknown"


def _sorted_components(components: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(components, key=lambda c: _component_label(c).lower())


def _sorted_findings(findings: list[Finding]) -> list[Finding]:
    return sorted(
        findings,
        key=lambda f: (
            SEVERITY_ORDER.index(_severity_bucket(f.severity)),
            -(f.cvss or 0),
            f.title or "",
        ),
    )


def _normalized_match(finding: Finding) -> str:
    """Whitespace-collapsed matched text, so trivially different formatting of
    the same construct still groups together."""
    return " ".join((finding.matched_text or "").split())[:200]


def _finding_group_key(finding: Finding) -> tuple:
    """Same-issue key: a stable catalog rule_id groups by itself. The generic
    'llm-extra' fallback (or a missing rule_id) has no cross-file stability on
    its own, so those group by CWE + a normalized title topic instead.

    Regex findings additionally key on the *matched text*. Without it, one
    generated-code pattern repeated across hundreds of near-identical files
    would collapse together with genuinely different code that happens to trip
    the same rule — a hardcoded test placeholder and a real credential would
    become one row. With it, identical constructs collapse and distinct ones
    stay visible. LLM findings keep the old rule-only key: their `matched_text`
    is empty, so including it would be a no-op anyway.

    A rule can opt out of the matched-text split via `group_by_rule` (denorm'd
    onto the Finding row from RegexRule.group_by_rule at write time) for
    constructs whose matched text is expected to differ every occurrence — a
    class name, say — while still describing exactly one issue.

    Every rule/topic key ends with the occurrence's semantic signature (what it
    calls, on which receiver type, with which kinds of arguments): members that
    differ there are different constructs and never share a verdict. Rows
    without facts all carry `unresolved`, so old scans group exactly as before."""
    verdict = (finding.verdict or "").strip().lower()
    rule_id = (finding.rule_id or "").strip()
    if (finding.kind or "").strip().lower() == "dast":
        # One issue = one check on one endpoint template and parameter, however many concrete
        # URLs showed it (the DAST writer already dedupes; this is the read-side guarantee).
        from app.dast.domain.dedupe import target_path_template

        # An authorization finding is also one issue per role pair (each pair is a matrix cell).
        return (verdict, "dast", rule_id, target_path_template(finding.target or finding.file_path or ""), finding.param or "", ">".join(finding.roles or ()))
    if (finding.kind or "").strip().lower() == "flow":
        # A flow's identity is its source→sink chain. Two paths reaching the
        # same kind of sink through different files are different
        # vulnerabilities, but they share a CWE and title and carry no matched
        # text, so every other key here collapses them into one row and hides
        # all but one. Hop line numbers are excluded so an edit that shifts a
        # method down the file doesn't split an existing group.
        hops = tuple(
            (str(hop.get("file") or ""), str(hop.get("method") or ""))
            for hop in (finding.flow_path or [])
        )
        if hops:
            return (verdict, "flow", (finding.cwe or "").strip().upper(), hops)
    signature = semantic_signature(finding.facts)
    if rule_id and rule_id.lower() != "llm-extra":
        if finding.group_by_rule:
            return (verdict, "rule", rule_id, signature)
        return (verdict, "rule", rule_id, _normalized_match(finding), signature)
    topic = re.sub(r"[^a-z0-9]+", " ", (finding.title or "").lower()).strip()
    topic = " ".join(topic.split()[:3])
    return (verdict, "topic", (finding.cwe or "").strip().upper(), topic,
            _normalized_match(finding), signature)


# Kinds whose title/rule_id are model-written and so vary between findings that
# describe one root cause. Regex rows are excluded on purpose (their matched
# text split is a design decision), as are dast/flow (structural identity).
_CONSOLIDATED_KINDS = frozenset({"llm", "joern"})
_MAX_AFFECTED_LINES = 200
_IDENT_TOKEN_RE = re.compile(r"[a-z][A-Z]|[_.()/\\]")
_TOPIC_PREPOSITIONS = frozenset({"in", "at", "within", "inside", "of", "on", "for", "from", "via", "by"})


def _cwe_number(finding: Finding) -> int | None:
    """The CWE as a number, whatever the spelling ('CWE-89', 'cwe 089', '89')."""
    match = re.search(r"\d+", finding.cwe or "")
    return int(match.group()) if match else None


def _issue_topic(finding: Finding) -> str:
    """Title with identifier tokens (method/class/variable names) and the
    preposition leading into them removed, so 'X in getActionCode' and
    'X in getPortMsgType' share a topic. Empty when too little is left."""
    tokens = (finding.title or "").split()
    kept: list[str] = []
    for token in tokens:
        if _IDENT_TOKEN_RE.search(token):
            while kept and kept[-1].lower() in _TOPIC_PREPOSITIONS:
                kept.pop()
            continue
        kept.append(token)
    topic = re.sub(r"[^a-z0-9]+", " ", " ".join(kept).lower()).split()
    return " ".join(topic) if len(topic) >= 2 else ""


def _consolidate_groups(groups: list[list[Finding]]) -> list[list[Finding]]:
    """Merge key-groups that describe one root cause, whatever rule_id/title the
    model gave each. Two scopes, both requiring the same verdict and CWE number:
    a single-file group merges with every other single-file group in that file;
    any group merges with groups sharing its identifier-free title topic."""
    parent = list(range(len(groups)))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    seen: dict[tuple, int] = {}
    for i, members in enumerate(groups):
        if not all((m.kind or "").strip().lower() in _CONSOLIDATED_KINDS for m in members):
            continue
        rep = members[0]
        cwe = _cwe_number(rep)
        if cwe is None:
            continue
        verdict = (rep.verdict or "").strip().lower()
        keys: list[tuple] = []
        files = {m.file_path or "" for m in members}
        if len(files) == 1:
            keys.append((verdict, cwe, "file", next(iter(files))))
        topic = _issue_topic(rep)
        if topic:
            keys.append((verdict, cwe, "topic", topic))
        for key in keys:
            if key in seen:
                parent[find(i)] = find(seen[key])
            else:
                seen[key] = i

    merged: dict[int, list[Finding]] = {}
    for i, members in enumerate(groups):
        merged.setdefault(find(i), []).extend(members)
    return list(merged.values())


def issue_groups(findings: list[Finding]) -> list[list[Finding]]:
    """The member lists `group_findings_by_issue` collapses `findings` into
    (key-groups, then consolidated). Reviews use this too, so a decision on a
    merged row covers every row it stands for."""
    groups: dict[tuple, list[Finding]] = {}
    for f in findings:
        groups.setdefault(_finding_group_key(f), []).append(f)
    return _consolidate_groups(list(groups.values()))


def construct_key(finding: Finding) -> str:
    """Identity of the *construct* a row is about — its group key without the
    leading verdict — as a short stable string. Rows that differ only in verdict
    share it: that is what "one construct, mixed outcomes" means."""
    return hashlib.sha1(repr(_finding_group_key(finding)[1:]).encode("utf-8")).hexdigest()[:12]


_MIX_BUCKETS = (("tp", "confirmed"), ("fp", "false positive"), ("unresolved", "unresolved"))


def _mix_bucket(finding: Finding) -> str:
    verdict = (finding.verdict or "").strip().lower()
    if verdict == "true_positive":
        return "tp"
    if verdict == "false_positive":
        return "fp"
    return "unresolved"


def construct_mix_text(mix: dict[str, int] | None) -> str:
    """"Mixed: 3 confirmed · 150 false positive · 3 unresolved" for a construct
    whose members did not all get the same verdict; empty for a uniform one."""
    if not mix or sum(1 for key, _ in _MIX_BUCKETS if mix.get(key)) < 2:
        return ""
    return "Mixed: " + " · ".join(f"{mix.get(key, 0)} {label}" for key, label in _MIX_BUCKETS if mix.get(key))


def review_coverage_text(finding: Finding) -> str:
    """"3 / 156" — occurrences the confirmation reviewed over the group's size."""
    coverage = getattr(finding, "review_coverage", None)
    if not coverage:
        return ""
    return f"{coverage.get('reviewed', 0)} / {coverage.get('total', 0)}"


def annotate_construct_mix(grouped: list[Finding]) -> list[Finding]:
    """Set `construct_key` and `construct_mix` ({tp, fp, unresolved} in MEMBERS,
    counted across every row of the construct) on each grouped row. Computed from
    the grouped list itself, so the API, both exports and the consistency check
    all read one answer. Like `occurrence_count` this is set on the ORM instance
    and never flushed."""
    mixes: dict[str, dict[str, int]] = {}
    keys: dict[int, str] = {}
    for finding in grouped:
        key = construct_key(finding)
        keys[id(finding)] = key
        mix = mixes.setdefault(key, {"tp": 0, "fp": 0, "unresolved": 0})
        mix[_mix_bucket(finding)] += _occurrences(finding)
    for finding in grouped:
        key = keys[id(finding)]
        finding.construct_key = key
        finding.construct_mix = dict(mixes[key])
    return grouped


def _occurrence_ids(finding: Finding) -> list[str]:
    """Occurrence ids a (possibly grouped) finding stands for, in member order."""
    members = getattr(finding, "member_occurrence_ids", None)
    if members:
        return list(members)
    return [finding.occurrence_id] if getattr(finding, "occurrence_id", None) else []


def group_findings_by_issue(findings: list[Finding]) -> list[Finding]:
    """Collapse findings that represent the same issue across multiple files
    into one representative Finding per group, with every affected file path
    listed in `affected_files` — so callers show one card/row per issue type
    instead of one per file occurrence.

    The representative is the highest-severity (then highest-cvss) member of
    the group, so its title/snippet/ai_fix/verdict_reason serve as the one
    generic example for the whole group. Singleton groups pass through with
    `affected_files` left unset, unchanged from today's behavior.

    Every returned finding carries `occurrence_count` (how many rows it
    stands for). Like `affected_files`, it is set on the ORM instance and
    must never be flushed — nothing commits after this runs."""
    result: list[Finding] = []
    for members in issue_groups(findings):
        reasons: dict[str, int] = {}
        for m in members:
            if _verdict_bucket(m) == UNVERIFIED_VERDICT:
                reason = m.unresolved_reason or UNKNOWN_BUCKET
                reasons[reason] = reasons.get(reason, 0) + 1
        if len(members) == 1:
            members[0].unresolved_reason_counts = reasons
            members[0].occurrence_count = 1
            members[0].affected_lines = None
            members[0].member_occurrence_ids = (
                [members[0].occurrence_id] if members[0].occurrence_id else None
            )
            result.append(members[0])
            continue
        # Among equals, prefer the member with the strongest reachability
        # evidence so the row shows it.
        representative = min(
            members,
            key=lambda f: (
                SEVERITY_ORDER.index(_severity_bucket(f.severity)),
                -(f.cvss or 0),
                _REACHABILITY_RANK.get((f.reachability or {}).get("status"), 9),
            ),
        )
        paths: list[str] = []
        seen: set[str] = set()
        for m in members:
            candidates = [m.file_path or ""]
            if (m.kind or "").lower() == "dast":
                candidates += list(m.affected_files or [])
            for candidate in candidates:
                p = candidate.strip()
                if p and p not in seen:
                    seen.add(p)
                    paths.append(p)
        locations = sorted(
            {((m.file_path or ""), int(m.line or 0)) for m in members if m.line}
        )[:_MAX_AFFECTED_LINES]
        representative.affected_files = paths
        representative.affected_lines = [{"file_path": p, "line": n} for p, n in locations] or None
        representative.occurrence_count = len(members)
        representative.unresolved_reason_counts = reasons
        member_ids: list[str] = []
        for m in members:
            if m.occurrence_id and m.occurrence_id not in member_ids:
                member_ids.append(m.occurrence_id)
        representative.member_occurrence_ids = member_ids or None
        result.append(representative)
    return _sorted_findings(annotate_construct_mix(result))


def _cwe_group_key(finding: Finding) -> str:
    """The CWE a grouped finding rolls up under ("CWE-89", whatever the spelling).
    One without a CWE stays with its own rule (or, for a model-written finding,
    its title topic), so unrelated issues never share a catch-all row."""
    number = _cwe_number(finding)
    if number is not None:
        return f"CWE-{number}"
    rule_id = (finding.rule_id or "").strip()
    if rule_id and rule_id.lower() != "llm-extra":
        return f"rule:{rule_id}"
    return "topic:" + (_issue_topic(finding) or " ".join((finding.title or "").lower().split()))


def _group_paths(finding: Finding) -> list[str]:
    return [p for p in (finding.affected_files or [finding.file_path or ""]) if p]


def group_by_cwe(grouped: list[Finding]) -> list[dict[str, Any]]:
    """Roll the issue groups (`group_findings_by_issue` output) up to one row per
    CWE: every file it touches, its occurrences, and per verdict bucket how many
    groups, occurrences and files fall there plus the ids of those groups. A
    read-side view only — each group, and each occurrence in it, keeps its own
    verdict; nothing here decides or copies one. Ordered by highest severity, then
    confirmed occurrences, then size."""
    rows: dict[str, dict[str, Any]] = {}
    for finding in grouped:
        key = _cwe_group_key(finding)
        occurrences = _occurrences(finding)
        paths = _group_paths(finding)
        row = rows.get(key)
        if row is None:
            row = rows[key] = {
                "key": key, "cwe": key if key.startswith("CWE-") else None,
                "title": finding.title or "", "owasp_category": finding.owasp_category,
                "severity": _severity_bucket(finding.severity), "groups": 0, "occurrences": 0,
                "files": [], "by_bucket": {}, "_seen": set(), "_lead": -1,
            }
        if occurrences > row["_lead"]:  # the largest group names the row
            row["_lead"] = occurrences
            row["title"] = finding.title or row["title"]
            row["owasp_category"] = finding.owasp_category or row["owasp_category"]
        severity = _severity_bucket(finding.severity)
        if SEVERITY_ORDER.index(severity) < SEVERITY_ORDER.index(row["severity"]):
            row["severity"] = severity
        row["groups"] += 1
        row["occurrences"] += occurrences
        for path in paths:
            if path not in row["_seen"]:
                row["_seen"].add(path)
                row["files"].append(path)
        cell = row["by_bucket"].setdefault(
            _verdict_bucket(finding),
            {"groups": 0, "occurrences": 0, "files": 0, "finding_ids": [], "_seen": set()},
        )
        cell["groups"] += 1
        cell["occurrences"] += occurrences
        cell["finding_ids"].append(finding.id)
        cell["_seen"].update(paths)

    result = []
    for row in rows.values():
        del row["_seen"], row["_lead"]
        for cell in row["by_bucket"].values():
            cell["files"] = len(cell.pop("_seen"))
        row["file_count"] = len(row["files"])
        result.append(row)
    result.sort(key=lambda r: (
        SEVERITY_ORDER.index(r["severity"]),
        -r["by_bucket"].get("true_positive", {}).get("occurrences", 0),
        -r["occurrences"],
        r["key"],
    ))
    return result


_REPORT_BUCKETS = (*VERDICT_BUCKETS, UNKNOWN_BUCKET)


def _report_order(findings: list[Finding], cwe_rows: list[dict[str, Any]]) -> list[Finding]:
    """The PDF's order (its index and its finding blocks share the numbering):
    verdict bucket (Confirmed, Unresolved, False Positive), then CWE in
    `group_by_cwe` order, then the usual severity order within a CWE."""
    cwe_rank = {row["key"]: i for i, row in enumerate(cwe_rows)}
    base = {id(f): i for i, f in enumerate(_sorted_findings(findings))}
    return sorted(findings, key=lambda f: (
        _REPORT_BUCKETS.index(_verdict_bucket(f)), cwe_rank.get(_cwe_group_key(f), len(cwe_rank)), base[id(f)],
    ))


# =============================================================================
# Excel
# =============================================================================

_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _xl_header_row(sheet: Any, row: int, headers: list[str], *, widths: list[int] | None = None) -> None:
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col, label in enumerate(headers, start=1):
        cell = sheet.cell(row, col, label)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor="152238")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN
    if widths:
        for idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = width


def _build_summary_sheet(sheet: Any, project: Project, scan: Scan, gate: dict, counts: dict[str, Any]) -> None:
    """Row positions are computed as we go. The previous version hardcoded
    `A14` for the severity header while the metadata block above it had a
    conditional extra row — one more metadata field would have overwritten it.

    `counts` is `count_findings` over the grouped findings: the headline counts
    are actionable groups, with each verdict bucket and issue class alongside."""
    actionable, dismissed = counts["actionable"], counts["dismissed"]
    title_font = Font(name="Calibri", size=18, bold=True, color="152238")
    section_font = Font(name="Calibri", size=13, bold=True, color="152238")
    label_font = Font(name="Calibri", size=11, color="64748B")
    value_font = Font(name="Calibri", size=11, color="152238")

    sheet["A1"] = "Secure Build — SAST Scan Report"
    sheet["A1"].font = title_font
    sheet.merge_cells("A1:B1")

    gate_failed = bool(gate.get("gate_failed"))
    sheet["A2"] = "QUALITY GATE FAILED" if gate_failed else "QUALITY GATE PASSED"
    sheet["A2"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    sheet["A2"].fill = PatternFill("solid", fgColor="B91C1C" if gate_failed else "047857")
    sheet["A2"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:B2")

    meta_rows: list[tuple[str, Any]] = [
        ("Project", project.name),
        ("Scan version", f"v{scan.version}"),
        ("Status", scan.status),
        ("Branch", scan.branch),
        ("Commit", (scan.commit_sha or "—")[:12]),
        ("Git URL", redact_git_url(scan.git_url) or "—"),
        ("Created", _fmt_dt(scan.created_at)),
        (
            "LLM",
            "None"
            if (scan.metadata_json or {}).get("use_llm") is False
            or (not scan.llm_provider and not scan.llm_model)
            else _llm_display(scan),
        ),
        ("Actionable issues (groups)", sum(v["groups"] for v in actionable.values())),
        ("Occurrences", sum(v["occurrences"] for v in actionable.values())),
    ]
    status_note = _scan_status_banner(scan)
    if status_note:
        meta_rows.insert(3, ("Note", status_note))

    row = 4
    for label, value in meta_rows:
        sheet.cell(row, 1, label).font = label_font
        sheet.cell(row, 2, _excel_scalar(value)).font = value_font
        row += 1
    # Each verdict bucket on its own row: groups, then occurrences.
    for bucket, cell in verdict_bucket_totals(counts).items():
        sheet.cell(row, 1, f"{BUCKET_ROW_PREFIX}{bucket.lower()}").font = label_font
        sheet.cell(row, 2, cell["groups"]).font = value_font
        sheet.cell(row, 3, cell["occurrences"]).font = value_font
        row += 1

    row += 1
    sheet.cell(row, 1, "Severity breakdown").font = section_font
    row += 1
    _xl_header_row(sheet, row, ["Severity", "Groups", "Occurrences"])
    row += 1
    for sev in SEVERITY_ORDER:
        cell = sheet.cell(row, 1, sev.upper())
        cell.fill = EXCEL_SEVERITY_FILLS[sev]
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.border = _THIN
        for col, value in ((2, actionable[sev]["groups"]), (3, actionable[sev]["occurrences"])):
            count_cell = sheet.cell(row, col, value)
            count_cell.alignment = Alignment(horizontal="center")
            count_cell.border = _THIN
        row += 1
    sheet.cell(row, 1, "False positive").border = _THIN
    for col, value in ((2, dismissed["groups"]), (3, dismissed["occurrences"])):
        count_cell = sheet.cell(row, col, value)
        count_cell.alignment = Alignment(horizontal="center")
        count_cell.border = _THIN
    row += 1

    row += 1
    sheet.cell(row, 1, "Issue class").font = section_font
    row += 1
    _xl_header_row(sheet, row, ["Issue class", "Groups", "Occurrences"])
    row += 1
    for name, cell in counts["by_issue_class"].items():
        sheet.cell(row, 1, name).border = _THIN
        for col, value in ((2, cell["groups"]), (3, cell["occurrences"])):
            count_cell = sheet.cell(row, col, value)
            count_cell.alignment = Alignment(horizontal="center")
            count_cell.border = _THIN
        row += 1

    row += 1
    sheet.cell(row, 1, "OWASP distribution").font = section_font
    row += 1
    _xl_header_row(sheet, row, ["Category", "Groups", "Critical + High", "Occurrences"])
    row += 1
    for entry in counts["by_owasp"]:
        label, total, severe, occurrences = (
            entry["category"], entry["groups"], entry["severe"], entry["occurrences"]
        )
        sheet.cell(row, 1, label).border = _THIN
        for col, value in ((2, total), (3, severe), (4, occurrences)):
            cell = sheet.cell(row, col, value)
            cell.alignment = Alignment(horizontal="center")
            cell.border = _THIN
        row += 1

    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 56
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 14


def _build_findings_sheet(sheet: Any, findings: list[Finding]) -> None:
    headers = [
        "#",
        "Severity",
        "Scanner Verdict",
        "Review State",
        "Issue Class",
        "Unresolved Reason",
        "Verdict (raw)",
        "Counted As",
        "OWASP Category",
        "CWE",
        "CVSS",
        "CVSS Vector",
        "Severity Basis",
        "Confidence",
        "Files",
        "Issue",
        "Location(s)",
        "Data Flow",
        "Risk",
        "Recommendation",
        "Example Fix",
        "Occurrence ID",
        "Construct Mix",
        "Reviewed / Total",
    ]
    widths = [5, 12, 24, 22, 12, 24, 16, 14, 32, 12, 8, 44, 48, 12, 8, 42, 44, 36, 44, 44, 46, 68, 40, 16]
    _xl_header_row(sheet, 1, headers, widths=widths)

    sorted_findings = _sorted_findings(findings)
    wrap_from = headers.index("Issue") + 1

    for idx, finding in enumerate(sorted_findings, start=1):
        sev = _severity_bucket(finding.severity)
        recommendation = _recommendation_text(finding)
        if not recommendation and finding.ai_fix and not _example_fix_text(finding):
            recommendation = (finding.ai_fix or "").strip()
        row_values = [
            idx,
            sev.upper(),
            _display_verdict(_scanner_verdict(finding), finding.unresolved_reason),
            _review_state_label(finding),
            finding.issue_class or "",
            finding.unresolved_reason or "",
            _scanner_verdict(finding) or "",
            _counted_as(finding),
            _owasp_display(finding),
            _cwe_display(finding),
            finding.cvss if finding.cvss is not None else "",
            finding.cvss_vector or "",
            _excel_cell(severity_basis_text(finding.severity_basis)),
            _confidence_display(finding),
            _affected_file_count(finding),
            _excel_cell(finding.title or ""),
            _excel_cell(_location_block(finding)),
            _excel_cell(_taint_flow_text(finding)),
            _excel_cell(_risk_text(finding)),
            _excel_cell(recommendation),
            _excel_cell(_example_fix_text(finding)),
            "\n".join(_occurrence_ids(finding)),
            construct_mix_text(getattr(finding, "construct_mix", None)),
            review_coverage_text(finding),
        ]
        sheet.append([_excel_scalar(v) for v in row_values])

        excel_row = idx + 1
        sev_cell = sheet.cell(excel_row, 2)
        sev_cell.fill = EXCEL_SEVERITY_FILLS[sev]
        sev_cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        sev_cell.alignment = Alignment(horizontal="center", vertical="top")

        for col in range(1, len(headers) + 1):
            cell = sheet.cell(excel_row, col)
            cell.border = _THIN
            cell.alignment = Alignment(vertical="top", wrap_text=col >= wrap_from)
            if excel_row % 2 == 0 and col != 2:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")

    sheet.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_col}{max(1, len(sorted_findings) + 1)}"


_CWE_SHEET_MAX_PATHS = 50


def _bucket_cell(row: dict[str, Any], bucket: str) -> str:
    cell = row["by_bucket"].get(bucket)
    if not cell:
        return ""
    return f"{cell['groups']} ({cell['occurrences']})"


def _build_cwe_sheet(sheet: Any, cwe_rows: list[dict[str, Any]]) -> None:
    """One row per CWE (`group_by_cwe`): what it is, how bad, where, and how its
    groups split by verdict — "groups (occurrences)" per bucket."""
    headers = [
        "CWE", "Title", "OWASP Category", "Highest Severity", "Confirmed", "Unresolved",
        "False Positive", "Groups", "Occurrences", "Files", "Affected Files",
    ]
    widths = [12, 48, 32, 14, 14, 14, 14, 10, 12, 8, 80]
    _xl_header_row(sheet, 1, headers, widths=widths)
    for idx, row in enumerate(cwe_rows, start=2):
        paths = [to_module_path(p) or p for p in row["files"]]
        shown = "\n".join(paths[:_CWE_SHEET_MAX_PATHS])
        if len(paths) > _CWE_SHEET_MAX_PATHS:
            shown += f"\n… and {len(paths) - _CWE_SHEET_MAX_PATHS} more file(s)"
        sev = row["severity"]
        values = [
            row["cwe"] or "—", _excel_cell(row["title"]), row["owasp_category"] or "",
            sev.upper(), _bucket_cell(row, "true_positive"), _bucket_cell(row, UNVERIFIED_VERDICT),
            _bucket_cell(row, "false_positive"), row["groups"], row["occurrences"], row["file_count"],
            _excel_cell(shown),
        ]
        sheet.append([_excel_scalar(v) for v in values])
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(idx, col)
            cell.border = _THIN
            cell.alignment = Alignment(vertical="top", wrap_text=col in (2, len(headers)))
        sev_cell = sheet.cell(idx, 4)
        sev_cell.fill = EXCEL_SEVERITY_FILLS[sev]
        sev_cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(cwe_rows) + 1)}"


def _build_sca_summary_sheet(sheet: Any, project: Project, scan: Scan, findings: list[Finding]) -> None:
    inventory = _sca_inventory(scan)
    counts = _count_by_severity(findings)
    title_font = Font(name="Calibri", size=18, bold=True, color="152238")
    section_font = Font(name="Calibri", size=13, bold=True, color="152238")
    label_font = Font(name="Calibri", size=11, color="64748B")
    value_font = Font(name="Calibri", size=11, color="152238")

    sheet["A1"] = "Secure Build — Dependency Vulnerability (SCA) Report"
    sheet["A1"].font = title_font
    sheet.merge_cells("A1:B1")

    meta_rows: list[tuple[str, Any]] = [
        ("Project", project.name),
        ("Scan version", f"v{scan.version}"),
        ("Status", scan.status),
        ("Branch", scan.branch),
        ("Commit", (scan.commit_sha or "—")[:12]),
        ("Git URL", redact_git_url(scan.git_url) or "—"),
        ("Created", _fmt_dt(scan.created_at)),
        ("Lookup source", "OSV.dev"),
        ("Dependencies", inventory["total_dependencies"]),
        ("Vulnerable packages", inventory["vulnerable_dependencies"]),
        ("Manifests", ", ".join(inventory["manifests_parsed"]) or "—"),
        ("Ecosystems", ", ".join(inventory["ecosystems"]) or "—"),
        ("Advisories", len(findings)),
    ]
    status_note = _scan_status_banner(scan)
    if status_note:
        meta_rows.insert(3, ("Note", status_note))

    row = 3
    for label, value in meta_rows:
        sheet.cell(row, 1, label).font = label_font
        sheet.cell(row, 2, _excel_scalar(value)).font = value_font
        row += 1

    row += 1
    sheet.cell(row, 1, "Severity breakdown").font = section_font
    row += 1
    _xl_header_row(sheet, row, ["Severity", "Count"])
    row += 1
    for sev in SEVERITY_ORDER:
        cell = sheet.cell(row, 1, sev.upper())
        cell.fill = EXCEL_SEVERITY_FILLS[sev]
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.border = _THIN
        count_cell = sheet.cell(row, 2, counts.get(sev, 0))
        count_cell.alignment = Alignment(horizontal="center")
        count_cell.border = _THIN
        row += 1

    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 72


def _build_sca_findings_sheet(sheet: Any, findings: list[Finding]) -> None:
    headers = [
        "#",
        "Severity",
        "Package",
        "Advisory",
        "CVSS",
        "Manifest",
        "Summary",
        "Suggested fix",
    ]
    widths = [5, 12, 28, 22, 8, 36, 52, 44]
    _xl_header_row(sheet, 1, headers, widths=widths)

    sorted_findings = _sorted_findings(findings)
    wrap_from = headers.index("Package") + 1

    for idx, finding in enumerate(sorted_findings, start=1):
        sev = _severity_bucket(finding.severity)
        row_values = [
            idx,
            sev.upper(),
            _excel_cell(_package_label(finding)),
            _excel_cell(_advisory_id(finding)),
            finding.cvss if finding.cvss is not None else "",
            _excel_cell(_location_paths_display(finding) or finding.file_path or "—"),
            _excel_cell(_advisory_summary(finding)),
            _excel_cell((finding.ai_fix or "").strip()),
        ]
        sheet.append([_excel_scalar(v) for v in row_values])

        excel_row = idx + 1
        sev_cell = sheet.cell(excel_row, 2)
        sev_cell.fill = EXCEL_SEVERITY_FILLS[sev]
        sev_cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        sev_cell.alignment = Alignment(horizontal="center", vertical="top")

        for col in range(1, len(headers) + 1):
            cell = sheet.cell(excel_row, col)
            cell.border = _THIN
            cell.alignment = Alignment(vertical="top", wrap_text=col >= wrap_from)
            if excel_row % 2 == 0 and col != 2:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")

    sheet.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_col}{max(1, len(sorted_findings) + 1)}"


def _build_sbom_summary_sheet(sheet: Any, project: Project, scan: Scan, data: dict[str, Any]) -> None:
    title_font = Font(name="Calibri", size=18, bold=True, color="152238")
    section_font = Font(name="Calibri", size=13, bold=True, color="152238")
    label_font = Font(name="Calibri", size=11, color="64748B")
    value_font = Font(name="Calibri", size=11, color="152238")

    sheet["A1"] = "Secure Build — Software Bill of Materials (SBOM) Report"
    sheet["A1"].font = title_font
    sheet.merge_cells("A1:B1")

    meta_rows: list[tuple[str, Any]] = [
        ("Project", project.name),
        ("Scan version", f"v{scan.version}"),
        ("Status", scan.status),
        ("Branch", scan.branch),
        ("Commit", (scan.commit_sha or "—")[:12]),
        ("Git URL", redact_git_url(scan.git_url) or "—"),
        ("Created", _fmt_dt(scan.created_at)),
        ("Lookup source", "deps.dev"),
        ("Total components", data["total_components"]),
        ("Manifests", ", ".join(data["manifests_parsed"]) or "—"),
        ("Ecosystems", ", ".join(data["ecosystems"]) or "—"),
    ]
    status_note = _scan_status_banner(scan)
    if status_note:
        meta_rows.insert(3, ("Note", status_note))

    row = 3
    for label, value in meta_rows:
        sheet.cell(row, 1, label).font = label_font
        sheet.cell(row, 2, _excel_scalar(value)).font = value_font
        row += 1

    row += 1
    sheet.cell(row, 1, "License / cost breakdown").font = section_font
    row += 1
    _xl_header_row(sheet, row, ["Classification", "Count"])
    row += 1
    for level in COST_ORDER:
        cell = sheet.cell(row, 1, COST_LABELS[level])
        cell.fill = EXCEL_COST_FILLS[level]
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.border = _THIN
        count_cell = sheet.cell(row, 2, data["license_summary"].get(level, 0))
        count_cell.alignment = Alignment(horizontal="center")
        count_cell.border = _THIN
        row += 1

    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 56


def _build_sbom_components_sheet(sheet: Any, components: list[dict[str, Any]]) -> None:
    headers = [
        "#",
        "Component Name",
        "Version",
        "Supplier",
        "Type",
        "License",
        "Dependency Relationship",
        "Unique Identifier (PURL/CPE)",
    ]
    widths = [5, 32, 16, 26, 14, 20, 20, 52]
    _xl_header_row(sheet, 1, headers, widths=widths)

    ordered = _sorted_components(components)
    wrap_from = headers.index("Unique Identifier (PURL/CPE)") + 1
    type_col = headers.index("Type") + 1

    for idx, comp in enumerate(ordered, start=1):
        type_label = _component_type_label(comp)
        row_values = [
            idx,
            _excel_cell(_component_label(comp)),
            _excel_cell(comp.get("version") or "—"),
            _excel_cell(_component_supplier(comp)),
            type_label,
            _excel_cell(_component_license_label(comp)),
            _excel_cell(_component_relationship(comp)),
            _excel_cell(comp.get("purl") or "—"),
        ]
        sheet.append([_excel_scalar(v) for v in row_values])

        excel_row = idx + 1
        type_cell = sheet.cell(excel_row, type_col)
        type_cell.fill = EXCEL_TYPE_FILLS[type_label]
        type_cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        type_cell.alignment = Alignment(horizontal="center", vertical="top")

        for col in range(1, len(headers) + 1):
            cell = sheet.cell(excel_row, col)
            cell.border = _THIN
            cell.alignment = Alignment(vertical="top", wrap_text=col >= wrap_from)
            if excel_row % 2 == 0 and col != type_col:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")

    sheet.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last_col}{max(1, len(ordered) + 1)}"


def _build_excel_sbom_report(project: Project, scan: Scan) -> bytes:
    data = _sbom_data(scan)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _build_sbom_summary_sheet(summary, project, scan, data)
    _build_sbom_components_sheet(wb.create_sheet("Components"), data["components"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_excel_sca_report(project: Project, scan: Scan, findings: list[Finding]) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _build_sca_summary_sheet(summary, project, scan, findings)
    _build_sca_findings_sheet(wb.create_sheet("Findings"), findings)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_excel_report(project: Project, scan: Scan, findings: list[Finding]) -> bytes:
    if _is_dast_scan(scan):
        from app.services.dast_export import build_dast_xlsx  # lazy: dast_export builds on this module

        return build_dast_xlsx(project, scan, findings)
    if _is_sca_scan(scan):
        return _build_excel_sca_report(project, scan, findings)
    if _is_sbom_scan(scan):
        return _build_excel_sbom_report(project, scan)

    counts = count_findings(findings)
    gate = build_gate_summary(project, findings, counts=counts)
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    _build_summary_sheet(summary, project, scan, gate, counts)

    _build_cwe_sheet(wb.create_sheet("By CWE"), group_by_cwe(findings))
    _build_findings_sheet(wb.create_sheet("Findings"), findings)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# PDF — building blocks
# =============================================================================

class _Bookmark(Flowable):
    """Zero-height flowable that registers a PDF outline entry, so the reader's
    sidebar becomes a working table of contents."""

    width = 0
    height = 0

    def __init__(self, title: str, key: str, level: int = 0):
        super().__init__()
        self.title = title
        self.key = key
        self.level = level

    def draw(self) -> None:
        self.canv.bookmarkPage(self.key)
        self.canv.addOutlineEntry(self.title[:120], self.key, level=self.level, closed=False)

    def wrap(self, *_args) -> tuple[float, float]:
        return (0, 0)


def _code_wrap_chars(width_pt: float, font_size: float = _CODE_FONT_SIZE) -> int:
    """Characters per line from actual Courier metrics rather than a fixed 88 —
    landscape gives roughly 140, so real code stops being mangled."""
    return max(40, int(width_pt / (font_size * _COURIER_ADVANCE)) - 1)


def _pill(text: str, bg: colors.Color, styles: Any, *, width: float) -> Table:
    pill = Table(
        [[Paragraph(f"<b>{_esc(text)}</b>", styles["BadgeText"])]],
        colWidths=[width],
    )
    pill.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), bg),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("ROUNDEDCORNERS", [3, 3, 3, 3]),
            ]
        )
    )
    return pill


def _labelled_block(
    label: str,
    text: str,
    *,
    styles: Any,
    width: float,
    label_color: str,
    accent: str | None = None,
) -> Table:
    """A labelled prose block. Deliberately borderless — it lives inside the
    finding card, and nesting a bordered box inside a bordered card is what
    made the old output look like scattered fragments."""
    body = _truncate_text(text or "", _PDF_SECTION_MAX)
    block = Table(
        [
            [Paragraph(f'<font color="{label_color}"><b>{_esc(label.upper())}</b></font>', styles["FindingLabel"])],
            [Paragraph(_esc(body).replace("\n", "<br/>"), styles["FindingBody"])],
        ],
        colWidths=[width],
        splitInRow=1,
    )
    cmds: list[tuple] = [
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (0, 0), 0),
        ("BOTTOMPADDING", (0, 0), (0, 0), 1.2 * mm),
        ("TOPPADDING", (0, 1), (0, 1), 0),
        ("BOTTOMPADDING", (0, 1), (0, 1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    if accent:
        cmds.extend([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(accent)),
            ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
            ("TOPPADDING", (0, 0), (0, 0), 2 * mm),
            ("BOTTOMPADDING", (0, 1), (0, 1), 2 * mm),
        ])
    block.setStyle(TableStyle(cmds))
    return block


def _code_box(label: str, snippet: str, *, styles: Any, width: float) -> Table:
    wrap_at = _code_wrap_chars(width - 6 * mm)
    lines = (snippet or "").replace("\t", "  ").splitlines() or [""]
    wrapped: list[str] = []
    for line in lines:
        while len(line) > wrap_at:
            wrapped.append(_esc(line[:wrap_at]))
            line = line[wrap_at:]
        wrapped.append(_esc(line))
    code_html = "<br/>".join(wrapped)
    box = Table(
        [
            [Paragraph(f'<font color="#334155"><b>{_esc(label.upper())}</b></font>', styles["FindingLabel"])],
            [Paragraph(
                f'<font face="Courier" size="{_CODE_FONT_SIZE}" color="#0f172a">{code_html}</font>',
                styles["FindingCode"],
            )],
        ],
        colWidths=[width],
        # Without this an oversized snippet raises LayoutError and takes the
        # whole report build down; with it the row flows to the next page.
        splitInRow=1,
    )
    box.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#e2e8f0")),
                ("BACKGROUND", (0, 1), (0, 1), colors.HexColor("#f8fafc")),
                ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#cbd5e1")),
                ("LEFTPADDING", (0, 0), (-1, -1), 3 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
                ("TOPPADDING", (0, 0), (0, 0), 1.8 * mm),
                ("BOTTOMPADDING", (0, 0), (0, 0), 1 * mm),
                ("TOPPADDING", (0, 1), (0, 1), 1.5 * mm),
                ("BOTTOMPADDING", (0, 1), (0, 1), 2.5 * mm),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return box


def _code_boxes(label: str, snippet: str, *, styles: Any, width: float) -> list[Any]:
    """Split long code into page-safe chunks. Chunk labels carry (n/total) so a
    split reads as continuation rather than as separate evidence."""
    capped = _truncate_text(snippet or "", _PDF_CODE_MAX)
    lines = capped.replace("\t", "  ").splitlines() or [""]
    if len(lines) <= _PDF_CODE_LINES_PER_CHUNK:
        return [_code_box(label, "\n".join(lines), styles=styles, width=width)]

    out: list[Any] = []
    total = (len(lines) + _PDF_CODE_LINES_PER_CHUNK - 1) // _PDF_CODE_LINES_PER_CHUNK
    for i in range(0, len(lines), _PDF_CODE_LINES_PER_CHUNK):
        chunk = lines[i : i + _PDF_CODE_LINES_PER_CHUNK]
        part = i // _PDF_CODE_LINES_PER_CHUNK + 1
        out.append(_code_box(f"{label} ({part}/{total})", "\n".join(chunk), styles=styles, width=width))
        if part < total:
            out.append(Spacer(1, 1.5 * mm))
    return out


def _rail(inner: Any, *, content_width: float, pad: float, sev_color: colors.Color,
          bg: colors.Color, first: bool = False, last: bool = False) -> Table:
    """Card chrome for one flowable strip of a finding.

    Only a left severity rail plus top/bottom edges on the first and last
    strips. Consecutive strips therefore read as one continuous card while
    still being independently page-breakable."""
    shell = Table([[inner]], colWidths=[content_width])
    cmds: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, -1), bg),
        ("LINEBEFORE", (0, 0), (0, 0), 3.5, sev_color),
        ("LINEAFTER", (0, 0), (0, 0), 0.6, HAIRLINE),
        ("LEFTPADDING", (0, 0), (-1, -1), pad),
        ("RIGHTPADDING", (0, 0), (-1, -1), pad),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    if first:
        cmds.append(("LINEABOVE", (0, 0), (-1, 0), 0.6, HAIRLINE))
    if last:
        cmds.append(("LINEBELOW", (0, -1), (-1, -1), 0.6, HAIRLINE))
    shell.setStyle(TableStyle(cmds))
    return shell


def _finding_story_block(
    idx: int,
    finding: Finding,
    *,
    styles: Any,
    content_width: float,
) -> list[Any]:
    """One finding as page-breakable flowables sharing a continuous rail."""
    pad = 5 * mm
    inner_width = content_width - (2 * pad)
    prose_width = min(inner_width, _PROSE_MAX_WIDTH)
    sev = _severity_bucket(finding.severity)
    sev_color = SEVERITY_COLORS[sev]
    tint = SEVERITY_TINTS[sev]
    scanner_verdict = _scanner_verdict(finding)
    verdict = _display_verdict(scanner_verdict)  # the pill is short: "Unresolved", not its reason
    verdict_key = (scanner_verdict or "").strip().lower().replace("-", "_")
    # Each verdict picks its colour explicitly; an unrecognized one must never
    # fall through to the red "confirmed" colour.
    if verdict_key == "false_positive":
        verdict_bg = colors.HexColor("#64748b")   # slate — the scanner dismissed it
    elif verdict_key == UNVERIFIED_VERDICT:
        verdict_bg = colors.HexColor("#b45309")   # amber — nobody has decided it
    elif verdict_key == "true_positive":
        verdict_bg = colors.HexColor("#b91c1c")   # red — confirmed
    else:
        verdict_bg = colors.HexColor("#334155")   # dark slate — outside the contract

    title = finding.title or "Untitled finding"
    description = _description_text(finding)
    taint_flow = _taint_flow_text(finding)
    recommendation = _recommendation_text(finding)
    example_fix = _example_fix_text(finding)
    if not recommendation and finding.ai_fix and not example_fix:
        recommendation = (finding.ai_fix or "").strip()
    risk = _risk_text(finding)
    location_paths = _location_paths(finding)
    location_code = _location_code(finding)

    sev_w, verdict_w = 26 * mm, 34 * mm
    title_w = max(40 * mm, inner_width - sev_w - verdict_w - 4 * mm)
    header = Table(
        [[
            Paragraph(f'<font color="#64748b"><b>FINDING {idx}</b></font>&nbsp;&nbsp;'
                      f"<b>{_esc(title)}</b>", styles["FindingTitle"]),
            _pill(sev.upper(), sev_color, styles, width=sev_w),
            _pill(verdict, verdict_bg, styles, width=verdict_w),
        ]],
        colWidths=[title_w, sev_w + 2 * mm, verdict_w + 2 * mm],
    )
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (-1, 0), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 4 * mm),
        ("RIGHTPADDING", (1, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5 * mm),
    ]))

    meta_fields = [
        ("OWASP", _owasp_display(finding)),
        ("CWE", _cwe_display(finding)),
        ("CVSS", f"{finding.cvss:.1f}" if finding.cvss is not None else "—"),
        ("Confidence", _confidence_display(finding) or "—"),
        ("Affected files", str(_affected_file_count(finding))),
    ]
    # OWASP labels are long; give that column the room instead of splitting
    # the width into five equal parts and wrapping it to three lines.
    meta_weights = [0.36, 0.16, 0.12, 0.18, 0.18]
    meta = Table(
        [[Paragraph(f'<font color="#64748b"><b>{_esc(label.upper())}</b></font><br/>{_esc(value)}',
                    styles["FindingMeta"]) for label, value in meta_fields]],
        colWidths=[inner_width * w for w in meta_weights],
    )
    meta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), 0.5, HAIRLINE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 1.8 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8 * mm),
    ]))

    head_inner = Table([[header], [meta]], colWidths=[inner_width])
    head_inner.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    strips: list[Any] = [head_inner]
    strips.append(_labelled_block(
        "Verdict",
        f"Scanner: {_display_verdict(scanner_verdict, finding.unresolved_reason)}"
        f" · Review: {_review_state_label(finding)}"
        f" · {COUNTED_AS_PREFIX}{_counted_as(finding)}"
        + (f" · Construct: {mix}" if (mix := construct_mix_text(getattr(finding, "construct_mix", None))) else "")
        + (f" · Reviewed / Total: {cov}" if (cov := review_coverage_text(finding)) else ""),
        styles=styles, width=inner_width, label_color="#334155"))
    severity_text = _severity_text(finding)
    if severity_text:
        strips.append(_labelled_block("Severity", severity_text, styles=styles, width=inner_width,
                                      label_color="#334155"))

    if description:
        strips.append(_labelled_block("Description", description, styles=styles,
                                      width=prose_width, label_color="#1e3a5f"))
    if risk and risk != description:
        strips.append(_labelled_block("Risk", risk, styles=styles, width=prose_width,
                                      label_color="#991b1b", accent="#fef2f2"))
    if taint_flow:
        strips.append(_labelled_block("Data flow", taint_flow, styles=styles,
                                      width=inner_width, label_color="#0369a1", accent="#f0f9ff"))
    if location_paths:
        strips.append(_labelled_block("Affected location(s)", _location_paths_display(finding),
                                      styles=styles, width=inner_width, label_color="#334155"))
    if location_code:
        strips.extend(_code_boxes("Code context", location_code, styles=styles, width=inner_width))
    if recommendation:
        strips.append(_labelled_block("Recommendation", recommendation, styles=styles,
                                      width=prose_width, label_color="#065f46", accent="#ecfdf5"))
    if example_fix:
        strips.extend(_code_boxes("Example fix", example_fix, styles=styles, width=inner_width))
    occurrence_ids = _occurrence_ids(finding)
    if occurrence_ids:
        shown = ", ".join(occurrence_ids[:6])
        if len(occurrence_ids) > 6:
            shown += f" (+{len(occurrence_ids) - 6} more)"
        strips.append(_labelled_block("Occurrence ID", shown, styles=styles,
                                      width=inner_width, label_color="#64748b"))

    # Spacers between strips would show as gaps in the rail; drop them and let
    # the strip padding do the spacing instead.
    strips = [s for s in strips if not isinstance(s, Spacer)]

    out: list[Any] = [
        _Bookmark(f"{idx}. [{sev.upper()}] {title}", f"finding-{idx}"),
        CondPageBreak(45 * mm),
    ]
    for pos, strip in enumerate(strips):
        out.append(_rail(
            strip, content_width=content_width, pad=pad, sev_color=sev_color, bg=tint,
            first=(pos == 0), last=(pos == len(strips) - 1),
        ))
    out.append(Spacer(1, 5 * mm))
    return out


# =============================================================================
# PDF — summary page
# =============================================================================

def _gate_banner(gate: dict, styles: Any, width: float) -> Table:
    failed = bool(gate.get("gate_failed"))
    label = "QUALITY GATE FAILED" if failed else "QUALITY GATE PASSED"
    bg = colors.HexColor("#b91c1c") if failed else colors.HexColor("#047857")
    banner = Table([[Paragraph(f"<b>{label}</b>", styles["GateText"])]], colWidths=[width])
    banner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), bg),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3 * mm),
    ]))
    return banner


def _meta_grid(project: Project, scan: Scan, styles: Any, width: float) -> Table:
    llm = (
        "None"
        if (scan.metadata_json or {}).get("use_llm") is False
        or (not scan.llm_provider and not scan.llm_model)
        else _llm_display(scan)
    )
    fields = [
        ("Project", project.name or "—"),
        ("Scan", f"v{scan.version}"),
        ("Status", scan.status or "—"),
        ("Branch", scan.branch or "—"),
        ("Commit", (scan.commit_sha or "—")[:12]),
        ("Created", _fmt_dt(scan.created_at)),
        ("Repository", redact_git_url(scan.git_url) or "—"),
        ("Analysis engine", llm),
    ]
    cols = 4
    rows: list[list[Any]] = []
    for i in range(0, len(fields), cols):
        rows.append([
            Paragraph(f'<font color="#64748b"><b>{_esc(label.upper())}</b></font><br/>{_esc(value)}',
                      styles["FindingMeta"])
            for label, value in fields[i : i + cols]
        ])
    while rows and len(rows[-1]) < cols:
        rows[-1].append(Paragraph("", styles["FindingMeta"]))

    grid = Table(rows, colWidths=[width / cols] * cols)
    grid.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2 * mm),
    ]))
    return grid


def _severity_strip(
    counts: dict[str, int], styles: Any, width: float, *, occurrences: dict[str, int] | None = None,
) -> Table:
    """One cell per severity. With `occurrences` (SAST), `counts` are groups
    and each cell also says how many occurrences they cover."""
    def _sub(sev: str) -> str:
        if occurrences is None:
            return ""
        n, occ = counts.get(sev, 0), occurrences.get(sev, 0)
        return f'<br/><font size="7">{"group" if n == 1 else "groups"} · {occ} {"occurrence" if occ == 1 else "occurrences"}</font>'

    cells = [
        Paragraph(f'<font size="15"><b>{counts.get(sev, 0)}</b></font><br/>'
                  f'<font size="8">{sev.upper()}</font>{_sub(sev)}', styles["BadgeText"])
        for sev in SEVERITY_ORDER
    ]
    strip = Table([cells], colWidths=[width / len(SEVERITY_ORDER)] * len(SEVERITY_ORDER))
    cmds: list[tuple] = [
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4 * mm),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, colors.white),
    ]
    for idx, sev in enumerate(SEVERITY_ORDER):
        cmds.append(("BACKGROUND", (idx, 0), (idx, 0), SEVERITY_COLORS[sev]))
    strip.setStyle(TableStyle(cmds))
    return strip


def _owasp_table(rows: list[dict[str, Any]], styles: Any, width: float) -> Table | None:
    if not rows:
        return None
    data: list[list[Any]] = [[
        Paragraph("<b>OWASP CATEGORY</b>", styles["TableHead"]),
        Paragraph("<b>GROUPS</b>", styles["TableHeadCenter"]),
        Paragraph("<b>CRITICAL + HIGH</b>", styles["TableHeadCenter"]),
        Paragraph("<b>OCCURRENCES</b>", styles["TableHeadCenter"]),
    ]]
    for entry in rows:
        label, total, severe, occurrences = (
            entry["category"], entry["groups"], entry["severe"], entry["occurrences"]
        )
        data.append([
            Paragraph(_esc(label), styles["TableCell"]),
            Paragraph(str(total), styles["TableCellCenter"]),
            Paragraph(str(severe), styles["TableCellCenter"]),
            Paragraph(str(occurrences), styles["TableCellCenter"]),
        ])
    table = Table(data, colWidths=[width * 0.52, width * 0.14, width * 0.17, width * 0.17], repeatRows=1)
    table.hAlign = "LEFT"
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 1.6 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6 * mm),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    return table


def _findings_index(ordered: list[Finding], styles: Any, width: float) -> Table | None:
    """The summary-page index, numbered like the finding blocks (`_report_order`)."""
    if not ordered:
        return None
    shown = ordered[:_INDEX_MAX_ROWS]
    data: list[list[Any]] = [[
        Paragraph("<b>#</b>", styles["TableHeadCenter"]),
        Paragraph("<b>SEVERITY</b>", styles["TableHeadCenter"]),
        Paragraph("<b>ISSUE</b>", styles["TableHead"]),
        Paragraph("<b>OWASP</b>", styles["TableHead"]),
        Paragraph("<b>LOCATION</b>", styles["TableHead"]),
        Paragraph("<b>VERDICT</b>", styles["TableHead"]),
    ]]
    cmds: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 1.4 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.4 * mm),
    ]
    for idx, finding in enumerate(shown, start=1):
        sev = _severity_bucket(finding.severity)
        data.append([
            Paragraph(str(idx), styles["TableCellCenter"]),
            Paragraph(f'<font color="{_hex(SEVERITY_COLORS[sev])}">'
                      f"<b>{sev.upper()}</b></font>", styles["TableCellCenter"]),
            Paragraph(_esc(_truncate_text(finding.title or "Untitled finding", 150)), styles["TableCell"]),
            Paragraph(_esc(_owasp_header(finding)), styles["TableCell"]),
            Paragraph(_esc(_primary_location(finding)), styles["TableCellMono"]),
            Paragraph(_esc(_display_verdict(_scanner_verdict(finding))), styles["TableCell"]),
        ])
        if idx % 2 == 0:
            cmds.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#f8fafc")))
        cmds.append(("LINEBEFORE", (0, idx), (0, idx), 2.5, SEVERITY_COLORS[sev]))

    table = Table(
        data,
        colWidths=[width * w for w in (0.04, 0.09, 0.34, 0.17, 0.24, 0.12)],
        repeatRows=1,
    )
    table.hAlign = "LEFT"
    table.setStyle(TableStyle(cmds))
    return table


def _cwe_table(cwe_rows: list[dict[str, Any]], styles: Any, width: float) -> Table | None:
    """Summary-page table, one row per CWE: occurrences per verdict bucket and files."""
    if not cwe_rows:
        return None
    shown = cwe_rows[:_INDEX_MAX_ROWS]
    data: list[list[Any]] = [[
        Paragraph("<b>CWE</b>", styles["TableHead"]),
        Paragraph("<b>ISSUE</b>", styles["TableHead"]),
        Paragraph("<b>SEVERITY</b>", styles["TableHeadCenter"]),
        Paragraph("<b>CONFIRMED</b>", styles["TableHeadCenter"]),
        Paragraph("<b>UNRESOLVED</b>", styles["TableHeadCenter"]),
        Paragraph("<b>FALSE POS.</b>", styles["TableHeadCenter"]),
        Paragraph("<b>FILES</b>", styles["TableHeadCenter"]),
    ]]
    cmds: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 1.4 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.4 * mm),
    ]

    def occurrences(row: dict[str, Any], bucket: str) -> str:
        return str(row["by_bucket"].get(bucket, {}).get("occurrences", 0))

    for idx, row in enumerate(shown, start=1):
        sev = row["severity"]
        data.append([
            Paragraph(_esc(row["cwe"] or "—"), styles["TableCell"]),
            Paragraph(_esc(_truncate_text(row["title"] or "Untitled finding", 150)), styles["TableCell"]),
            Paragraph(f'<font color="{_hex(SEVERITY_COLORS[sev])}"><b>{sev.upper()}</b></font>',
                      styles["TableCellCenter"]),
            Paragraph(occurrences(row, "true_positive"), styles["TableCellCenter"]),
            Paragraph(occurrences(row, UNVERIFIED_VERDICT), styles["TableCellCenter"]),
            Paragraph(occurrences(row, "false_positive"), styles["TableCellCenter"]),
            Paragraph(str(row["file_count"]), styles["TableCellCenter"]),
        ])
        if idx % 2 == 0:
            cmds.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#f8fafc")))
        cmds.append(("LINEBEFORE", (0, idx), (0, idx), 2.5, SEVERITY_COLORS[sev]))
    table = Table(
        data, colWidths=[width * w for w in (0.09, 0.43, 0.1, 0.1, 0.1, 0.1, 0.08)], repeatRows=1,
    )
    table.hAlign = "LEFT"
    table.setStyle(TableStyle(cmds))
    return table


def _cwe_heading(row: dict[str, Any], bucket: str, styles: Any) -> Paragraph:
    """The line above one CWE's finding blocks within a verdict section. It must
    never read like a block heading or verdict line to `_check_pdf`."""
    cell = row["by_bucket"].get(bucket, {})
    groups, occ, files = cell.get("groups", 0), cell.get("occurrences", 0), cell.get("files", 0)
    label = row["cwe"] or "No CWE"
    return Paragraph(
        f"<b>{_esc(label)}</b>&nbsp;&nbsp;{_esc(_truncate_text(row['title'] or '', 120))}"
        f'&nbsp;&nbsp;<font color="#64748b">{groups} {"group" if groups == 1 else "groups"} · '
        f'{occ} {"occurrence" if occ == 1 else "occurrences"} · {files} {"file" if files == 1 else "files"}</font>',
        styles["FindingTitle"],
    )


# =============================================================================
# PDF — canvas / build
# =============================================================================

def _make_canvas(
    project_name: str,
    version: int,
    left: float,
    right: float,
    page: tuple,
    *,
    footer_prefix: str = "Secure Build",
):
    """Two-pass canvas so the footer can print 'Page X of Y'."""

    class _NumberedCanvas(pdfcanvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages: list[dict] = []

        def showPage(self) -> None:
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self) -> None:
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._footer(total)
                super().showPage()
            super().save()

        def _footer(self, total: int) -> None:
            self.saveState()
            self.setStrokeColor(HAIRLINE)
            self.setLineWidth(0.6)
            self.setFont("Helvetica", 8)
            self.setFillColor(MUTED)
            y = 9 * mm
            self.line(left, y, page[0] - right, y)
            self.drawString(left, y - 4.5 * mm, f"{footer_prefix} — {project_name} v{version}")
            self.drawCentredString(page[0] / 2.0, y - 4.5 * mm, "Confidential")
            self.drawRightString(page[0] - right, y - 4.5 * mm,
                                 f"Page {self._pageNumber} of {total}")
            self.restoreState()

    return _NumberedCanvas


def _build_styles() -> Any:
    styles = getSampleStyleSheet()

    def add(name: str, parent: str, **kwargs: Any) -> None:
        styles.add(ParagraphStyle(name=name, parent=styles[parent], **kwargs))

    add("ReportTitle", "Heading1", fontName="Helvetica-Bold", fontSize=24,
        textColor=INK, spaceAfter=1 * mm, leading=28)
    add("ReportSub", "Normal", fontName="Helvetica", fontSize=11,
        textColor=MUTED, spaceAfter=3 * mm, leading=14)
    add("SectionHeading", "Heading2", fontName="Helvetica-Bold", fontSize=13,
        textColor=INK, spaceBefore=1 * mm, spaceAfter=3 * mm, leading=17)
    add("FindingTitle", "Normal", fontName="Helvetica", fontSize=12,
        textColor=INK, leading=16)
    add("FindingLabel", "Normal", fontName="Helvetica-Bold", fontSize=7.5,
        textColor=MUTED, leading=10, spaceAfter=0)
    add("FindingBody", "Normal", fontName="Helvetica", fontSize=10,
        textColor=INK, leading=14)
    add("FindingMeta", "Normal", fontName="Helvetica", fontSize=9,
        textColor=INK, leading=12)
    add("FindingCode", "Normal", fontName="Courier", fontSize=_CODE_FONT_SIZE,
        textColor=INK, leading=11, wordWrap="CJK")
    add("BadgeText", "Normal", fontName="Helvetica-Bold", fontSize=8.5,
        textColor=colors.white, alignment=TA_CENTER, leading=11)
    add("GateText", "Normal", fontName="Helvetica-Bold", fontSize=13,
        textColor=colors.white, alignment=TA_CENTER, leading=16)
    add("TableHead", "Normal", fontName="Helvetica-Bold", fontSize=8,
        textColor=colors.white, alignment=TA_LEFT, leading=11)
    add("TableHeadCenter", "Normal", fontName="Helvetica-Bold", fontSize=8,
        textColor=colors.white, alignment=TA_CENTER, leading=11)
    add("TableCell", "Normal", fontName="Helvetica", fontSize=8.5,
        textColor=INK, leading=11.5)
    add("TableCellCenter", "Normal", fontName="Helvetica", fontSize=8.5,
        textColor=INK, alignment=TA_CENTER, leading=11.5)
    add("TableCellMono", "Normal", fontName="Courier", fontSize=7.5,
        textColor=INK, leading=10, wordWrap="CJK")
    return styles


def _sca_meta_grid(project: Project, scan: Scan, styles: Any, width: float) -> Table:
    inventory = _sca_inventory(scan)
    fields = [
        ("Project", project.name or "—"),
        ("Scan", f"v{scan.version}"),
        ("Status", scan.status or "—"),
        ("Branch", scan.branch or "—"),
        ("Commit", (scan.commit_sha or "—")[:12]),
        ("Created", _fmt_dt(scan.created_at)),
        ("Repository", redact_git_url(scan.git_url) or "—"),
        ("Lookup source", "OSV.dev"),
        ("Dependencies", str(inventory["total_dependencies"])),
        ("Vulnerable packages", str(inventory["vulnerable_dependencies"])),
        ("Manifests", ", ".join(inventory["manifests_parsed"]) or "—"),
        ("Ecosystems", ", ".join(inventory["ecosystems"]) or "—"),
    ]
    cols = 4
    rows: list[list[Any]] = []
    for i in range(0, len(fields), cols):
        rows.append([
            Paragraph(f'<font color="#64748b"><b>{_esc(label.upper())}</b></font><br/>{_esc(value)}',
                      styles["FindingMeta"])
            for label, value in fields[i : i + cols]
        ])
    while rows and len(rows[-1]) < cols:
        rows[-1].append(Paragraph("", styles["FindingMeta"]))

    grid = Table(rows, colWidths=[width / cols] * cols)
    grid.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2 * mm),
    ]))
    return grid


def _sca_findings_index(findings: list[Finding], styles: Any, width: float) -> Table | None:
    ordered = _sorted_findings(findings)
    if not ordered:
        return None
    shown = ordered[:_INDEX_MAX_ROWS]
    data: list[list[Any]] = [[
        Paragraph("<b>#</b>", styles["TableHeadCenter"]),
        Paragraph("<b>SEVERITY</b>", styles["TableHeadCenter"]),
        Paragraph("<b>PACKAGE</b>", styles["TableHead"]),
        Paragraph("<b>ADVISORY</b>", styles["TableHead"]),
        Paragraph("<b>CVSS</b>", styles["TableHeadCenter"]),
        Paragraph("<b>MANIFEST</b>", styles["TableHead"]),
    ]]
    cmds: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 1.4 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.4 * mm),
    ]
    for idx, finding in enumerate(shown, start=1):
        sev = _severity_bucket(finding.severity)
        cvss = f"{finding.cvss:.1f}" if finding.cvss is not None else "—"
        data.append([
            Paragraph(str(idx), styles["TableCellCenter"]),
            Paragraph(f'<font color="{_hex(SEVERITY_COLORS[sev])}">'
                      f"<b>{sev.upper()}</b></font>", styles["TableCellCenter"]),
            Paragraph(_esc(_truncate_text(_package_label(finding), 80)), styles["TableCellMono"]),
            Paragraph(_esc(_advisory_id(finding)), styles["TableCell"]),
            Paragraph(_esc(cvss), styles["TableCellCenter"]),
            Paragraph(_esc(_primary_location(finding)), styles["TableCellMono"]),
        ])
        if idx % 2 == 0:
            cmds.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#f8fafc")))
        cmds.append(("LINEBEFORE", (0, idx), (0, idx), 2.5, SEVERITY_COLORS[sev]))

    table = Table(
        data,
        colWidths=[width * w for w in (0.04, 0.10, 0.28, 0.20, 0.08, 0.30)],
        repeatRows=1,
    )
    table.hAlign = "LEFT"
    table.setStyle(TableStyle(cmds))
    return table


def _sbom_meta_grid(project: Project, scan: Scan, data: dict[str, Any], styles: Any, width: float) -> Table:
    fields = [
        ("Project", project.name or "—"),
        ("Scan", f"v{scan.version}"),
        ("Status", scan.status or "—"),
        ("Branch", scan.branch or "—"),
        ("Commit", (scan.commit_sha or "—")[:12]),
        ("Created", _fmt_dt(scan.created_at)),
        ("Repository", redact_git_url(scan.git_url) or "—"),
        ("Lookup source", "deps.dev"),
        ("Total components", str(data["total_components"])),
        ("Manifests", ", ".join(data["manifests_parsed"]) or "—"),
        ("Ecosystems", ", ".join(data["ecosystems"]) or "—"),
    ]
    cols = 4
    rows: list[list[Any]] = []
    for i in range(0, len(fields), cols):
        rows.append([
            Paragraph(f'<font color="#64748b"><b>{_esc(label.upper())}</b></font><br/>{_esc(value)}',
                      styles["FindingMeta"])
            for label, value in fields[i : i + cols]
        ])
    while rows and len(rows[-1]) < cols:
        rows[-1].append(Paragraph("", styles["FindingMeta"]))

    grid = Table(rows, colWidths=[width / cols] * cols)
    grid.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2 * mm),
    ]))
    return grid


def _cost_strip(license_summary: dict[str, int], styles: Any, width: float) -> Table:
    cells = [
        Paragraph(f'<font size="15"><b>{license_summary.get(level, 0)}</b></font><br/>'
                  f'<font size="8">{COST_LABELS[level].upper()}</font>', styles["BadgeText"])
        for level in COST_ORDER
    ]
    strip = Table([cells], colWidths=[width / len(COST_ORDER)] * len(COST_ORDER))
    cmds: list[tuple] = [
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4 * mm),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, colors.white),
    ]
    for idx, level in enumerate(COST_ORDER):
        cmds.append(("BACKGROUND", (idx, 0), (idx, 0), COST_COLORS[level]))
    strip.setStyle(TableStyle(cmds))
    return strip


def _sbom_components_table(components: list[dict[str, Any]], styles: Any, width: float) -> Table | None:
    ordered = _sorted_components(components)
    if not ordered:
        return None
    shown = ordered[:_INDEX_MAX_ROWS]
    data: list[list[Any]] = [[
        Paragraph("<b>#</b>", styles["TableHeadCenter"]),
        Paragraph("<b>COMPONENT NAME</b>", styles["TableHead"]),
        Paragraph("<b>VERSION</b>", styles["TableHeadCenter"]),
        Paragraph("<b>SUPPLIER</b>", styles["TableHead"]),
        Paragraph("<b>TYPE</b>", styles["TableHeadCenter"]),
        Paragraph("<b>LICENSE</b>", styles["TableHead"]),
        Paragraph("<b>DEPENDENCY RELATIONSHIP</b>", styles["TableHead"]),
        Paragraph("<b>UNIQUE IDENTIFIER (PURL/CPE)</b>", styles["TableHead"]),
    ]]
    cmds: list[tuple] = [
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("BOX", (0, 0), (-1, -1), 0.5, RULE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.2 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 1.4 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.4 * mm),
    ]
    for idx, comp in enumerate(shown, start=1):
        type_label = _component_type_label(comp)
        type_color = TYPE_COLORS[type_label]
        data.append([
            Paragraph(str(idx), styles["TableCellCenter"]),
            Paragraph(_esc(_truncate_text(_component_label(comp), 80)), styles["TableCellMono"]),
            Paragraph(_esc(comp.get("version") or "—"), styles["TableCellCenter"]),
            Paragraph(_esc(_component_supplier(comp)), styles["TableCell"]),
            Paragraph(f'<font color="{_hex(type_color)}"><b>{_esc(type_label.upper())}</b></font>',
                      styles["TableCellCenter"]),
            Paragraph(_esc(_component_license_label(comp)), styles["TableCell"]),
            Paragraph(_esc(_component_relationship(comp)), styles["TableCell"]),
            Paragraph(_esc(comp.get("purl") or "—"), styles["TableCellMono"]),
        ])
        if idx % 2 == 0:
            cmds.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#f8fafc")))
        cmds.append(("LINEBEFORE", (0, idx), (0, idx), 2.5, type_color))

    table = Table(
        data,
        colWidths=[width * w for w in (0.03, 0.18, 0.09, 0.16, 0.09, 0.12, 0.15, 0.18)],
        repeatRows=1,
    )
    table.hAlign = "LEFT"
    table.setStyle(TableStyle(cmds))
    return table


def _build_pdf_sbom_report(project: Project, scan: Scan) -> bytes:
    data = _sbom_data(scan)
    buf = io.BytesIO()
    page = landscape(A4)
    left = right = 13 * mm
    top, bottom = 12 * mm, 18 * mm
    content_width = page[0] - left - right

    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=left,
        rightMargin=right,
        topMargin=top,
        bottomMargin=bottom,
        title=f"{project.name} — Secure Build SBOM Report v{scan.version}",
        author="Secure Build",
        subject="Software bill of materials report",
    )

    styles = _build_styles()
    story: list[Any] = [
        _Bookmark("Summary", "summary"),
        Paragraph(_esc(project.name or "Application"), styles["ReportTitle"]),
        Paragraph(
            f"Secure Build — Software Bill of Materials (SBOM) Report"
            f"&nbsp;&nbsp;·&nbsp;&nbsp;v{scan.version}"
            f"&nbsp;&nbsp;·&nbsp;&nbsp;{_esc(_fmt_dt(scan.created_at))}",
            styles["ReportSub"],
        ),
    ]

    status_note = _scan_status_banner(scan)
    if status_note:
        story.append(Paragraph(f'<font color="#b45309"><b>{_esc(status_note)}</b></font>',
                               styles["ReportSub"]))

    story.append(_sbom_meta_grid(project, scan, data, styles, content_width))
    story.append(Spacer(1, 5 * mm))
    story.append(_cost_strip(data["license_summary"], styles, content_width))
    story.append(Spacer(1, 6 * mm))

    components = data["components"]
    table = _sbom_components_table(components, styles, content_width)
    if table is not None:
        extra = max(0, len(components) - _INDEX_MAX_ROWS)
        story.append(Paragraph("COMPONENT INVENTORY", styles["SectionHeading"]))
        story.append(table)
        if extra:
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph(
                f'<font color="#64748b">… and {extra} further component(s); '
                f"see the Excel export for the full list.</font>", styles["FindingBody"]))
    else:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph("No components were found in this project's dependency manifests.",
                               styles["FindingBody"]))

    doc.build(
        story,
        canvasmaker=_make_canvas(
            project.name or "Application",
            scan.version,
            left,
            right,
            page,
            footer_prefix="Secure Build SBOM",
        ),
    )
    return buf.getvalue()


def _sca_finding_story_block(
    idx: int,
    finding: Finding,
    *,
    styles: Any,
    content_width: float,
) -> list[Any]:
    """One SCA advisory as a page-breakable card — package / CVE / upgrade, not
    file/line/CWE/OWASP/data-flow/code-snippet."""
    pad = 5 * mm
    inner_width = content_width - (2 * pad)
    prose_width = min(inner_width, _PROSE_MAX_WIDTH)
    sev = _severity_bucket(finding.severity)
    sev_color = SEVERITY_COLORS[sev]
    tint = SEVERITY_TINTS[sev]
    package = _package_label(finding)
    advisory = _advisory_id(finding)

    sev_w = 26 * mm
    title_w = max(40 * mm, inner_width - sev_w - 4 * mm)
    header = Table(
        [[
            Paragraph(f'<font color="#64748b"><b>ADVISORY {idx}</b></font>&nbsp;&nbsp;'
                      f"<b>{_esc(package)}</b>", styles["FindingTitle"]),
            _pill(sev.upper(), sev_color, styles, width=sev_w),
        ]],
        colWidths=[title_w, sev_w + 2 * mm],
    )
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 4 * mm),
        ("RIGHTPADDING", (1, 0), (1, 0), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5 * mm),
    ]))

    cvss = f"{finding.cvss:.1f}" if finding.cvss is not None else "—"
    meta_fields = [
        ("Advisory", advisory),
        ("CVSS", cvss),
        ("Manifest", _primary_location(finding)),
        ("Affected manifests", str(_affected_file_count(finding))),
    ]
    meta = Table(
        [[Paragraph(f'<font color="#64748b"><b>{_esc(label.upper())}</b></font><br/>{_esc(value)}',
                    styles["FindingMeta"]) for label, value in meta_fields]],
        colWidths=[inner_width * w for w in (0.28, 0.12, 0.42, 0.18)],
    )
    meta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), 0.5, HAIRLINE),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, HAIRLINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 1.8 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8 * mm),
    ]))

    head_inner = Table([[header], [meta]], colWidths=[inner_width])
    head_inner.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    strips: list[Any] = [head_inner]
    summary = _advisory_summary(finding)
    if summary:
        strips.append(_labelled_block("Advisory summary", summary, styles=styles,
                                      width=prose_width, label_color="#1e3a5f"))
    locations = _location_paths_display(finding)
    if locations and _affected_file_count(finding) > 1:
        strips.append(_labelled_block("Declared in", locations, styles=styles,
                                      width=inner_width, label_color="#334155"))
    recommendation = (finding.ai_fix or "").strip()
    if recommendation:
        strips.append(_labelled_block("Suggested upgrade", recommendation, styles=styles,
                                      width=prose_width, label_color="#065f46", accent="#ecfdf5"))

    out: list[Any] = [
        _Bookmark(f"{idx}. [{sev.upper()}] {package} {advisory}", f"finding-{idx}"),
        CondPageBreak(45 * mm),
    ]
    for pos, strip in enumerate(strips):
        out.append(_rail(
            strip, content_width=content_width, pad=pad, sev_color=sev_color, bg=tint,
            first=(pos == 0), last=(pos == len(strips) - 1),
        ))
    out.append(Spacer(1, 5 * mm))
    return out


def _build_pdf_sca_report(project: Project, scan: Scan, findings: list[Finding]) -> bytes:
    counts = _count_by_severity(findings)
    buf = io.BytesIO()
    page = landscape(A4)
    left = right = 13 * mm
    top, bottom = 12 * mm, 18 * mm
    content_width = page[0] - left - right

    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=left,
        rightMargin=right,
        topMargin=top,
        bottomMargin=bottom,
        title=f"{project.name} — Secure Build SCA Report v{scan.version}",
        author="Secure Build",
        subject="Software composition analysis report",
    )

    styles = _build_styles()
    story: list[Any] = [
        _Bookmark("Summary", "summary"),
        Paragraph(_esc(project.name or "Application"), styles["ReportTitle"]),
        Paragraph(
            f"Secure Build — Dependency Vulnerability (SCA) Report"
            f"&nbsp;&nbsp;·&nbsp;&nbsp;v{scan.version}"
            f"&nbsp;&nbsp;·&nbsp;&nbsp;{_esc(_fmt_dt(scan.created_at))}",
            styles["ReportSub"],
        ),
    ]

    status_note = _scan_status_banner(scan)
    if status_note:
        story.append(Paragraph(f'<font color="#b45309"><b>{_esc(status_note)}</b></font>',
                               styles["ReportSub"]))

    story.append(_sca_meta_grid(project, scan, styles, content_width))
    story.append(Spacer(1, 5 * mm))
    story.append(_severity_strip(counts, styles, content_width))
    story.append(Spacer(1, 6 * mm))

    index = _sca_findings_index(findings, styles, content_width)
    if index is not None:
        extra = max(0, len(findings) - _INDEX_MAX_ROWS)
        story.append(Paragraph("ADVISORIES INDEX", styles["SectionHeading"]))
        story.append(index)
        if extra:
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph(
                f'<font color="#64748b">… and {extra} further advisory(ies); '
                f"see the detail pages and the Excel export.</font>", styles["FindingBody"]))

    if not findings:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph("No known vulnerabilities were identified in declared dependencies.",
                               styles["FindingBody"]))
    else:
        story.append(PageBreak())
        story.append(_Bookmark("Advisories", "findings"))
        story.append(Paragraph("DETAILED ADVISORIES", styles["SectionHeading"]))
        for idx, finding in enumerate(_sorted_findings(findings), start=1):
            story.extend(_sca_finding_story_block(idx, finding, styles=styles,
                                                  content_width=content_width))

    doc.build(
        story,
        canvasmaker=_make_canvas(
            project.name or "Application",
            scan.version,
            left,
            right,
            page,
            footer_prefix="Secure Build SCA",
        ),
    )
    return buf.getvalue()


def build_pdf_report(project: Project, scan: Scan, findings: list[Finding]) -> bytes:
    """Landscape PDF: summary page (gate, metadata, severity, OWASP, index),
    then one card per finding. SCA scans use a separate dependency-advisory
    layout instead of the SAST finding card; SBOM scans use a component-
    inventory layout instead (see module docstring). DAST scans use the
    report built by `dast_export` (evidence, unverified section, diff)."""
    if _is_dast_scan(scan):
        from app.services.dast_export import build_dast_pdf  # lazy: dast_export builds on this module

        return build_dast_pdf(project, scan, findings)
    if _is_sca_scan(scan):
        return _build_pdf_sca_report(project, scan, findings)
    if _is_sbom_scan(scan):
        return _build_pdf_sbom_report(project, scan)

    summary = count_findings(findings)
    group_counts = {sev: v["groups"] for sev, v in summary["actionable"].items()}
    occurrence_counts = {sev: v["occurrences"] for sev, v in summary["actionable"].items()}
    gate = build_gate_summary(project, findings, counts=summary)
    buf = io.BytesIO()
    page = landscape(A4)
    left = right = 13 * mm
    top, bottom = 12 * mm, 18 * mm
    content_width = page[0] - left - right

    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=left,
        rightMargin=right,
        topMargin=top,
        bottomMargin=bottom,
        title=f"{project.name} — Secure Build SAST Report v{scan.version}",
        author="Secure Build",
        subject="Static application security testing report",
    )

    styles = _build_styles()
    story: list[Any] = [
        _Bookmark("Summary", "summary"),
        Paragraph(_esc(project.name or "Application"), styles["ReportTitle"]),
        Paragraph(f"Secure Build — SAST Scan Report&nbsp;&nbsp;·&nbsp;&nbsp;v{scan.version}"
                  f"&nbsp;&nbsp;·&nbsp;&nbsp;{_esc(_fmt_dt(scan.created_at))}", styles["ReportSub"]),
    ]

    status_note = _scan_status_banner(scan)
    if status_note:
        story.append(Paragraph(f'<font color="#b45309"><b>{_esc(status_note)}</b></font>',
                               styles["ReportSub"]))

    story.append(_gate_banner(gate, styles, content_width))
    story.append(Spacer(1, 4 * mm))
    story.append(_meta_grid(project, scan, styles, content_width))
    story.append(Spacer(1, 5 * mm))
    story.append(_severity_strip(group_counts, styles, content_width, occurrences=occurrence_counts))
    story.append(Spacer(1, 2 * mm))
    bucket_text = " · ".join(
        f"{label}: {t['groups']} {'group' if t['groups'] == 1 else 'groups'} "
        f"({t['occurrences']} {'occurrence' if t['occurrences'] == 1 else 'occurrences'})"
        for label, t in verdict_bucket_totals(summary).items()
    )
    story.append(Paragraph(
        f'<font color="#64748b">Actionable issue groups (confirmed + unresolved) by severity. '
        f"{_esc(bucket_text)}.</font>",
        styles["FindingBody"]))
    story.append(Spacer(1, 6 * mm))

    owasp_table = _owasp_table(summary["by_owasp"], styles, content_width * 0.7)
    if owasp_table is not None:
        story.append(KeepTogether([
            Paragraph("OWASP DISTRIBUTION", styles["SectionHeading"]),
            owasp_table,
        ]))
        story.append(Spacer(1, 6 * mm))

    cwe_rows = group_by_cwe(findings)
    ordered = _report_order(findings, cwe_rows)
    cwe_table = _cwe_table(cwe_rows, styles, content_width)
    if cwe_table is not None:
        story.append(Paragraph("FINDINGS BY CWE", styles["SectionHeading"]))
        story.append(cwe_table)
        if len(cwe_rows) > _INDEX_MAX_ROWS:
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph(
                f'<font color="#64748b">… and {len(cwe_rows) - _INDEX_MAX_ROWS} further CWE(s); '
                f'see the "By CWE" sheet of the Excel export.</font>', styles["FindingBody"]))
        story.append(Spacer(1, 6 * mm))

    index = _findings_index(ordered, styles, content_width)
    if index is not None:
        extra = max(0, len(findings) - _INDEX_MAX_ROWS)
        story.append(Paragraph("FINDINGS INDEX", styles["SectionHeading"]))
        story.append(index)
        if extra:
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph(
                f'<font color="#64748b">… and {extra} further finding(s); '
                f"see the detail pages and the Excel export.</font>", styles["FindingBody"]))

    if not findings:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph("No vulnerabilities were identified in this scan.",
                               styles["FindingBody"]))
    else:
        story.append(PageBreak())
        story.append(_Bookmark("Findings", "findings"))
        story.append(Paragraph("DETAILED FINDINGS", styles["SectionHeading"]))
        rows_by_key = {row["key"]: row for row in cwe_rows}
        section: tuple[str, str] | None = None
        for idx, finding in enumerate(ordered, start=1):
            bucket, key = _verdict_bucket(finding), _cwe_group_key(finding)
            if section is None or section[0] != bucket:
                story.append(Paragraph(_BUCKET_LABELS[bucket].upper(), styles["SectionHeading"]))
            if section != (bucket, key):
                story.append(_cwe_heading(rows_by_key[key], bucket, styles))
                story.append(Spacer(1, 2 * mm))
                section = (bucket, key)
            story.extend(_finding_story_block(idx, finding, styles=styles,
                                              content_width=content_width))

    doc.build(story, canvasmaker=_make_canvas(project.name or "Application", scan.version,
                                              left, right, page))
    return buf.getvalue()


# =============================================================================
# Report consistency
# =============================================================================

_COVER_BUCKET_RE = re.compile(
    r"(Confirmed|Unresolved|False Positive|Unknown): (\d+) groups? \((\d+) occurrences?\)"
)


def _recount_rows(rows: list[dict[str, Any]]) -> dict[str, dict[str, dict[str, int]]]:
    """Per-verdict, per-severity counts recomputed from serialized rows, without
    `count_findings` — so a disagreement between the API's numbers and its own
    rows cannot hide behind a shared bug."""
    out: dict[str, dict[str, dict[str, int]]] = {}
    for row in rows:
        verdict = (row.get("verdict") or "").strip().lower()
        bucket = verdict if verdict in VERDICT_BUCKETS else UNKNOWN_BUCKET
        cell = out.setdefault(bucket, {}).setdefault(
            _severity_bucket(row.get("severity")), {"groups": 0, "occurrences": 0}
        )
        cell["groups"] += 1
        cell["occurrences"] += row.get("occurrence_count") or 1
    return out


def _check_construct_mix(rows: list[dict[str, Any]]) -> list[str]:
    """Each row's `construct_mix` must equal a recount, from the rows themselves,
    of the members behind every row that shares its `construct_key`."""
    totals: dict[str, dict[str, int]] = {}
    for row in rows:
        verdict = (row.get("verdict") or "").strip().lower()
        bucket = "tp" if verdict == "true_positive" else "fp" if verdict == "false_positive" else "unresolved"
        totals.setdefault(row.get("construct_key") or "", {"tp": 0, "fp": 0, "unresolved": 0})[bucket] += (
            row.get("occurrence_count") or 1
        )
    problems = []
    for row in rows:
        expected = totals.get(row.get("construct_key") or "")
        if row.get("construct_mix") != expected:
            problems.append(
                f"API construct mix: {row.get('title')!r} says {row.get('construct_mix')}, "
                f"the rows sharing its construct give {expected}"
            )
    return problems


def _check_api(counts: dict[str, Any], grouped: list[Finding]) -> list[str]:
    from app.schemas import FindingOut  # lazy: keeps this module import-light

    rows = [FindingOut.model_validate(f).model_dump() for f in grouped]
    recomputed = _recount_rows(rows)
    problems: list[str] = []
    for bucket in (*VERDICT_BUCKETS, UNKNOWN_BUCKET):
        for sev in SEVERITY_ORDER:
            stated = (
                counts["by_verdict"][bucket][sev] if bucket in counts["by_verdict"]
                else {"groups": 0, "occurrences": 0}
            )
            if bucket == UNKNOWN_BUCKET:
                stated = {"groups": 0, "occurrences": 0}  # counted in `actionable`, see below
            actual = recomputed.get(bucket, {}).get(sev, {"groups": 0, "occurrences": 0})
            if bucket != UNKNOWN_BUCKET and stated != actual:
                problems.append(
                    f"API counts: {_BUCKET_LABELS[bucket]} / {sev} says {stated}, its rows give {actual}"
                )
    problems += _check_construct_mix(rows)
    for sev in SEVERITY_ORDER:
        expected_groups = sum(
            recomputed.get(b, {}).get(sev, {"groups": 0})["groups"]
            for b in ("true_positive", UNVERIFIED_VERDICT, UNKNOWN_BUCKET)
        )
        if counts["actionable"][sev]["groups"] != expected_groups:
            problems.append(
                f"API counts: actionable {sev} says {counts['actionable'][sev]['groups']} groups, "
                f"its rows give {expected_groups}"
            )
    return problems


def _check_xlsx(data: bytes, counts: dict[str, Any], grouped: list[Finding] | None = None) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data))
    summary, findings_sheet = workbook["Summary"], workbook["Findings"]
    stated: dict[str, tuple[int, int]] = {}
    for row in summary.iter_rows(min_row=1, values_only=True):
        label = row[0]
        if isinstance(label, str) and label.startswith(BUCKET_ROW_PREFIX):
            stated[label[len(BUCKET_ROW_PREFIX):]] = (row[1], row[2])
    headers = [c.value for c in findings_sheet[1]]
    column = headers.index("Counted As")
    rows_by_bucket: dict[str, int] = {}
    for row in findings_sheet.iter_rows(min_row=2, values_only=True):
        rows_by_bucket[row[column]] = rows_by_bucket.get(row[column], 0) + 1

    problems: list[str] = []
    if grouped is not None:
        mix_col, cov_col = headers.index("Construct Mix"), headers.index("Reviewed / Total")
        sheet_rows = list(findings_sheet.iter_rows(min_row=2, values_only=True))
        expected_rows = _sorted_findings(grouped)
        if len(sheet_rows) != len(expected_rows):
            problems.append(f"Excel findings sheet: {len(sheet_rows)} rows, the API has {len(expected_rows)}")
        for number, (row, finding) in enumerate(zip(sheet_rows, expected_rows), start=1):
            want_mix = construct_mix_text(getattr(finding, "construct_mix", None))
            if (row[mix_col] or "") != want_mix:
                problems.append(f"Excel row {number}: Construct Mix says {row[mix_col]!r}, the API says {want_mix!r}")
            if (row[cov_col] or "") != review_coverage_text(finding):
                problems.append(
                    f"Excel row {number}: Reviewed / Total says {row[cov_col]!r}, "
                    f"the API says {review_coverage_text(finding)!r}"
                )
    for label, cell in verdict_bucket_totals(counts).items():
        summary_groups, summary_occurrences = stated.get(label.lower(), (None, None))
        if (summary_groups, summary_occurrences) != (cell["groups"], cell["occurrences"]):
            problems.append(
                f"Excel summary: {label} says {summary_groups} groups / {summary_occurrences} "
                f"occurrences, the counts say {cell['groups']} / {cell['occurrences']}"
            )
        listed = rows_by_bucket.get(label, 0)
        if listed != cell["groups"]:
            problems.append(f"Excel findings sheet: {label} lists {listed} rows, the summary says {cell['groups']}")
    extra = set(rows_by_bucket) - set(verdict_bucket_totals(counts))
    for label in sorted(extra):
        problems.append(f"Excel findings sheet: {rows_by_bucket[label]} row(s) counted as {label!r}, which the summary has no bucket for")
    return problems


def _check_pdf(data: bytes, counts: dict[str, Any]) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError:  # pypdf is a runtime dependency now; if it is missing say so
        logger.warning("pypdf is not installed; the PDF half of the consistency check was skipped")
        return []
    reader = PdfReader(io.BytesIO(data))
    text = " ".join(" ".join(page.extract_text() or "" for page in reader.pages).split())
    cover = {m.group(1): (int(m.group(2)), int(m.group(3))) for m in _COVER_BUCKET_RE.finditer(text)}
    blocks: dict[str, int] = {}
    for label in re.findall(re.escape(COUNTED_AS_PREFIX) + r"(Confirmed|Unresolved|False Positive|Unknown)", text):
        blocks[label] = blocks.get(label, 0) + 1
    headings = len(re.findall(r"FINDING \d+", text))

    problems: list[str] = []
    expected_total = sum(cell["groups"] for cell in verdict_bucket_totals(counts).values())
    for label, cell in verdict_bucket_totals(counts).items():
        shown = cover.get(label)
        if shown != (cell["groups"], cell["occurrences"]):
            problems.append(
                f"PDF cover: {label} says {shown}, the counts say "
                f"({cell['groups']}, {cell['occurrences']})"
            )
        if blocks.get(label, 0) != cell["groups"]:
            problems.append(
                f"PDF findings: {label} has {blocks.get(label, 0)} finding block(s), "
                f"the cover says {cell['groups']}"
            )
    if headings != expected_total:
        problems.append(f"PDF findings: {headings} finding heading(s), the cover totals {expected_total}")
    return problems


def check_report_consistency(db: Any, scan: Scan) -> list[str]:
    """Contradictions between the numbers a scan reports in its API, its Excel
    export and its PDF, as a list of messages; empty means consistent.

    It groups the scan's findings once (as the API and exports do) and checks
    (1) the API's counts against a recount of the API's own rows, (2) the Excel
    summary against the rows of the Excel findings sheet, and (3) the PDF cover
    against the finding blocks rendered in the PDF. A SAST scan only; other scan
    types have their own reports and return no contradictions."""
    if (getattr(scan, "scan_type", None) or "sast").strip().lower() != "sast":
        return []
    from app.services.feedback import effective_scan_findings  # lazy: feedback imports this module

    grouped = group_findings_by_issue(effective_scan_findings(db, scan))
    counts = count_findings(grouped)
    project = scan.project
    problems = _check_api(counts, grouped)
    problems += _check_xlsx(build_excel_report(project, scan, grouped), counts, grouped)
    problems += _check_pdf(build_pdf_report(project, scan, grouped), counts)
    return problems
