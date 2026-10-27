"""The rescan decision diff (SAST-00-090).

Two scans of one project are compared occurrence by occurrence — joined only on
`occurrence_id`, which survives rescans and line shifts — and every change of the
scanner's decision is listed, in both directions, with the stage that made it. A
scanner change is then judged by what it changed, not by how many rows it removed:
a fix that quietly turns confirmed vulnerabilities into false positives shows up
first.

`diff()` is pure over plain records, so the release gate reuses it unchanged to
compare a corpus run with its baseline. `records_for_scan()` builds the records
from the scan's current `Finding` rows (the scanner's own verdict, not the
reviewer overlay) and each finding's latest `FindingDecision`.
"""

from __future__ import annotations

import io
from typing import Any, Iterable, Mapping

from sqlalchemy.orm import Session

from app.models import UNVERIFIED_VERDICT, Finding, FindingDecision, Scan

ADDED, REMOVED = "added", "removed"
VERDICT_CHANGED, CLASS_CHANGED, UNCHANGED = "verdict_changed", "issue_class_changed", "unchanged"
CHANGES = (ADDED, REMOVED, VERDICT_CHANGED, CLASS_CHANGED, UNCHANGED)

# Summary transitions lead with the ones a reviewer must see first.
_LEAD_TRANSITIONS = (("true_positive", "false_positive"), ("true_positive", REMOVED))
_PREFERENCE = {"true_positive": 0, UNVERIFIED_VERDICT: 1, "false_positive": 2}

_STATE_FIELDS = ("verdict", "issue_class", "unresolved_reason", "confidence")


class DiffError(ValueError):
    """The two scans cannot be compared; the message says why (an HTTP 422)."""


# --- records ------------------------------------------------------------------

def latest_decisions(
    db: Session, scan_id: str, *, stages: frozenset[str] | None = None
) -> dict[str, FindingDecision]:
    """{finding id: its newest `FindingDecision`}, optionally among `stages` only."""
    query = db.query(FindingDecision).filter(FindingDecision.scan_id == scan_id)
    if stages is not None:
        query = query.filter(FindingDecision.stage.in_(stages))
    latest: dict[str, FindingDecision] = {}
    for decision in query.order_by(FindingDecision.at.desc(), FindingDecision.id.desc()).all():
        latest.setdefault(decision.finding_id, decision)
    return latest


def records_for_scan(db: Session, scan: Scan) -> tuple[list[dict[str, Any]], int]:
    """(records, rows without an occurrence id). One record per finding row that has
    an `occurrence_id`, carrying its latest decision's stage and reasoning."""
    rows = db.query(Finding).filter(Finding.scan_id == scan.id).all()
    with_identity = [row for row in rows if row.occurrence_id]
    latest = latest_decisions(db, scan.id) if with_identity else {}
    records = []
    for row in with_identity:
        decision = latest.get(row.id)
        records.append({
            "occurrence_id": row.occurrence_id,
            "rule_id": row.rule_id,
            "title": row.title,
            "severity": row.severity,
            "file_path": row.file_path,
            "line": row.line,
            "identity_quality": row.identity_quality,
            "verdict": row.verdict,
            "issue_class": row.issue_class,
            "unresolved_reason": row.unresolved_reason,
            "confidence": row.confidence,
            "stage": decision.stage if decision else None,
            "verdict_reason": (decision.verdict_reason if decision else None) or row.verdict_reason,
        })
    return records, len(rows) - len(with_identity)


def _representative(records: list[Mapping[str, Any]]) -> Mapping[str, Any]:
    """When one occurrence id has several rows, the one that says the most: a
    confirmation over an undecided answer over a dismissal, then the surest."""
    return min(
        records,
        key=lambda r: (_PREFERENCE.get(r.get("verdict") or "", 3), -(r.get("confidence") or 0)),
    )


def _index(records: Iterable[Mapping[str, Any]]) -> dict[str, tuple[Mapping[str, Any], int]]:
    grouped: dict[str, list[Mapping[str, Any]]] = {}
    for record in records:
        grouped.setdefault(record["occurrence_id"], []).append(record)
    return {key: (_representative(group), len(group)) for key, group in grouped.items()}


# --- the diff -----------------------------------------------------------------

def _state(record: Mapping[str, Any] | None) -> dict[str, Any] | None:
    return None if record is None else {name: record.get(name) for name in _STATE_FIELDS}


def _classify(old: Mapping[str, Any] | None, new: Mapping[str, Any] | None) -> str:
    if old is None:
        return ADDED
    if new is None:
        return REMOVED
    if old.get("verdict") != new.get("verdict"):
        return VERDICT_CHANGED
    if old.get("issue_class") != new.get("issue_class"):
        return CLASS_CHANGED
    return UNCHANGED


def diff(
    base_records: Iterable[Mapping[str, Any]],
    new_records: Iterable[Mapping[str, Any]],
    *,
    include_unchanged: bool = False,
) -> dict[str, Any]:
    """Compare two record lists. Returns {entries, counts, summary}; see the module
    docstring. `entries` omits `unchanged` unless `include_unchanged`."""
    old_index, new_index = _index(base_records), _index(new_records)
    entries: list[dict[str, Any]] = []
    counts = {name: 0 for name in CHANGES}
    transitions: dict[tuple[str, str], int] = {}
    by_rule: dict[str, dict[str, int]] = {}

    for occurrence_id in sorted(set(old_index) | set(new_index)):
        old, old_n = old_index.get(occurrence_id, (None, 0))
        new, new_n = new_index.get(occurrence_id, (None, 0))
        change = _classify(old, new)
        counts[change] += 1
        anchor = new or old
        if change != UNCHANGED:
            source = (old.get("verdict") if old else ADDED) or "none"
            target = (new.get("verdict") if new else REMOVED) or "none"
            if change != CLASS_CHANGED:
                transitions[(source, target)] = transitions.get((source, target), 0) + 1
            rule_counts = by_rule.setdefault(anchor.get("rule_id") or "unknown", {c: 0 for c in CHANGES[:4]})
            rule_counts[change] += 1
        if change == UNCHANGED and not include_unchanged:
            continue
        entries.append({
            "occurrence_id": occurrence_id,
            "change": change,
            "rule_id": anchor.get("rule_id"),
            "title": anchor.get("title"),
            "severity": anchor.get("severity"),
            "file_path": anchor.get("file_path"),
            "line": anchor.get("line"),
            "identity_quality": anchor.get("identity_quality"),
            "old": _state(old),
            "new": _state(new),
            # The *new* decision: which stage made it, and why.
            "stage": new.get("stage") if new else None,
            "verdict_reason": new.get("verdict_reason") if new else None,
            "members": {"old": old_n, "new": new_n},
        })

    def order(item: tuple[tuple[str, str], int]) -> tuple[int, int, str]:
        key, count = item
        lead = _LEAD_TRANSITIONS.index(key) if key in _LEAD_TRANSITIONS else len(_LEAD_TRANSITIONS)
        return (lead, -count, f"{key[0]}>{key[1]}")

    entries.sort(key=lambda e: (CHANGES.index(e["change"]), e["file_path"] or "", e["line"] or 0, e["occurrence_id"]))
    return {
        "entries": entries,
        "counts": counts,
        "summary": {
            "transitions": [
                {"from": src, "to": dst, "count": count}
                for (src, dst), count in sorted(transitions.items(), key=order)
            ],
            "by_rule": dict(sorted(by_rule.items(), key=lambda item: (-sum(item[1].values()), item[0]))),
        },
    }


# --- two scans ----------------------------------------------------------------

def _scan_block(scan: Scan) -> dict[str, Any]:
    return {
        "scan_id": scan.id,
        "version": scan.version,
        "commit_sha": scan.commit_sha,
        "created_at": scan.created_at,
        "evidence_version": dict((scan.metadata_json or {}).get("versions") or {}),
    }


def compare_scans(db: Session, scan: Scan, base: Scan, *, include_unchanged: bool = False) -> dict[str, Any]:
    """The decision diff of `scan` against `base`. Raises `DiffError` when they are
    not comparable: another project, a scan type with its own diff, or a scan that
    predates occurrence identity."""
    if scan.project_id != base.project_id:
        raise DiffError("Both scans must belong to the same project.")
    for each in (scan, base):
        if (each.scan_type or "sast") != "sast":
            raise DiffError("The decision diff compares SAST scans; DAST scans have their own diff.")
    base_records, base_missing = records_for_scan(db, base)
    new_records, new_missing = records_for_scan(db, scan)
    for label, each, records, missing in (("base", base, base_records, base_missing),
                                           ("current", scan, new_records, new_missing)):
        if missing and not records:
            raise DiffError(
                f"The {label} scan (v{each.version}) predates occurrence identity: none of its "
                f"{missing} finding(s) has an occurrence id, so it cannot be compared. Rescan it."
            )
    result = diff(base_records, new_records, include_unchanged=include_unchanged)
    return {
        "scan": _scan_block(scan),
        "base": _scan_block(base),
        "commit_differs": (scan.commit_sha or "") != (base.commit_sha or ""),
        "file_level_joins": sum(
            1 for e in result["entries"] if e["identity_quality"] == "file_level"
        ),
        "without_identity": {"base": base_missing, "current": new_missing},
        **result,
    }


# --- Excel --------------------------------------------------------------------

def build_decision_diff_xlsx(result: Mapping[str, Any]) -> bytes:
    """The same data as the JSON: a Summary sheet and one Changes row per entry."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    bold = Font(bold=True)
    rows: list[tuple[str, Any]] = [
        ("Scan", f"v{result['scan']['version']} ({result['scan']['scan_id']})"),
        ("Compared with", f"v{result['base']['version']} ({result['base']['scan_id']})"),
        ("Commit differs", "yes" if result["commit_differs"] else "no"),
        *[(f"Count — {name}", result["counts"][name]) for name in CHANGES],
    ]
    for label, value in rows:
        summary.append([label, value])
    summary.append([])
    summary.append(["Transition", "Count"])
    summary.cell(summary.max_row, 1).font = bold
    for item in result["summary"]["transitions"]:
        summary.append([f"{item['from']} → {item['to']}", item["count"]])
    summary.append([])
    summary.append(["Rule", "Added", "Removed", "Verdict changed", "Issue class changed"])
    summary.cell(summary.max_row, 1).font = bold
    for rule, counts in result["summary"]["by_rule"].items():
        summary.append([rule, counts[ADDED], counts[REMOVED], counts[VERDICT_CHANGED], counts[CLASS_CHANGED]])
    summary.column_dimensions["A"].width = 40

    changes = wb.create_sheet("Changes")
    headers = [
        "Change", "Rule", "File", "Line", "Old verdict", "New verdict", "Old issue class",
        "New issue class", "Old unresolved reason", "New unresolved reason", "Old confidence",
        "New confidence", "Stage", "Reason", "Identity", "Occurrence ID",
    ]
    changes.append(headers)
    for cell in changes[1]:
        cell.font = bold
    for e in result["entries"]:
        old, new = e["old"] or {}, e["new"] or {}
        changes.append([
            e["change"], e["rule_id"], e["file_path"], e["line"],
            old.get("verdict"), new.get("verdict"), old.get("issue_class"), new.get("issue_class"),
            old.get("unresolved_reason"), new.get("unresolved_reason"),
            old.get("confidence"), new.get("confidence"),
            e["stage"], (e["verdict_reason"] or "")[:2000], e["identity_quality"], e["occurrence_id"],
        ])
    changes.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
